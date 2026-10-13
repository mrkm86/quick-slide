#!/usr/bin/env python3
"""image-slide のテスト。`python3 test_logic.py <名前>` で1つ走らせる（run.sh が全部回す）。

本物の Codex は呼ばない。偽物（fake_codex.py）を IMAGE_SLIDE_CODEX で差し、実ホームも偽の CODEX_HOME にする。
"""
from __future__ import annotations

import base64
import json
import os
import pathlib
import shutil
import subprocess
import sys
import tempfile
import time

HERE = pathlib.Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path.insert(0, str(SKILL / "scripts"))
ENTRY = SKILL / "scripts" / "image_slide.py"
SAMPLE = SKILL / "examples" / "sample-pitch.json"
HTML_SLIDE = SKILL.parent / "html-slide"

import compose as C  # noqa: E402
import layout as L  # noqa: E402
import prompt as P  # noqa: E402
import script_io as S  # noqa: E402
import styles as ST  # noqa: E402


def check(cond, msg):
    if not cond:
        raise AssertionError(msg)


def runnable(script: pathlib.Path) -> str:
    """偽物（`.py`）を、コマンドとして起こせる形で返す。Windows は `.py` を直接起こせないので、同じ Python で
    走らせる `.cmd` の包みを一時フォルダに作る（cmd.exe は包みを OEM のコードページで読むので、その字で書く）。
    偽物は `-X utf8` で起こす（本物の codex・claude と同じく、stdin・stdout を UTF-8 で読み書きする）。"""
    if os.name != "nt":
        return str(script)
    cmd = pathlib.Path(tempfile.mkdtemp(prefix="image-slide-test-launch-")) / f"{script.stem}.cmd"
    cmd.write_text(f'@"{sys.executable}" -X utf8 "{script}" %*\r\n', encoding="oem", newline="")
    return str(cmd)


def tmpdir() -> pathlib.Path:
    return pathlib.Path(tempfile.mkdtemp(prefix="image-slide-test."))


def fake_jwt(exp: float) -> str:
    enc = lambda d: base64.urlsafe_b64encode(json.dumps(d).encode()).decode().rstrip("=")  # noqa: E731
    return f"{enc({'alg': 'none'})}.{enc({'exp': int(exp), 'iat': int(time.time())})}.sig"


def fake_home(root: pathlib.Path, age_days=1.0, left_days=5.0) -> pathlib.Path:
    home = root / "codex-home"
    home.mkdir(parents=True, exist_ok=True)
    last = time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime(time.time() - age_days * 86400))
    (home / "auth.json").write_text(json.dumps({
        "auth_mode": "chatgpt", "OPENAI_API_KEY": None, "last_refresh": last,
        "tokens": {"id_token": "x", "access_token": fake_jwt(time.time() + left_days * 86400), "refresh_token": "r",
                   "account_id": "a"}}), encoding="utf-8")
    return home


def run(args, root: pathlib.Path, mode="ok", extra_env=None, home=None):
    """偽物の codex で入口を流す。(rc, stdout, stderr, 呼び出しの記録)"""
    log = root / "calls"
    env = {**os.environ, "IMAGE_SLIDE_CODEX": runnable(HERE / "fake_codex.py"), "FAKE_CODEX_LOG": str(log),
           "FAKE_CODEX_MODE": mode, "CODEX_HOME": str(home or fake_home(root)),
           "IMAGE_SLIDE_CONFIG": str(root / "no-config.yaml"), "OPENAI_API_KEY": "sk-test-should-not-pass",
           "IMAGE_SLIDE_CACHE": str(root / "cache"), "QUICK_SLIDE_HOME": str(root / "home"),
           **(extra_env or {})}
    r = subprocess.run([sys.executable, str(ENTRY), *args], capture_output=True, text=True, env=env, timeout=300, encoding="utf-8")
    calls = [json.loads(p.read_text(encoding="utf-8")) for p in sorted(log.glob("*.json"))] if log.exists() else []
    return r.returncode, r.stdout, r.stderr, calls


def rawdir(root: pathlib.Path, out: pathlib.Path) -> pathlib.Path:
    """その出力フォルダの描いた絵の置き場（キャッシュの区画・ADR 0015）。"""
    import cache as K
    return K.area_path(root / "cache", out)


def write_script(root: pathlib.Path, data: dict) -> pathlib.Path:
    p = root / "deck.json"
    p.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    return p


# --- 受け口（html-slide と同じ台本） -------------------------------------------------------

def t_same_types_as_html_slide():
    if not (HTML_SLIDE / "scripts" / "deck_types.py").is_file():
        print("html-slide が無いので飛ばす")
        return
    sys.path.insert(0, str(HTML_SLIDE / "scripts"))
    import deck as HD
    import deck_types as HT
    check(S.TOP_KEYS == HD.TOP_KEYS, f"一番外のキーがずれた: {S.TOP_KEYS ^ HD.TOP_KEYS}")
    check(S.COMMON_KEYS == HT.COMMON_KEYS, f"共通のキーがずれた: {S.COMMON_KEYS ^ HT.COMMON_KEYS}")
    check(S.THEME_KEYS == HD.THEME_KEYS, "テーマのキーがずれた")
    check(list(S.TYPES) == list(HT.TYPES), f"型の並びがずれた: {set(S.TYPES) ^ set(HT.TYPES)}")
    for k, (keys, chrome, _) in S.TYPES.items():
        _, hkeys, hchrome = HT.TYPES[k]
        check(keys == hkeys, f"{k} のキーがずれた: {keys ^ hkeys}")
        check(chrome == hchrome, f"{k} の帯の有無がずれた")
    check(set(P.TYPE_NOTES) == set(S.TYPES), "中身の形の言葉が全部の型に無い")
    for name in ("default", "forest", "warm"):
        a = (SKILL / "themes" / f"{name}.yaml").read_text(encoding="utf-8").split("\n")
        b = (HTML_SLIDE / "themes" / f"{name}.yaml").read_text(encoding="utf-8").split("\n")
        va = [x for x in a if x and not x.startswith("#")]
        vb = [x for x in b if x and not x.startswith("#")]
        check(va == vb, f"テーマ {name} の値が html-slide とずれた")


def t_sample_passes_both():
    root = tmpdir()
    rc, out, err, _ = run(["check", "--script", str(SAMPLE), "--prompts", str(root / "p")], root)
    check(rc == 0, f"image-slide の check が落ちた: {err}")
    check(len(list((root / "p").glob("slide-*.txt"))) == 5, "描く指示が5枚ぶん出ていない")
    if (HTML_SLIDE / "scripts" / "html_slide.py").is_file():
        r = subprocess.run([sys.executable, str(HTML_SLIDE / "scripts" / "html_slide.py"), "check", "--script", str(SAMPLE)],
                           capture_output=True, text=True, encoding="utf-8")
        check(r.returncode == 0, f"同じ台本が html-slide で通らない: {r.stderr}")


def t_script_errors():
    root = tmpdir()
    cases = [
        ({"title": "x", "logo": "a.png", "slides": [{"type": "statement", "headline": "h"}]}, "引数で渡す"),
        ({"slides": [{"type": "nope"}]}, "知らない型"),
        ({"slides": [{"type": "statement"}]}, "headline が無い"),
        ({"slides": [{"type": "statement", "headline": "h", "colour": 1}]}, "知らないキー"),
        ({"slides": [{"type": "statement", "headline": "h", "_visual": 3}]}, "_visual は文字"),
        ({"slides": [{"type": "closing", "headline": "h"}]}, "items か contact"),
        ({"slides": []}, "slides が無い"),
        ({"slides": [{"type": "photo", "headline": "h", "image": {"path": "missing.png"}}]}, "画像が無い"),
        ({"slides": [{"type": "photo", "headline": "h", "image": "https://example.com/a.png"}]}, "URL"),
    ]
    for data, want in cases:
        p = write_script(root, data)
        rc, _, err, calls = run(["check", "--script", str(p)], root)
        check(rc == 2 and want in err, f"{data} → rc={rc} err={err!r}（{want} を期待）")
        check(not calls, "check で codex が呼ばれた")
    p = write_script(root, {"slides": [{"type": "statement", "headline": "h", "_note": "html-slide も読み飛ばす"}]})
    rc, _, err, _ = run(["check", "--script", str(p)], root)
    check(rc == 0, f"_ で始まるキーで落ちた: {err}")


def t_prompt_content():
    deck = S.load(SAMPLE)
    theme = S.load_theme(None, deck["dir"], {})
    flow = deck["slides"][2]
    b = P.build(deck, flow, theme, style_note="水彩風", band=True)
    t = b["text"]
    for want in ("朝会は3つのことだけを話す", "助けがほしいこと", "このページは 3枚目", "全5枚", "型（中身の形）は flow", "構図: 指定しない", "下端から 7%",
                 "スタイルへの追加の注文: 水彩風", "スタイル「visual」", "輪になって話している", theme["primary"], "1. 朝の15分で、迷いを減らす"):
        check(want in t, f"描く指示に {want!r} が無い")
    for bad in ("朝会で話すのは3つだけです", "talk", "_visual", str(SKILL), "/Users/"):
        check(bad not in t, f"描く指示に入れてはいけないものがある: {bad!r}")
    nb = P.build(deck, flow, theme, band=False)["text"]
    check("下端から" not in nb, "帯が要らないのに帯の指示が入った")
    c = S.content_of({"type": "photo", "headline": "h", "layout": "left", "split": "1/3", "image": {"prompt": "海"}}, [
        {"slot": "image", "path": None, "prompt": "海", "alt": ""}])
    check("layout" not in c and "split" not in c and "海" in c["image"], f"中身から描き方のキーが外れていない: {c}")
    k1 = P.page_key(t, [])
    check(k1 == P.page_key(t, []) and k1 != P.page_key(t + "x", []), "鍵が中身で変わらない／同じ中身で変わる")


def t_image_refs():
    root = tmpdir()
    from PIL import Image
    (root / "images").mkdir()
    Image.new("RGB", (64, 64), "red").save(root / "images" / "s01-image.png")
    Image.new("RGB", (64, 64), "blue").save(root / "ref.png")
    p = write_script(root, {"slides": [{"type": "photo", "headline": "海の写真", "image": {"prompt": "海"}}]})
    deck = S.load(p)
    imgs = S.images_of(deck["slides"][0]["spec"], deck["dir"], 1)
    check(imgs[0]["path"] == (root / "images" / "s01-image.png").resolve(), f"html-slide と同じ置き場を見ていない: {imgs}")
    b = P.build(deck, deck["slides"][0], S.load_theme(None, root, {}), refs_common=1)
    check("添付の参照画像1の内容" in b["text"] and "添付の参照画像2は、資料全体" in b["text"], "参照画像の言い方がずれた")
    rc, _, err, _ = run(["check", "--script", str(p)] + sum((["--ref", str(root / "ref.png")] for _ in range(4)), []), root)
    check(rc == 0, f"枠1＋--ref 4 の5枚で止まった（Codex の上限は5）: {err}")
    rc, _, err, _ = run(["check", "--script", str(p)] + sum((["--ref", str(root / "ref.png")] for _ in range(5)), []), root)
    check(rc == 2 and "5 枚まで" in err, f"参照画像が多すぎても通った: {err}")


def t_characters():
    root = tmpdir()
    from PIL import Image
    (root / "kit").mkdir()
    Image.new("RGBA", (40, 80), (0, 0, 0, 0)).save(root / "kit" / "hana-fullbody-think.png")     # 透過のキャラ
    Image.new("RGB", (40, 80), "green").save(root / "kit" / "pochi-fullbody-surprised.png")
    Image.new("RGB", (64, 64), "red").save(root / "photo.png")
    Image.new("RGB", (64, 64), "blue").save(root / "ref.png")
    hana = {"name": "Hana", "image": "kit/hana-fullbody-think.png", "pose": "考える"}
    pochi = {"name": "Pochi", "image": str(root / "kit" / "pochi-fullbody-surprised.png")}
    data = {"slides": [{"type": "photo", "headline": "海", "image": "photo.png", "_characters": [hana, pochi]},
                       {"type": "statement", "headline": "主張"}]}
    p = write_script(root, data)
    deck = S.load(p)
    th = S.load_theme(None, root, {})
    b = P.build(deck, deck["slides"][0], th, refs_common=1)
    check([c["name"] for c in b["characters"]] == ["Hana", "Pochi"], b["characters"])
    for want in ("添付の参照画像1の内容", "添付の参照画像2は Hana（ポーズ・表情: 考える）", "添付の参照画像3は Pochi\n",
                 "参照画像のまま描き", "添付の参照画像4は、資料全体"):
        check(want in b["text"], f"描く指示に {want!r} が無い:\n{b['text'][-600:]}")
    check(str(root) not in b["text"], "描く指示にファイルのパスが入った")
    check(not P.build(deck, deck["slides"][1], th)["characters"], "キャラの無いページにキャラが付いた")
    # 枚数は枠・キャラ・--ref を合わせて数える（1＋2＋2 = 5 は通る、＋1 で止まる）
    refs = lambda n: sum((["--ref", str(root / "ref.png")] for _ in range(n)), [])  # noqa: E731
    rc, _, err, _ = run(["check", "--script", str(p)] + refs(2), root)
    check(rc == 0, err)
    rc, _, err, _ = run(["check", "--script", str(p)] + refs(3), root)
    check(rc == 2 and "キャラ" in err and "5 枚まで" in err, f"6枚で止まらない: {err}")
    # 並び（枠 → キャラ → --ref）と、透過は白に合成して渡す（元のファイルは触らない）
    out = root / "out"
    rc, _, se, calls = run(["deck", "--script", str(p), "--out", str(out), "--no-pdf"] + refs(1), root)
    check(rc == 0, se)
    ex = sorted(_execs(calls), key=lambda c: c["page"])
    check(len(ex[0]["refs"]) == 4 and len(ex[1]["refs"]) == 1, [c["refs"] for c in ex])
    check(ex[0]["refs"][0].endswith("photo.png") and ex[0]["refs"][2].endswith("pochi-fullbody-surprised.png")
          and ex[0]["refs"][3].endswith("ref.png"), ex[0]["refs"])
    check(ex[0]["ref_modes"][1] == "RGB" and not ex[0]["refs"][1].endswith("hana-fullbody-think.png"),
          f"透過のキャラが白に合成されずに渡った: {ex[0]['refs']} {ex[0]['ref_modes']}")
    check(Image.open(root / "kit" / "hana-fullbody-think.png").mode == "RGBA", "元の絵を書き換えた")
    # 書き方の誤りは読むときに止める
    for bad, word in (([{"image": "kit/hana-fullbody-think.png"}], "name"),
                      ([{"name": "Hana"}], "image"),
                      ([{"name": "Hana", "image": "kit/none.png"}], "絵が無い"),
                      ([{"name": "Hana", "image": "kit/hana-fullbody-think.png", "size": 2}], "知らないキー"),
                      ({"name": "Hana"}, "配列"), ([], "配列")):
        q = write_script(root, {"slides": [{"type": "statement", "headline": "x", "_characters": bad}]})
        rc, _, err, _ = run(["check", "--script", str(q)], root)
        check(rc == 2 and word in err, f"{bad} で止まらない／言い方が違う: {err}")
    # 1枚に凝縮するスタイルでは、全ページのキャラを（同じ絵は1枚だけ）添え、区画に誰が出るかを書く
    data["slides"][1]["_characters"] = [hana]
    p = write_script(root, data)
    rc, so, se, _ = run(["check", "--script", str(p), "--style", "infographic"], root)
    check(rc == 0, se)
    b = P.build_one_page(S.load(p), th, style=ST.load("infographic"))
    check(len(b["characters"]) == 2 and '"登場するキャラ"' in b["text"] and "添付の参照画像2は Hana" in b["text"], b["text"][-500:])


def t_normalize():
    root = tmpdir()
    from PIL import Image
    for size, op in (((1672, 941), "resize"), ((1774, 887), "crop"), ((1024, 1536), "pad"), ((2400, 900), "pad")):
        p = root / f"{size[0]}x{size[1]}.png"
        Image.new("RGB", size, "white").save(p)
        im, info = C.normalize(p, "#ffffff")
        check(im.size == (L.W, L.H), f"{size} → {im.size}")
        check(info["op"] == op, f"{size} は {op} のはずが {info['op']}")


def t_band_check():
    from PIL import Image, ImageDraw
    im = Image.new("RGB", (L.W, L.H), "white")
    check(C.band_check(im)["ok"], "白い帯で NG")
    d = ImageDraw.Draw(im)
    d.rectangle((300, L.BAND_TOP + 10, 1500, L.H), fill=(20, 40, 90))
    r = C.band_check(im)
    check(not r["ok"] and r["share"] > 0.3, f"帯に物があるのに OK: {r}")
    dark = Image.new("RGB", (L.W, L.H), (10, 20, 40))
    r = C.band_check(dark)
    check(r["ok"] and r["dark"], f"無地の暗い帯は OK で、文字は白にする: {r}")
    check(C.text_color(r, {"ink": "#1a2b4c"}, False) == "#ffffff", "暗い帯の文字が白でない")


def t_logo_rules_match_html_slide():
    """ロゴの寸法と規則の数字が html-slide と同じ（ADR 0008。片方だけ変えたら落ちる）。"""
    if not (HTML_SLIDE / "scripts" / "deck.py").is_file():
        print("html-slide が無いので飛ばす")
        return
    sys.path.insert(0, str(HTML_SLIDE / "scripts"))
    import deck as HD
    for k in ("LOGO_BOX_W", "LOGO_BOX_H", "LOGO_RIGHT", "LOGO_GAP", "LOGO_MIN_SIDE", "LOGO_RING", "LOGO_TOP_MID_Y", "LOGO_POSITIONS",
              "WATERMARK_PX", "WATERMARK_COLOR", "WATERMARK_BOTTOM"):
        check(getattr(L, k) == getattr(HD, k), f"{k} が html-slide とずれた: {getattr(L, k)} と {getattr(HD, k)}")
    for w, h in ((444, 59), (370, 100), (100, 20), (329, 44), (2000, 2000)):
        a, b = L.logo_geometry(w, h), HD.logo_geometry(w, h)
        for k in ("x", "w", "h", "scale", "warning"):
            check(a[k] == b[k], f"{w}×{h} の {k} が html-slide とずれた: {a[k]} と {b[k]}")


def t_logo_place():
    from PIL import Image, ImageDraw
    root = tmpdir()
    # 透明な余白を切ってから枠に収める（透明な余白の多い横長のロゴの形）・縮めるだけ
    lg = Image.new("RGBA", (595, 365), (0, 0, 0, 0))
    ImageDraw.Draw(lg).rectangle((113, 132, 482, 231), fill=(20, 20, 20, 255))
    lg.save(root / "pad.png")
    logo = C.load_logo(root / "pad.png", root / "out")
    check((logo["w"], logo["h"], logo["trim"]) == (163, 44, [113, 132, 483, 232]), f"余白を切っていない: {logo}")
    check(Image.open(logo["file"]).size == (370, 100), "切った写しを書いていない")
    Image.new("RGBA", (100, 20), (20, 20, 20, 255)).save(root / "small.png")
    _, g = C.read_logo(root / "small.png")
    check((g["w"], g["h"]) == (100, 20) and g["warning"] and "小さい" in g["warning"], f"引き伸ばした／警告が無い: {g}")
    # 白い地: 右下に、白の地を敷いて置く
    white = Image.new("RGB", (L.W, L.H), (254, 254, 254))
    r = C.place_logo(white, logo, band=None)
    check(r["corner"] == "bottom-right" and r["x"] == L.W - 120 - 163 and r["plate"] == "#fefefe", f"白い地: {r}")
    # 右下に絵: 左下へ（x は本文の余白）
    busy = white.copy()
    d = ImageDraw.Draw(busy)
    for i in range(0, 700, 12):
        d.line((1200 + i, 900, 1200 + i, L.H), fill=(30, 90, 40), width=5)
    r = C.place_logo(busy, logo, band=None)
    check(r["corner"] == "bottom-left" and r["x"] == 120 and r["plate"] == "#fefefe", f"右下がふさがったのに左下でない: {r}")
    # 両方ふさがる → 置かない
    for i in range(0, 700, 12):
        d.line((i, 900, i, L.H), fill=(30, 90, 40), width=5)
    check(C.place_logo(busy, logo, band=None) == {"off": True}, "両方ふさがったのに置いた")
    # 帯を頼んだ本文のページは、帯が空いていれば右下（帯の色で地を敷く）・板を敷いたら板の上の右下（地は要らない）
    r = C.place_logo(white, logo, band="ok")
    check(r["corner"] == "bottom-right" and r["plate"] == "#fefefe", f"帯の上: {r}")
    r = C.place_logo(busy, logo, band="plate")
    check(r["corner"] == "bottom-right" and r["plate"] is None, f"帯の板の上: {r}")
    # 平らだが1色ではない地（右へ明るくなるグラデーション・ぼけた写真の代わり）: 右下は避けて、1色の左下へ
    grad = Image.new("RGB", (L.W, L.H))
    gd = ImageDraw.Draw(grad)
    for x in range(L.W):
        v = 180 + max(0, x - 1320) * 60 // 600
        gd.line((x, 0, x, L.H), fill=(v, v - 30, v - 60))
    r = C.place_logo(grad, logo, band=None)
    check(r["corner"] == "bottom-left" and r["plate"] == "#b49678", f"グラデーションの上に置いた: {r}")
    # --logo-position: 左下は 左下 → 右下、右上は 右上 → 右下（帯の中の右下は帯の検査で済ませる）
    r = C.place_logo(white, logo, band=None, position="bottom-left")
    check(r["corner"] == "bottom-left" and r["x"] == 120, f"左下の指定: {r}")
    r = C.place_logo(white, logo, band=None, position="top-right")
    check(r["corner"] == "top-right" and (r["x"], r["y"]) == (L.W - 120 - 163, 24), f"右上の指定: {r}")
    top = white.copy()
    td = ImageDraw.Draw(top)
    for i in range(1500, L.W, 12):
        td.line((i, 0, i, 120), fill=(30, 90, 40), width=5)
    r = C.place_logo(top, logo, band="ok", position="top-right")
    check(r["corner"] == "bottom-right" and r["plate"] == "#fefefe", f"右上がふさがったのに右下へ逃げない: {r}")
    both = top.copy()
    bd = ImageDraw.Draw(both)
    for i in range(1300, L.W, 12):
        bd.line((i, 900, i, L.H), fill=(30, 90, 40), width=5)
    check(C.place_logo(both, logo, band=None, position="top-right") == {"off": True}, "右上 → 右下 のあと左下へ逃げた")
    # 全面に描かせたページ（bleed）: 1色の角が無ければ、平らな角（グラデーション）に地を敷かずに重ねる
    r = C.place_logo(grad, logo, band=None, position="top-right", bleed=True)
    check(r["corner"] == "top-right" and r["plate"] is None and r["variant"] == "normal", f"全面のページで平らな角に置いていない: {r}")
    # どの角もでこぼこでも、明るくて読めるなら指定の角にそのまま重ねる
    r = C.place_logo(both, logo, band=None, position="top-right", bleed=True)
    check(r["corner"] == "top-right" and r["plate"] is None and not r["sticker"], f"全面のページで置かなかった: {r}")
    # 暗くてでこぼこ（夜の岩）: ロゴは加工せず、ステッカー（余白つきの角丸の地・地の色）を敷く
    night = Image.new("RGB", (L.W, L.H), (15, 20, 30))
    nd = ImageDraw.Draw(night)
    for i in range(0, L.W, 12):
        nd.line((i, 0, i, L.H), fill=(70, 80, 60), width=5)
    r = C.place_logo(night, logo, band=None, bleed=True, background="#fefefe")
    pad = L.LOGO_STICKER_PAD
    check(r["corner"] == "bottom-right" and r["sticker"] and r["plate"] == "#fefefe" and r["variant"] == "normal"
          and r["plate_box"] == (r["x"] - pad, r["y"] - pad, r["w"] + 2 * pad, r["h"] + 2 * pad), f"暗い絵の上のステッカー: {r}")
    check(C.logo_file(logo, r) == logo["file"] and logo["white_file"] is None, "ロゴを加工した（白抜きを作った）")
    # 暗い1色の地（濃い面）でも黒いロゴは沈むので、地の色の地ではなくステッカー
    # 帯に淡い枠が入りこんだ本文のページ: 置くが、地は敷かない（枠の角を四角く欠かない）
    cream = white.copy()
    ImageDraw.Draw(cream).rounded_rectangle((900, 700, 1850, 1030), radius=24, fill=(255, 248, 236))
    r = C.place_logo(cream, logo, band="ok")
    check(r["corner"] == "bottom-right" and r["plate"] is None and not r["sticker"], f"淡い枠の上に地を敷いた: {r}")
    navy = Image.new("RGB", (L.W, L.H), (10, 30, 60))
    r = C.place_logo(navy, logo, band="ok", background="#fefefe")
    check(r["sticker"] and r["plate"] == "#fefefe", f"暗い帯の上: {r}")
    check(C.place_logo(navy, logo, band=None) == {"off": True}, "帯の無いページの暗い1色の角に、沈むロゴを置いた")
    # 全面の絵の上の文字の色は、それぞれの下の明るさで決める
    dark = Image.new("RGB", (L.W, L.H), (250, 250, 250))
    ImageDraw.Draw(dark).rectangle((0, 980, 700, L.H), fill=(20, 30, 40))
    ink = C.inks(dark, theme={"ink": "#111111"}, band=C.band_check(dark), plate=False, bleed=True,
                 boxes={"footer": (120, 1031, 300, 22), "number": (1700, 1031, 20, 22)})
    check(ink == {"footer": "#ffffff", "number": "#111111"}, f"文字の色: {ink}")
    # 確認用の絵: 地とロゴが焼かれ、ページ番号は中央下
    r = C.place_logo(white, logo, band=None)
    shown = C.overlay(white, theme={"background": "#ffffff", "ink": "#111111"}, plate=False, footer="", number="3",
                      ink={"footer": "#111111", "number": "#111111"}, logo=logo, place=r, font_file=None)
    px = shown.getpixel((r["x"] + 50, r["y"] + 22))
    check(px[0] < 60, f"ロゴが焼かれていない: {px}")
    nx = L.page_number_box(10)
    check(nx[0] == L.W // 2 - 5, f"ページ番号が中央下にない: {nx}")


def t_logo_deck():
    from PIL import Image
    from pptx import Presentation
    root = tmpdir()
    Image.new("RGBA", (444, 59), (200, 30, 30, 255)).save(root / "logo.png")
    out = root / "out"
    # 1回目の絵は下端に横長の物がある（右下も左下もふさがる）。描き直さない
    rc, so, se, _ = run(["deck", "--script", str(SAMPLE), "--style", "standard", "--out", str(out), "--no-pdf", "--retries", "0",
                         "--logo", str(root / "logo.png")], root, mode="band_first")
    check(rc == 0, se)
    b = json.loads((out / "build.json").read_text(encoding="utf-8"))
    pages = b["logo"]["pages"]
    # 表紙は帯を頼まない（全面の絵を下で途切れさせない）→ 角がふさがっているので置かない。本文は帯の板の上の右下
    check(pages[0] == {"no": 1, "off": True}, f"表紙: {pages[0]}")
    check(all(p.get("corner") == "bottom-right" and p["plate"] is None for p in pages[1:]), f"本文: {pages[1:]}")
    check((b["logo"]["w"], b["logo"]["h"]) == (329, 44), f"329×44 に収めていない: {b['logo']}")
    check("下端から" not in (out / "prompts" / "slide-01.txt").read_text(encoding="utf-8"), "--logo で表紙にも帯を頼んだ")
    check("下端から" in (out / "prompts" / "slide-02.txt").read_text(encoding="utf-8"), "本文に帯を頼んでいない")
    check("ロゴは置かない" in se, "置かなかったことを知らせていない")
    prs = Presentation(str(out / "deck.pptx"))
    names = [[s.name for s in sl.shapes] for sl in prs.slides]
    check("ロゴ" not in names[0] and all("ロゴ" in n for n in names[1:]), f"PPTX のロゴ: {names}")
    # 白い絵: 表紙も右下に、白の地つきで置く。地とロゴは別の部品
    out2 = root / "out2"
    rc, so, se, _ = run(["deck", "--script", str(SAMPLE), "--style", "standard", "--out", str(out2), "--no-pdf", "--logo", str(root / "logo.png")], root)
    b = json.loads((out2 / "build.json").read_text(encoding="utf-8"))
    check(rc == 0 and all(p.get("corner") == "bottom-right" and p["plate"] == "#ffffff" for p in b["logo"]["pages"]), f"{b.get('logo')}")
    sl = Presentation(str(out2 / "deck.pptx")).slides[1]
    shapes = {s.name: s for s in sl.shapes}
    check(list(shapes).index("ロゴの地") < list(shapes).index("ロゴ"), "ロゴの地がロゴの下にない")
    check(shapes["ロゴの地"].left == shapes["ロゴ"].left and shapes["ロゴの地"].width == shapes["ロゴ"].width, "地とロゴの大きさが違う")
    mid = shapes["ページ番号"].left + shapes["ページ番号"].width // 2
    check(abs(mid - 1920 * 6350 // 2) < 6350 and shapes["フッタ"].left + shapes["フッタ"].width <= (1920 // 2 - 30) * 6350,
          "ページ番号が中央下にない／フッタが番号に届く")
    # 左下の指定: フッタはロゴの右へ、ページ番号は中央下のまま
    out3 = root / "out3"
    rc, so, se, _ = run(["deck", "--script", str(SAMPLE), "--style", "standard", "--out", str(out3), "--no-pdf", "--logo", str(root / "logo.png"),
                         "--logo-position", "bottom-left"], root)
    b = json.loads((out3 / "build.json").read_text(encoding="utf-8"))
    check(rc == 0 and b["logo"]["position"] == "bottom-left" and all(p.get("corner") == "bottom-left" for p in b["logo"]["pages"]), f"{b.get('logo')} {se}")
    shapes = {s.name: s for s in Presentation(str(out3 / "deck.pptx")).slides[1].shapes}
    check(shapes["フッタ"].left > shapes["ロゴ"].left + shapes["ロゴ"].width, "フッタがロゴの右にない")
    check(abs(shapes["ページ番号"].left + shapes["ページ番号"].width // 2 - 1920 * 6350 // 2) < 6350, "ページ番号が中央下でない")
    rc, so, se, _ = run(["check", "--script", str(SAMPLE), "--logo", str(root / "logo.png"), "--logo-position", "top-left"], root)
    check(rc == 2, f"左上が通った: {rc}")
    # 受け口は html-slide と同じ（PNG・JPEG だけ）
    Image.new("RGBA", (444, 59), (200, 30, 30, 255)).save(root / "logo.webp")
    rc, so, se, _ = run(["check", "--script", str(SAMPLE), "--logo", str(root / "logo.webp")], root)
    check(rc == 2 and "PNG か JPEG" in se, f"WebP のロゴを通した: {rc} {se}")


def t_visual_bleed():
    """visual は本文のページも全面に描かせる（帯を頼まない・帯の検査で描き直さない・板を敷かない）。standard は帯のまま。"""
    from PIL import Image
    root = tmpdir()
    Image.new("RGBA", (444, 59), (200, 30, 30, 255)).save(root / "logo.png")
    out = root / "out"
    rc, so, se, calls = run(["deck", "--script", str(SAMPLE), "--out", str(out), "--no-pdf", "--logo", str(root / "logo.png")],
                            root, mode="band_first")
    check(rc == 0, se)
    t = (out / "prompts" / "slide-02.txt").read_text(encoding="utf-8")
    check("全面に使ってよい" in t and "1色の無地にする" not in t, "visual の本文に帯を頼んだ")
    check(not any(c["n"] == 2 for c in _execs(calls)), "全面のページを帯の検査で描き直した")
    b = json.loads((out / "build.json").read_text(encoding="utf-8"))
    check(not any(s["plate"] for s in b["slides"]), "全面のページに帯の板を敷いた")
    check(all(s["band"] is None for s in b["slides"]), "全面のページを帯として記録した")
    st = P.build(S.load(SAMPLE), S.load(SAMPLE)["slides"][1], S.load_theme(None, SAMPLE.parent, {}),
                 style=ST.load("standard"), band=True)["text"]
    check("1色の無地にする" in st, "standard の本文に帯を頼んでいない")
    # 公式の白いロゴ（--logo-white）は余白を切って使う・--logo なしでは止める
    lw = Image.new("RGBA", (500, 120), (0, 0, 0, 0))
    lw.paste(Image.new("RGBA", (444, 59), (255, 255, 255, 255)), (20, 30))
    lw.save(root / "white.png")
    Image.new("RGBA", (444, 59), (20, 20, 20, 255)).save(root / "dark.png")      # 黒いロゴ（暗い地に沈む）
    logo = C.load_logo(root / "dark.png", root / "w", root / "white.png")
    check(Image.open(logo["white_file"]).size == (444, 59), "白いロゴの余白を切っていない")
    night = Image.new("RGB", (L.W, L.H), (15, 20, 30))
    r = C.place_logo(night, logo, band="ok", background="#fefefe")
    check(r["variant"] == "white" and not r["sticker"] and C.logo_file(logo, r) == logo["white_file"], f"公式の白いロゴを使っていない: {r}")
    Image.new("RGBA", (300, 100), (255, 255, 255, 255)).save(root / "white-sq.png")
    rc, so, se, _ = run(["check", "--script", str(SAMPLE), "--logo", str(root / "logo.png"), "--logo-white", str(root / "white-sq.png")], root)
    check(rc == 2 and "縦横比" in se, f"縦横比の違う白いロゴを通した: {se}")
    rc, so, se, _ = run(["check", "--script", str(SAMPLE), "--logo-white", str(root / "white.png")], root)
    check(rc == 2 and "--logo と一緒" in se, f"--logo なしの --logo-white を通した: {se}")


def t_watermark():
    """--watermark: 全ページ（表紙も）の中央下の端に固定のグレーで焼く・PPTX の絵にも焼く。
    下の行は、フッタか番号があるページだけ上げる（透かしの有無では動かない・表紙のロゴやフッタの無い資料のロゴは帯の中心）。"""
    from PIL import Image
    from pptx import Presentation
    import io
    root = tmpdir()
    Image.new("RGBA", (444, 59), (200, 30, 30, 255)).save(root / "logo.png")
    gray = tuple(int(L.WATERMARK_COLOR[i:i + 2], 16) for i in (1, 3, 5))

    def marks(im):
        """中央下の端の、透かしの色に近い画素の数と、その一番下の y。"""
        im = im.convert("RGB")
        ys = [y for y in range(L.H - 40, L.H) for x in range(L.W // 2 - 120, L.W // 2 + 120)
              if sum(abs(a - b) for a, b in zip(im.getpixel((x, y)), gray)) < 12]   # 灰色でも青みのない文字のにじみは外す
        return len(ys), max(ys, default=None)

    out = root / "out"
    rc, so, se, _ = run(["deck", "--script", str(SAMPLE), "--style", "standard", "--out", str(out), "--no-pdf",
                         "--logo", str(root / "logo.png"), "--watermark", "example.com"], root)
    check(rc == 0, se)
    b = json.loads((out / "build.json").read_text(encoding="utf-8"))
    check(b.get("watermark") == "example.com", f"build.json に透かしが無い: {b.get('watermark')}")
    n = len(b["slides"])
    for no in range(1, n + 1):
        cnt, low = marks(Image.open(out / "slides" / f"slide-{no:02d}.png"))
        check(cnt > 5 and L.H - L.WATERMARK_BOTTOM - 4 <= low <= L.H - L.WATERMARK_BOTTOM, f"{no}枚目に透かしが無い・位置が違う: {cnt} {low}")
    check(marks(Image.open(rawdir(root, out) / "slide-02.png").resize((L.W, L.H)))[0] == 0, "描いた絵（キャッシュ）に透かしを焼いた")
    check(not (out / "base").exists(), "そろえた絵（base）を出力フォルダに残した")
    prs = Presentation(str(out / "deck.pptx"))
    for sl in prs.slides:
        pic = next(sh for sh in sl.shapes if sh.name.startswith("絵 "))
        check(marks(Image.open(io.BytesIO(pic.image.blob)))[0] > 5, "PPTX の絵に透かしが焼かれていない")
    check(b["logo"]["y"] == round(L.ROW_MID_RAISED - 22) and b["logo"]["y"] >= L.BAND_TOP, f"ロゴの行が上がっていない: {b['logo']}")
    shapes = {sh.name: sh for sh in prs.slides[1].shapes}
    no_top, no_h = shapes["ページ番号"].top // 6350, shapes["ページ番号"].height // 6350      # 文字の枠（上下に 8px の余白つき）
    check(no_top + no_h // 2 == round(L.ROW_MID_RAISED), f"ページ番号の行が上がっていない: {no_top}+{no_h}/2")
    check(no_top + no_h < L.H - L.WATERMARK_BOTTOM - L.WATERMARK_PX, "ページ番号の枠が透かしにかかる")
    tops = [round(next(sh for sh in sl.shapes if sh.name == "ロゴ").top / 6350) for sl in prs.slides]
    check(tops[0] == round(L.LOGO_MID_Y - 22) and tops[1] == round(L.ROW_MID_RAISED - 22),
          f"表紙（下の文字の行が無い）のロゴまで上げた・本文を上げていない: {tops}")
    check("透かし「example.com」" in (out / "inspect.md").read_text(encoding="utf-8"), "inspect.md に透かしが無い")
    # 付けなくても行の高さは同じ（フッタ・番号があるので上がる）・透かしは出ない
    out2 = root / "out2"
    rc, so, se, _ = run(["deck", "--script", str(SAMPLE), "--style", "standard", "--out", str(out2), "--no-pdf", "--logo", str(root / "logo.png")], root)
    b2 = json.loads((out2 / "build.json").read_text(encoding="utf-8"))
    check(rc == 0 and b2["logo"]["y"] == b["logo"]["y"] and "watermark" not in b2, f"透かしの有無で行が動いた: {b2.get('logo')}")
    check(marks(Image.open(out2 / "slides" / "slide-02.png"))[0] == 0, "透かしなしで透かしが出た")
    # フッタも番号も無い資料（週報のポスターなど）は、ロゴだけなので帯の中心のまま
    data = json.loads(SAMPLE.read_text(encoding="utf-8"))
    data.pop("footer", None)
    data["page_numbers"] = False
    out3 = root / "out3"
    rc, so, se, _ = run(["deck", "--script", str(write_script(root, data)), "--style", "standard", "--out", str(out3), "--no-pdf",
                         "--logo", str(root / "logo.png"), "--watermark", "example.com"], root)
    b3 = json.loads((out3 / "build.json").read_text(encoding="utf-8"))
    tops = [round(next(sh for sh in sl.shapes if sh.name == "ロゴ").top / 6350) for sl in Presentation(str(out3 / "deck.pptx")).slides]
    check(rc == 0 and b3["logo"]["y"] == round(L.LOGO_MID_Y - 22) and set(tops) == {round(L.LOGO_MID_Y - 22)},
          f"フッタも番号も無いのにロゴを上げた: {b3.get('logo')} {tops} {se}")


def t_page_numbers_default_off():
    """ページ番号は既定でオフ（台本に page_numbers: true と書いたときだけ・html-slide と同じ）。"""
    root = tmpdir()
    data = json.loads(SAMPLE.read_text(encoding="utf-8"))
    data.pop("page_numbers", None)
    deck = S.load(write_script(root, data))
    check(deck["page_numbers"] is False, "既定でページ番号が出る")
    check(S.load(SAMPLE)["page_numbers"] is True, "page_numbers: true が効かない")


# --- 認証と並列（偽物の codex） --------------------------------------------------------------

def _execs(calls):
    """終わりまで走った exec（始まりの記録は除く）。"""
    return [c for c in calls if c["cmd"] == "exec" and "phase" not in c]


def _starts(calls):
    return [c for c in calls if c["cmd"] == "exec" and c.get("phase") == "start"]


def max_concurrency(ex) -> int:
    """始まりと終わりの時刻から、同時に走っていた本数の最大。"""
    ev = sorted([(c["t0"], 1) for c in ex] + [(c["t1"], -1) for c in ex], key=lambda x: (x[0], x[1]))
    cur = best = 0
    for _, d in ev:
        cur += d
        best = max(best, cur)
    return best


def t_deck_parallel():
    root = tmpdir()
    home = fake_home(root)
    before = (home / "auth.json").read_bytes()
    out = root / "out"
    rc, so, se, calls = run(["deck", "--script", str(SAMPLE), "--out", str(out), "--parallel", "5"], root, home=home)
    check(rc == 0, f"rc={rc}: {se}")
    logins = [c for c in calls if c["cmd"] == "login-status"]
    ex = _execs(calls)
    check(len(logins) == 1, f"認証の確認が1回でない: {len(logins)}")
    check(len(ex) == 5, f"5枚なのに exec が {len(ex)} 回")
    conc = max_concurrency(ex)
    check(conc >= 4, f"並列になっていない（同時に走ったのは最大 {conc} 本）")
    check(not any(c["api_key"] for c in ex), "OPENAI_API_KEY が codex に渡った（従量課金になる）")
    check(all(c["sandbox"] == "read-only" for c in ex), "sandbox が read-only でない")
    check(all(c["home"] != str(home) for c in ex), "ワーカーが実ホームで走った")
    check(len({c["home"] for c in ex}) == 5, "ワーカーが隔離ホームを共有した")
    check(not any(pathlib.Path(c["home"]).exists() for c in ex), "隔離ホームが消えていない（auth.json のコピーが残る）")
    check((home / "auth.json").read_bytes() == before, "実 auth.json が書き換わった")
    for f in ("deck.pptx", "deck.pdf", "build.json", "inspect.md"):
        check((out / f).is_file(), f"{f} が無い")
    check(len(list((out / "slides").glob("slide-*.png"))) == 5, "確認用の PNG が5枚無い")
    b = json.loads((out / "build.json").read_text(encoding="utf-8"))
    check(b["auth"]["checked"] == 1 and all(s["status"] == "generated" for s in b["slides"]), f"build.json: {b['auth']}")
    check(all(s["normalize"]["op"] == "resize" for s in b["slides"]), "1672x941 が伸縮でそろっていない")


def t_deck_cache_and_redo():
    root = tmpdir()
    home = fake_home(root)
    out = root / "out"
    data = json.loads(SAMPLE.read_text(encoding="utf-8"))
    p = write_script(root, data)
    rc, _, se, _ = run(["deck", "--script", str(p), "--out", str(out), "--no-pdf"], root, home=home)
    check(rc == 0, se)
    shutil.rmtree(root / "calls")
    rc, so, se, calls = run(["deck", "--script", str(p), "--out", str(out), "--no-pdf"], root, home=home)
    check(rc == 0 and not calls, f"中身が同じなのに codex が呼ばれた（認証の確認も要らない）: {calls}")
    check("cached=5" in so, so)
    data["slides"][1]["headline"] = "困りごとの半分は、前の日に分かっていた"
    p = write_script(root, data)
    rc, so, se, calls = run(["deck", "--script", str(p), "--out", str(out), "--no-pdf"], root, home=home)
    ex = _execs(calls)
    check(rc == 0 and [c["page"] for c in ex] == [2], f"変えた2枚目だけを描き直していない: {[c['page'] for c in ex]}")
    shutil.rmtree(root / "calls")
    rc, so, se, calls = run(["deck", "--script", str(p), "--out", str(out), "--no-pdf", "--redo", "4-5"], root, home=home)
    check(sorted(c["page"] for c in _execs(calls)) == [4, 5], "--redo 4-5 が効いていない")
    rc, _, se, _ = run(["deck", "--script", str(p), "--out", str(out), "--redo", "9"], root, home=home)
    check(rc == 2 and "範囲の外" in se, se)




def t_deck_keep():
    """--keep: 置いてある絵を、描く指示（鍵）が変わっていても描き直さずに使う。鍵は前のまま（付けなければ描き直す）・--redo が勝つ。"""
    root = tmpdir()
    home = fake_home(root)
    out = root / "out"
    data = json.loads(SAMPLE.read_text(encoding="utf-8"))
    p = write_script(root, data)
    rc, _, se, _ = run(["deck", "--script", str(p), "--out", str(out), "--no-pdf"], root, home=home)
    check(rc == 0, se)
    area = rawdir(root, out)
    before = (area / "slide-02.png").read_bytes()
    shutil.rmtree(root / "calls")
    data["slides"][1]["headline"] = "困りごとの半分は、前の日に分かっていた"
    p = write_script(root, data)
    rc, so, se, calls = run(["deck", "--script", str(p), "--out", str(out), "--no-pdf", "--keep", "2"], root, home=home)
    check(rc == 0 and not calls, f"--keep のページを描き直した: {calls}")
    check((area / "slide-02.png").read_bytes() == before, "--keep のページの絵が替わった")
    b = json.loads((out / "build.json").read_text(encoding="utf-8"))
    check(b["slides"][1].get("kept") is True and b["slides"][1]["status"] == "cached", b["slides"][1])
    rc, so, se, calls = run(["deck", "--script", str(p), "--out", str(out), "--no-pdf", "--keep", "2", "--redo", "2"], root, home=home)
    check(rc == 0 and [c["page"] for c in _execs(calls)] == [2], "--redo が --keep に勝っていない")
    rc, _, se, _ = run(["deck", "--script", str(p), "--out", str(out), "--keep", "9"], root, home=home)
    check(rc == 2 and "--keep" in se and "範囲の外" in se, se)


def t_cache_area():
    """描いた絵はキャッシュの区画に置き、出力フォルダに raw/・base/ を残さない。前の形の raw/ は区画へ移して描き直さない（ADR 0015）。"""
    import cache as K
    root = tmpdir()
    home = fake_home(root)
    out = root / "out"
    p = write_script(root, json.loads(SAMPLE.read_text(encoding="utf-8")))
    rc, _, se, _ = run(["deck", "--script", str(p), "--out", str(out), "--no-pdf"], root, home=home)
    check(rc == 0, se)
    area = rawdir(root, out)
    check(not (out / "raw").exists() and not (out / "base").exists(), f"出力フォルダに中間物が残った: {sorted(x.name for x in out.iterdir())}")
    check(len(list(area.glob("slide-*.png"))) == 5 and len(list(area.glob("slide-*.json"))) == 5, f"区画に絵と鍵が無い: {list(area.iterdir())}")
    check(json.loads((area / K.STAMP).read_text(encoding="utf-8"))["owner"] == str(out.resolve()), "区画の印に持ち主が無い")
    check(json.loads((out / "build.json").read_text(encoding="utf-8"))["cache"] == str(area), "build.json に区画の場所が無い")
    check(len(list((out / "slides").glob("slide-*.png"))) == 5 and (out / "deck.pptx").is_file(), "ページの画像か PPTX が無い")
    # 前の形: 出力フォルダの raw/ と base/ → raw/ は区画へ移して使い回す・base/ は消す
    (out / "raw" / "history").mkdir(parents=True)
    for f in area.glob("slide-*"):
        f.replace(out / "raw" / f.name)
    (area / "history").mkdir(exist_ok=True)
    (out / "base").mkdir()
    (out / "base" / "slide-01.png").write_bytes(b"x")
    shutil.rmtree(root / "calls")
    rc, so, se, calls = run(["deck", "--script", str(p), "--out", str(out), "--no-pdf"], root, home=home)
    check(rc == 0 and not calls and "cached=5" in so, f"前の形の raw/ を使い回さずに描いた: {so} {se} {calls}")
    check(not (out / "raw").exists() and not (out / "base").exists(), "前の形の raw/・base/ が残った")
    check(len(list(area.glob("slide-*.png"))) == 5, "前の形の絵を区画へ移していない")


def t_cache_trim():
    """上限を超えたら、最後に使った時刻の古い区画から消す。今の区画・少し前に使った区画・印の無いフォルダは消さない。0 は消さない。
    cache.py は html-slide と同じ。"""
    import cache as K
    base = tmpdir() / "cache"
    env = {"IMAGE_SLIDE_CACHE": str(base)}
    old_env = {k: os.environ.get(k) for k in env}
    os.environ.update(env)
    try:
        mb = 1024 * 1024
        made = {}
        for name, age, size in (("old", 9000, 3), ("mid", 5000, 3), ("cur", 8000, 3), ("recent", 60, 3)):
            d = K.area("image-slide", {}, base.parent / name)
            (d / "slide-01.png").write_bytes(b"x" * size * mb)
            t = time.time() - age
            os.utime(d / K.STAMP, (t, t))
            made[name] = d
        stray = base / "not-ours"
        stray.mkdir()
        (stray / "big.bin").write_bytes(b"x" * 4 * mb)
        check(K.trim("image-slide", {"cache_max_mb": 0}, made["cur"]) == [], "上限 0 で消した")
        check(K.trim("image-slide", {"cache_max_mb": 100}, made["cur"]) == [], "上限の内なのに消した")
        gone = K.trim("image-slide", {"cache_max_mb": 10}, made["cur"])
        check(gone == [str((base.parent / "old").resolve())], f"古い区画から1つだけ消していない: {gone}")
        check(made["cur"].is_dir() and made["recent"].is_dir() and made["mid"].is_dir() and stray.is_dir(), "消してはいけないものを消した")
        gone = K.trim("image-slide", {"cache_max_mb": 1}, made["cur"])
        check(gone == [str((base.parent / "mid").resolve())] and made["cur"].is_dir() and made["recent"].is_dir() and stray.is_dir(),
              f"今の区画・少し前の区画・印の無いフォルダまで消した: {gone}")
    finally:
        for k, v in old_env.items():
            os.environ.pop(k, None) if v is None else os.environ.__setitem__(k, v)
    check(K.area_path(base, base.parent / "a b/日本語") != K.area_path(base, base.parent / "a_b/日本語"), "別のフォルダが同じ区画になった")
    check((SKILL / "scripts" / "cache.py").read_bytes() == (HTML_SLIDE / "scripts" / "cache.py").read_bytes(),
          "cache.py が html-slide とずれた（同じファイルを写す）")


def t_places():
    """置き場の根っこは QUICK_SLIDE_HOME（既定 ~/.quick-slide）。設定・--out を省いた出力はその下の engines/<スキル>/。
    キャッシュは ~/Library/Caches/quick-slide/<スキル>（macOS）。環境変数と設定が優先する。places.py は html-slide と同じ。"""
    import places as PL
    keys = ("QUICK_SLIDE_HOME", "IMAGE_SLIDE_CONFIG", "IMAGE_SLIDE_CACHE")
    old_env = {k: os.environ.get(k) for k in keys}
    try:
        for k in keys:
            os.environ.pop(k, None)
        check(PL.home() == pathlib.Path.home() / ".quick-slide", f"既定の根っこが ~/.quick-slide でない: {PL.home()}")
        root = tmpdir()
        os.environ["QUICK_SLIDE_HOME"] = str(root)
        check(PL.config_path("image-slide") == root / "engines" / "image-slide" / "config.yaml", "設定が根っこの下にない")
        check(PL.load_config("image-slide") == {}, "設定が無いのに空でない")
        check(PL.out_root("image-slide", {}) == root / "engines" / "image-slide" / "output", "出力が根っこの下にない")
        check(PL.out_root("image-slide", {"out": str(root / "o")}) == root / "o", "設定の out が効かない")
        if sys.platform == "darwin":
            check(PL.cache_root("image-slide", {}) == pathlib.Path.home() / "Library" / "Caches" / "quick-slide" / "image-slide",
                  f"キャッシュの既定がずれた: {PL.cache_root('image-slide', {})}")
        check(PL.cache_root("image-slide", {"cache_dir": str(root / "c")}) == root / "c", "設定の cache_dir が効かない")
        cfg = root / "elsewhere.yaml"
        cfg.write_text("out: " + str(root / "x") + "\n", encoding="utf-8")
        os.environ["IMAGE_SLIDE_CONFIG"] = str(cfg)
        check(PL.load_config("image-slide") == {"out": str(root / "x")}, "IMAGE_SLIDE_CONFIG が効かない")
        os.environ["IMAGE_SLIDE_CACHE"] = str(root / "env-cache")
        check(PL.cache_root("image-slide", {"cache_dir": str(root / "c")}) == root / "env-cache", "IMAGE_SLIDE_CACHE が設定より弱い")
    finally:
        for k, v in old_env.items():
            os.environ.pop(k, None) if v is None else os.environ.__setitem__(k, v)
    check((SKILL / "scripts" / "places.py").read_bytes() == (HTML_SLIDE / "scripts" / "places.py").read_bytes(),
          "places.py が html-slide とずれた（同じファイルを写す）")
    check((SKILL / "scripts" / "oscompat.py").read_bytes() == (HTML_SLIDE / "scripts" / "oscompat.py").read_bytes(),
          "oscompat.py が html-slide とずれた（同じファイルを写す。サーバの src/quick_slide/oscompat.py も同じ）")


def t_auth_dead_stops():
    root = tmpdir()
    rc, so, se, calls = run(["deck", "--script", str(SAMPLE), "--out", str(root / "out"), "--parallel", "1"], root, mode="auth")
    check(rc == 4, f"失効で rc=4 にならない: {rc} {se}")
    check(len(_execs(calls)) == 1, f"失効のあとも描き続けた: {len(_execs(calls))} 回")
    check("codex login" in se, "入り直す案内が無い")
    check(not (root / "out" / "deck.pptx").exists(), "失敗したのに PPTX を書いた")


def t_auth_stops_running_workers():
    root = tmpdir()
    t0 = time.monotonic()
    rc, so, se, calls = run(["deck", "--script", str(SAMPLE), "--out", str(root / "out"), "--parallel", "5"], root,
                            mode="auth_page1", extra_env={"FAKE_CODEX_SLOW": "30"})
    took = time.monotonic() - t0
    starts = _starts(calls)
    check(rc == 4 and len(starts) == 5, f"rc={rc} 始まった数={len(starts)} {se}")
    check(took < 15, f"失効のあと、走っているワーカーを止めていない（{took:.1f} 秒かかった）")
    check(not any(pathlib.Path(c["home"]).exists() for c in starts), "止めたワーカーの隔離ホームが残った")


def t_not_logged_in():
    root = tmpdir()
    rc, so, se, calls = run(["deck", "--script", str(SAMPLE), "--out", str(root / "out")], root,
                            extra_env={"FAKE_CODEX_LOGGEDOUT": "1"})
    check(rc == 4 and not _execs(calls) and "codex login" in se, f"rc={rc} {se}")
    root2 = tmpdir()
    home = root2 / "empty-home"
    home.mkdir()
    rc, so, se, calls = run(["auth"], root2, home=home)
    check(rc == 4 and "auth.json" in se, f"auth.json が無いのに通った: {se}")


def t_quick_auth_looks_without_calling_codex():
    # ログインの記録があれば通る。codex は呼ばない（古いトークンでも更新しに行かない）
    root = tmpdir()
    rc, so, se, calls = run(["auth", "--quick"], root, home=fake_home(root, age_days=9, left_days=0.05))
    check(rc == 0 and so.startswith("OK:") and calls == [], f"rc={rc} {so} {se} {calls}")
    # 記録が無い・API キーのログイン・codex が見つからない、はどれも 4
    root2 = tmpdir()
    empty = root2 / "empty-home"
    empty.mkdir()
    rc, so, se, calls = run(["auth", "--quick"], root2, home=empty)
    check(rc == 4 and "auth.json" in se and calls == [], f"記録が無いのに通った: {se}")
    apikey = root2 / "apikey-home"
    apikey.mkdir()
    (apikey / "auth.json").write_text(json.dumps({"OPENAI_API_KEY": "sk-x", "tokens": None}), encoding="utf-8")
    rc, so, se, calls = run(["auth", "--quick"], root2, home=apikey)
    check(rc == 4 and "API キー" in se, f"API キーのログインで通った: {se}")
    rc, so, se, calls = run(["auth", "--quick"], root2, extra_env={"IMAGE_SLIDE_CODEX": str(root2 / "no-such-codex")})
    check(rc == 4 and "codex が見つからない" in se, f"codex が無いのに通った: {se}")


def t_stale_token_refreshes_once():
    root = tmpdir()
    home = fake_home(root, age_days=9, left_days=0.05)          # 期限まで1.2時間
    rc, so, se, calls = run(["deck", "--script", str(SAMPLE), "--out", str(root / "out"), "--no-pdf"], root, home=home)
    check(rc == 0, se)
    ref = [c for c in calls if c["cmd"] == "refresh"]
    check(len(ref) == 1 and ref[0]["home"] == str(home), f"古いトークンの更新が実ホームで1回でない: {ref}")
    b = json.loads((root / "out" / "build.json").read_text(encoding="utf-8"))
    check(b["auth"]["refreshed_here"] is True and b["auth"]["hours_left"] > 24, b["auth"])
    # 最後の更新から7日を過ぎても、期限まで十分なら更新させない（codex は期限の5分前まで更新しないので、させても変わらない）
    root2 = tmpdir()
    home2 = fake_home(root2, age_days=7.03, left_days=3)
    rc, so, se, calls = run(["deck", "--script", str(SAMPLE), "--out", str(root2 / "out"), "--no-pdf"], root2, home=home2,
                            extra_env={"FAKE_CODEX_NO_REFRESH": "1"})
    check(rc == 0 and not [c for c in calls if c["cmd"] == "refresh"], f"日数だけで更新しに行った／止まった: rc={rc} {se}")
    # 期限が近いのに更新されなければ rc=4（入り直しを案内する）
    root3 = tmpdir()
    home3 = fake_home(root3, age_days=1, left_days=0.05)
    rc, so, se, calls = run(["deck", "--script", str(SAMPLE), "--out", str(root3 / "out"), "--no-pdf"], root3, home=home3,
                            extra_env={"FAKE_CODEX_NO_REFRESH": "1"})
    check(rc == 4 and "codex login" in se and not _execs(calls), f"更新されないまま描きに行った: rc={rc} {se}")
    sys.path.insert(0, str(SKILL / "scripts"))
    import codex_run as X
    raw = json.dumps({"tokens": {"access_token": fake_jwt(time.time() + 3600), "refresh_token": "r"},
                      "last_refresh": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())}).encode()
    check(not X.freshness(raw)["fresh"], "期限まで1時間なのに新しい扱い")
    noexp = lambda days: json.dumps({"tokens": {"access_token": "opaque", "refresh_token": "r"},  # noqa: E731
                                     "last_refresh": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime(time.time() - days * 86400))}).encode()
    check(X.freshness(noexp(1))["fresh"] and not X.freshness(noexp(8))["fresh"], "期限が読めないトークンで日数の予備が効いていない")
    check(not X.freshness(json.dumps({"OPENAI_API_KEY": "k", "tokens": None}).encode())["chatgpt"], "API キーのログインを通した")


def t_retry_on_failure():
    root = tmpdir()
    rc, so, se, calls = run(["deck", "--script", str(SAMPLE), "--out", str(root / "out"), "--no-pdf"], root, mode="fail_first")
    ex = _execs(calls)
    check(rc == 0 and len(ex) == 10, f"失敗のあと1回描き直していない: rc={rc} exec={len(ex)} {se}")
    root = tmpdir()
    rc, so, se, calls = run(["deck", "--script", str(SAMPLE), "--out", str(root / "out"), "--retries", "0"], root, mode="none")
    check(rc == 1 and "描けなかった" in se and not (root / "out" / "deck.pptx").exists(), f"rc={rc} {se}")


def t_rate_limit_waits():
    root = tmpdir()
    rc, so, se, calls = run(["deck", "--script", str(SAMPLE), "--out", str(root / "out"), "--no-pdf"], root, mode="rate_first")
    ex = _execs(calls)
    check(rc == 0 and "枠切れ" in se and len(ex) == 10, f"枠切れのあと待って続けていない: rc={rc} {se}")
    first_limit = min(c["t1"] for c in ex if c["n"] == 1)
    again = [c["t0"] for c in ex if c["n"] == 2]
    # retry-after: 1 なので、描き直しはどれも最初の枠切れから1秒（多少の揺れを見て 0.9 秒）以上あとに始まる
    check(min(again) - first_limit >= 0.9, f"枠切れのあと待たずに描き直した（{min(again) - first_limit:.2f} 秒後）")


def t_band_retry_and_plate():
    root = tmpdir()
    rc, so, se, calls = run(["deck", "--script", str(SAMPLE), "--style", "standard", "--out", str(root / "out"), "--no-pdf"], root, mode="band_first")
    ex = _execs(calls)
    # 帯が要るのはフッタがある中身のページ（2〜5枚目）。表紙は帯が要らないので描き直さない
    again = sorted(c["page"] for c in ex if c["n"] == 2)
    check(rc == 0 and again == [2, 3, 4, 5], f"帯が空いていないページだけ描き直していない: {again} {se}")
    check(all(c["retry_notice"] for c in ex if c["n"] == 2), "描き直しの指示に理由が入っていない")
    check("plates=0" in so, so)
    root = tmpdir()
    rc, so, se, calls = run(["deck", "--script", str(SAMPLE), "--style", "standard", "--out", str(root / "out"), "--retries", "0"], root, mode="band_first")
    b = json.loads((root / "out" / "build.json").read_text(encoding="utf-8"))
    check(rc == 0 and [s["plate"] for s in b["slides"]] == [False, True, True, True, True], f"板: {[s['plate'] for s in b['slides']]}")
    check("帯が空いていない" in se, "板を敷いたのに警告が無い")


def t_worker_refresh_is_reported():
    root = tmpdir()
    rc, so, se, calls = run(["deck", "--script", str(SAMPLE), "--out", str(root / "out"), "--no-pdf"], root, mode="refresh")
    b = json.loads((root / "out" / "build.json").read_text(encoding="utf-8"))
    check(rc == 0 and b["auth"]["worker_refreshed"] == 5 and "トークンが更新された" in se, f"{b['auth']} {se}")


def t_pptx_structure():
    root = tmpdir()
    from PIL import Image
    from pptx import Presentation
    Image.new("RGBA", (400, 100), (200, 30, 30, 255)).save(root / "logo.png")
    out = root / "out"
    rc, so, se, _ = run(["deck", "--script", str(SAMPLE), "--out", str(out), "--logo", str(root / "logo.png")], root)
    check(rc == 0, se)
    prs = Presentation(str(out / "deck.pptx"))
    check(len(prs.slides) == 5 and prs.slide_width == 12192000 and prs.slide_height == 6858000, "16:9 の5枚でない")
    data = json.loads(SAMPLE.read_text(encoding="utf-8"))
    for i, (slide, spec) in enumerate(zip(prs.slides, data["slides"]), 1):
        names = [s.name for s in slide.shapes]
        pic = slide.shapes[0]
        check(pic.shape_type == 13 and pic.width == prs.slide_width and pic.left == 0, f"{i}枚目: 全面の絵が一番下にない")
        check(pic._element.nvPicPr.cNvPr.get("descr") == S.headline_of(spec), f"{i}枚目: 代替テキストが見出しでない")
        check("ロゴ" in names, f"{i}枚目: ロゴが無い {names}")
        chrome = spec["type"] != "cover"
        check(("フッタ" in names) == chrome and ("ページ番号" in names) == chrome, f"{i}枚目: 帯の部品 {names}")
        if chrome:
            num = next(s for s in slide.shapes if s.name == "ページ番号")
            check('type="slidenum"' in num._element.xml, f"{i}枚目: ページ番号がフィールドでない")
            foot = next(s for s in slide.shapes if s.name == "フッタ")
            check(foot.text_frame.text == data["footer"], f"{i}枚目: フッタの文字 {foot.text_frame.text!r}")
        notes = slide.notes_slide.notes_text_frame.text
        check(notes == "\n".join(spec["talk"]), f"{i}枚目: ノートが talk と違う")
    insp = (out / "inspect.md").read_text(encoding="utf-8")
    check("一字一句" in insp and "slide-05.png" in insp, "inspect.md に検品の手がかりが無い")
    check(f"フッタ「{data['footer']}」・ページ番号「2」・ロゴ" in insp, "inspect.md に帯に載せた部品が無い")


# --- Codex レビュー r1 の指摘の再発防止 ---------------------------------------------------

def t_content_keeps_nested_and_layout():
    c = S.content_of({"type": "big-number", "headline": "h", "value": "3", "side": [{"value": "20", "text": "前年"}],
                      "role": "answer"}, [])
    check(c.get("side") == [{"value": "20", "text": "前年"}] and "role" not in c, f"大きな数字の脇の数字が消えた／role が残った: {c}")
    c = S.content_of({"type": "team", "headline": "h", "members": [{"name": "A", "role": "責任者"}]}, [])
    check(c["members"][0].get("role") == "責任者", f"体制の肩書きが消えた: {c}")
    c = S.content_of({"type": "bullets", "headline": "h", "items": ["a", "b"], "side": "left", "split": "1/2"}, [])
    check("side" not in c and "split" not in c, f"箇条書きの写真の側が中身に入った: {c}")
    root = tmpdir()
    keys, texts = [], []
    for lay in ("left", "right"):
        p = write_script(root, {"slides": [{"type": "photo", "headline": "海", "layout": lay, "image": {"prompt": "海"}}]})
        deck = S.load(p)
        b = P.build(deck, deck["slides"][0], S.load_theme(None, root, {}))
        keys.append(P.page_key(b["key_text"], []))
        texts.append(b["text"])
    check(keys[0] != keys[1], "layout を変えても鍵が同じ（描き直されない）")
    check('"layout": "right"' in texts[1] and "配置の指定" in texts[1], "配置の指定が描く指示に入っていない")


def t_image_slots_match_html_slide():
    if not (HTML_SLIDE / "examples" / "image-deck.json").is_file():
        print("html-slide が無いので飛ばす")
        return
    sys.path.insert(0, str(HTML_SLIDE / "scripts"))
    import deck as HD
    script = HTML_SLIDE / "examples" / "image-deck.json"
    theirs = [(x["slide"], x["slot"], pathlib.Path(x["default_path"]).with_suffix("")) for x in HD.image_slots(script)]
    deck = S.load(script)
    ours = []
    for pg in deck["slides"]:
        for x in S.images_of(pg["spec"], deck["dir"], pg["no"]):
            ours.append((pg["no"], x["slot"]))
    check([(a, b) for a, b, _ in theirs] == ours, f"画像の枠の名前が html-slide とずれた: {set(ours) ^ {(a, b) for a, b, _ in theirs}}")


def t_failed_redo_keeps_old_pair():
    root = tmpdir()
    home = fake_home(root)
    out = root / "out"
    rc, _, se, _ = run(["deck", "--script", str(SAMPLE), "--style", "standard", "--out", str(out), "--no-pdf"], root, home=home)
    check(rc == 0, se)
    raw2, meta2 = (rawdir(root, out) / "slide-02.png").read_bytes(), (rawdir(root, out) / "slide-02.json").read_text(encoding="utf-8")
    shutil.rmtree(root / "calls")
    rc, _, se, _ = run(["deck", "--script", str(SAMPLE), "--style", "standard", "--out", str(out), "--no-pdf", "--redo", "2"], root, home=home,
                       mode="band_then_fail")
    check(rc == 1, f"描き直しの失敗が rc=1 にならない: {rc} {se}")
    check((rawdir(root, out) / "slide-02.png").read_bytes() == raw2, "失敗した描き直しの絵が raw に入った（次の実行で使い回される）")
    check((rawdir(root, out) / "slide-02.json").read_text(encoding="utf-8") == meta2, "鍵が書き換わった")
    check(not list(rawdir(root, out).glob("*.try.png")), "候補の絵が残った")


def t_private_home_cleanup():
    import tempfile as tf
    import codex_run as X
    real = tf.mkdtemp
    made = []

    def boom(prefix="", **kw):
        if prefix.startswith("image-slide-work."):
            raise OSError("ディスクがいっぱい（試験）")
        d = real(prefix=prefix, **kw)
        made.append(d)
        return d
    X.tempfile.mkdtemp = boom
    try:
        root = tmpdir()
        try:
            X.generate("codex", b'{"tokens":{"refresh_token":"r"}}', "x", [], root / "a.png", root / "logs", timeout=5)
            raise AssertionError("失敗しなかった")
        except X.CodexError as e:
            check(e.kind == "gen", e)
    finally:
        X.tempfile.mkdtemp = real
    homes = [d for d in made if "image-slide-home." in d]
    check(homes and not any(pathlib.Path(d).exists() for d in homes), f"準備の途中で落ちたのに隔離ホームが残った: {homes}")


def t_auth_malformed_is_rc4():
    import codex_run as X
    check(X.freshness(b"[]")["chatgpt"] is False, "配列の auth.json で例外")
    f = X.freshness(json.dumps({"tokens": {"refresh_token": "r", "access_token": fake_jwt(time.time() + 5 * 86400)},
                                "last_refresh": time.strftime("%Y-%m-%dT%H:%M:%S", time.gmtime())}).encode())
    check(f["fresh"], f"タイムゾーンの無い時刻で新しいと読めない: {f}")
    real = X.subprocess.run

    def slow(*a, **k):
        raise X.subprocess.TimeoutExpired(a[0], k.get("timeout"))
    X.subprocess.run = slow
    root = tmpdir()
    os.environ["CODEX_HOME"] = str(fake_home(root))
    try:
        X.check_auth("codex", log=lambda *_: None)
        raise AssertionError("時間切れが通った")
    except X.CodexError as e:
        check(e.kind == "auth" and "秒で終わらなかった" in e.detail, e.detail)
    finally:
        X.subprocess.run = real


def t_secrets_redacted():
    root = tmpdir()
    home = fake_home(root)
    tokens = json.loads((home / "auth.json").read_text(encoding="utf-8"))["tokens"]
    tokens["refresh_token"] = "rt_SECRETSECRETSECRET123"
    a = json.loads((home / "auth.json").read_text(encoding="utf-8"))
    a["tokens"] = tokens
    (home / "auth.json").write_text(json.dumps(a), encoding="utf-8")
    out = root / "out"
    rc, so, se, _ = run(["deck", "--script", str(SAMPLE), "--out", str(out), "--retries", "0", "--parallel", "2"], root,
                        mode="leak", home=home)
    check(rc == 1, f"rc={rc}")
    blobs = [se, (out / "build.json").read_text(encoding="utf-8")] + [x.read_text(encoding="utf-8") for x in (out / "logs").glob("*")]
    for b in blobs:
        check("SECRETSECRET" not in b and "eyJhbGciOiJub25lIn0" not in b, "トークンがログ・build.json・画面に出た")
    check(any("［伏せた］" in b for b in blobs), "伏せた印が無い（伏せる処理を通っていない）")


def t_warning_words_not_misread():
    root = tmpdir()
    rc, so, se, calls = run(["deck", "--script", str(SAMPLE), "--out", str(root / "out"), "--no-pdf"], root, mode="warn_words")
    check(rc == 0 and len(_execs(calls)) == 5, f"警告の中の語を失効・枠切れと読み違えた: rc={rc} {se}")


def t_pptx_text_size():
    root = tmpdir()
    out = root / "out"
    rc, _, se, _ = run(["deck", "--script", str(SAMPLE), "--out", str(out), "--no-pdf"], root)
    check(rc == 0, se)
    from pptx import Presentation
    s = Presentation(str(out / "deck.pptx")).slides[1]
    foot = next(x for x in s.shapes if x.name == "フッタ")
    check('sz="1100"' in foot._element.xml, "帯の文字が 11pt（確認用の絵の 22px と同じ）でない")


# --- Codex レビュー r2 の指摘の再発防止 ---------------------------------------------------

def t_stopped_worker_logs_redacted():
    root = tmpdir()
    home = fake_home(root)
    a = json.loads((home / "auth.json").read_text(encoding="utf-8"))
    a["tokens"]["refresh_token"] = "rt_STOPPEDSECRET1234567"
    (home / "auth.json").write_text(json.dumps(a), encoding="utf-8")
    out = root / "out"
    rc, so, se, calls = run(["deck", "--script", str(SAMPLE), "--out", str(out), "--parallel", "5"], root, home=home,
                            mode="auth_page1", extra_env={"FAKE_CODEX_SLOW": "30", "FAKE_CODEX_LEAK_EARLY": "1"})
    check(rc == 4, f"rc={rc}")
    logs = list((out / "logs").glob("*"))
    names = sorted(x.name for x in logs if x.name.endswith(".stderr.txt"))
    check(names == [f"slide-{i:02d}.stderr.txt" for i in range(1, 6)], f"止めたワーカーのログがページの名前で写されていない: {names}")
    check(not any("STOPPEDSECRET" in x.read_text(encoding="utf-8") for x in logs), "止めたワーカーのログにトークンが残った")


def t_auth_with_image_stops():
    root = tmpdir()
    rc, so, se, calls = run(["deck", "--script", str(SAMPLE), "--out", str(root / "out"), "--parallel", "1"], root,
                            mode="auth_with_image")
    check(rc == 4 and len(_execs(calls)) == 1, f"絵が返ったときの失効を見逃した: rc={rc} exec={len(_execs(calls))}")


def t_broken_meta_redraws():
    root = tmpdir()
    home = fake_home(root)
    out = root / "out"
    rc, _, se, _ = run(["deck", "--script", str(SAMPLE), "--out", str(out), "--no-pdf"], root, home=home)
    check(rc == 0, se)
    (rawdir(root, out) / "slide-02.json").write_text('{"key": ', encoding="utf-8")
    shutil.rmtree(root / "calls")
    rc, so, se, calls = run(["deck", "--script", str(SAMPLE), "--out", str(out), "--no-pdf"], root, home=home)
    check(rc == 0 and [c["page"] for c in _execs(calls)] == [2], f"壊れた鍵のページだけを描き直していない: rc={rc} {se}")


def t_failure_lines_reading():
    import codex_run as X
    cases = [("2026-10-03T00:00:00Z ERROR codex: token_revoked", True, False),
             ("Error: reference image 429.png unavailable", False, False),
             ("WARNING: caption mentions token_revoked", False, False),
             ("ERROR: 429 Too Many Requests", False, True),
             ("ERROR: You've hit your usage limit. Try again in 30s", False, True)]
    for line, auth, rate in cases:
        f = X.failure_lines("", line)
        check(bool(X.AUTH_DEAD.search(f)) == auth and bool(X.RATE.search(f)) == rate, f"読み違え: {line!r}")
    f = X.failure_lines('{"type":"turn.failed","error":{"message":"token_revoked"}}\n{"type":"item.completed","text":"token_revoked"}', "")
    check(f.count("token_revoked") == 1, "--json の失敗のイベントだけを拾っていない")


def t_unreadable_auth_is_rc4():
    if os.name == "nt":
        print("skip: Windows は chmod(0) で読めないファイルを作れない（Quick Slide の ADR 0012）")
        return
    root = tmpdir()
    home = fake_home(root)
    (home / "auth.json").chmod(0)
    try:
        rc, so, se, calls = run(["auth"], root, home=home)
    finally:
        (home / "auth.json").chmod(0o600)
    check(rc == 4 and "読めない" in se, f"読めない auth.json で rc=4 にならない: rc={rc} {se}")


# --- 構図は任せる（ADR 0005） -------------------------------------------------------------

def t_composition_left_to_codex():
    root = tmpdir()
    p = write_script(root, {"slides": [{"type": "statement", "headline": "主張"}]})
    deck = S.load(p)
    th = S.load_theme(None, root, {})
    pitch = P.build(deck, deck["slides"][0], th)
    import styles as ST
    biz = P.build(deck, deck["slides"][0], th, style=ST.load("standard"))
    check("構図: 指定しない" in pitch["text"] and "center-focus" not in pitch["text"], "構図を任せていない")
    check("主役を1つに絞り" in pitch["text"] and "主役を1つに絞り" not in biz["text"], "作風の既定がピッチ寄りでない")
    check(P.page_key(pitch["key_text"], []) != P.page_key(biz["key_text"], []), "作風を変えても鍵が同じ")
    rc, so, se, _ = run(["layouts"], root)
    check(rc == 2, "layouts の入口が残っている")
    check(not (SKILL / "layouts").exists() and not (SKILL / "layout-map.yaml").exists(), "製品のレイアウトの棚が残っている")


# --- 持ち込みのレイアウト（ADR 0006） -------------------------------------------------------

def _layout_file(root, name="weekly", ref=True, parts=None):
    from PIL import Image
    if ref:
        Image.new("RGB", (160, 90), "orange").save(root / f"{name}-sample.png")
    import yaml
    d = {"name": name, "composition": "上にヘッダ、右2/3に主役の図解、下に学びの帯"}
    if ref:
        d["reference"] = f"{name}-sample.png"
    d["parts"] = parts if parts is not None else {"week_badge": "ヘッダ右上のバッジ", "takeaway": "下部の主役色の帯"}
    (root / f"{name}.yaml").write_text(yaml.safe_dump(d, allow_unicode=True), encoding="utf-8")
    return root / f"{name}.yaml"


def t_custom_layout_prompt():
    import layouts as Y
    root = tmpdir()
    lf = _layout_file(root)
    lays = Y.load([lf])
    spec = {"type": "flow", "headline": "AIが右腕、芯はメモ", "items": [{"label": "1", "title": "家を整える"},
            {"label": "2", "title": "驚き"}, {"label": "3", "title": "芯"}],
            "_layout": {"name": "weekly", "parts": {"week_badge": "2026-W40", "takeaway": "芯を支えるのはメモ"}}}
    p = write_script(root, {"slides": [spec, {"type": "statement", "headline": "任せるページ"}]})
    deck = S.load(p)
    th = S.load_theme(None, root, {})
    b = P.build(deck, deck["slides"][0], th, layouts=lays, refs_common=1)
    t = b["text"]
    for want in ("ユーザーが指定した「weekly」", "上にヘッダ", "2026-W40", "芯を支えるのはメモ", "ヘッダ右上のバッジ",
                 "添付の参照画像1は、このレイアウトの構図の見本", "添付の参照画像2は、資料全体"):
        check(want in t, f"描く指示に {want!r} が無い")
    check("構図: 指定しない" not in t, "持ち込みのレイアウトのページで、構図を任せる言葉が残った")
    free = P.build(deck, deck["slides"][1], th, layouts=lays)
    check("構図: 指定しない" in free["text"] and "weekly" not in free["text"], "_layout の無いページにレイアウトが効いた")
    # 鍵: 参照画像の枚数などほかの条件をそろえて、構図・部品の置き場・部品の文字を1つずつ変える（Codex レビュー）
    import copy
    import yaml as _y

    def key_with(change):
        r2 = tmpdir()
        d = _y.safe_load(lf.read_text(encoding="utf-8"))
        d.pop("reference", None)
        change(d)
        text = d.pop("_text", None)
        (r2 / "l.yaml").write_text(_y.safe_dump(d, allow_unicode=True), encoding="utf-8")
        lay = Y.load([r2 / "l.yaml"])
        sp = copy.deepcopy(spec)
        if text is not None:
            sp["_layout"]["parts"]["week_badge"] = text
        pp = write_script(r2, {"slides": [sp]})
        dk = S.load(pp)
        return P.page_key(P.build(dk, dk["slides"][0], th, layouts=lay, refs_common=1)["key_text"], [])
    base = key_with(lambda d: None)
    check(key_with(lambda d: None) == base, "同じ中身で鍵が変わった")
    check(key_with(lambda d: d.update(composition="別の構図")) != base, "構図を変えても鍵が同じ")
    check(key_with(lambda d: d["parts"].update(week_badge="ヘッダ左のバッジ")) != base, "部品の置き場を変えても鍵が同じ")
    check(key_with(lambda d: d.update(_text="2026-W41")) != base, "部品の文字を変えても鍵が同じ")


def t_custom_layout_cli_and_errors():
    root = tmpdir()
    lf = _layout_file(root)
    ok = {"slides": [{"type": "statement", "headline": "h", "_layout": {"name": "weekly", "parts": {"week_badge": "W40"}}}]}
    p = write_script(root, ok)
    rc, so, se, _ = run(["check", "--script", str(p), "--layout", str(lf)], root)
    check(rc == 0 and "レイアウト weekly" in so, f"rc={rc} {so} {se}")
    cases = [
        ([], ok, "渡されていない"),
        ([str(lf)], {"slides": [{"type": "statement", "headline": "h", "_layout": {"name": "weekly", "parts": {"nazo": "x"}}}]}, "無い部品"),
        ([str(lf)], {"slides": [{"type": "statement", "headline": "h", "_layout": 3}]}, "名前か"),
        ([str(root / "none.yaml")], ok, "ファイルが無い"),
    ]
    for lays, data, want in cases:
        p = write_script(root, data)
        rc, so, se, calls = run(["check", "--script", str(p)] + sum((["--layout", x] for x in lays), []), root)
        check(rc == 2 and want in se and not calls, f"{want}: rc={rc} {se}")
    # 部品の文字・ファイルの値の型（Codex レビュー: null・真偽値・空・配列を黙って通さない）
    for v, want in ((None, "None"), (True, "True"), ("", "空"), ([], "空の配列"), ({}, "空の辞書")):
        p = write_script(root, {"slides": [{"type": "statement", "headline": "h", "_layout": {"name": "weekly", "parts": {"week_badge": v}}}]})
        rc, so, se, _ = run(["check", "--script", str(p), "--layout", str(lf)], root)
        check(rc == 2 and want in se, f"部品の文字 {v!r}: rc={rc} {se}")
    for body, want in (("name: b\ncomposition: c\nparts: false\n", "parts は"), ("name: b\ncomposition: c\nreference: false\n", "reference は"),
                       ("name: b\ncomposition: c\n1: a\nz: b\n", "知らないキー")):
        (root / "badv.yaml").write_text(body, encoding="utf-8")
        rc, so, se, _ = run(["check", "--script", str(write_script(root, ok)), "--layout", str(root / "badv.yaml")], root)
        check(rc == 2 and want in se, f"{body!r}: rc={rc} {se}")
    bad = root / "bad.yaml"
    bad.write_text("name: b\nlayout: x\n", encoding="utf-8")
    rc, so, se, _ = run(["check", "--script", str(write_script(root, ok)), "--layout", str(bad)], root)
    check(rc == 2 and "知らないキー" in se, se)
    # html-slide は _layout を読み飛ばす（台本は共通のまま）
    if (HTML_SLIDE / "scripts" / "html_slide.py").is_file():
        r = subprocess.run([sys.executable, str(HTML_SLIDE / "scripts" / "html_slide.py"), "check", "--script", str(write_script(root, ok))],
                           capture_output=True, text=True, encoding="utf-8")
        check(r.returncode == 0, f"_layout つきの台本が html-slide で通らない: {r.stderr}")


def t_custom_layout_deck():
    root = tmpdir()
    lf = _layout_file(root)
    data = json.loads(SAMPLE.read_text(encoding="utf-8"))
    data["slides"][1]["_layout"] = {"name": "weekly", "parts": {"week_badge": "2026-W40"}}
    p = write_script(root, data)
    out = root / "out"
    rc, so, se, calls = run(["deck", "--script", str(p), "--out", str(out), "--layout", str(lf), "--no-pdf"], root)
    check(rc == 0, se)
    ex = _execs(calls)
    check([len(c["refs"]) for c in sorted(ex, key=lambda c: c["page"])] == [0, 1, 0, 0, 0], "構図の見本が2枚目だけに添えられていない")
    b = json.loads((out / "build.json").read_text(encoding="utf-8"))
    check([s["layout"] for s in b["slides"]] == [None, "weekly", None, None, None], b["slides"])
    insp = (out / "inspect.md").read_text(encoding="utf-8")
    check("（レイアウト weekly）" in insp and "2026-W40" in insp, "inspect.md にレイアウトと部品の文字が無い")


def t_revise():
    root = tmpdir()
    from PIL import Image
    Image.new("RGB", (40, 40), "blue").save(root / "hint.png")
    data = json.loads(SAMPLE.read_text(encoding="utf-8"))
    p = write_script(root, data)
    out = root / "out"
    rv = lambda *x, mode="ok": run(["revise", "--script", str(p), "--out", str(out), "--no-pdf", *x], root, mode=mode)  # noqa: E731
    # 描く前の微修正は土台が無いので止める（作り直しは、ほかのページが無いので止める）
    rc, _, err, calls = rv("--page", "2", "--note", "x")
    check(rc == 2 and "deck を流す" in err and not _execs(calls), f"描く前の revise が止まらない: {err}")
    rc, _, se, _ = run(["deck", "--script", str(p), "--out", str(out), "--no-pdf"], root)
    check(rc == 0, se)
    shutil.rmtree(root / "calls", ignore_errors=True)
    before = (rawdir(root, out) / "slide-02.png").read_bytes()
    # 微修正（既定）: 2枚目だけ描く。土台（前の絵の写し）を最後に添え、直しの節を足す
    rc, so, se, calls = rv("--page", "2", "--note", "見出しを大きく")
    check(rc == 0, se)
    ex = _execs(calls)
    check([c["page"] for c in ex] == [2], f"2枚目だけを描いていない: {[c['page'] for c in ex]}")
    hist = sorted((rawdir(root, out) / "history").glob("slide-02-*.png"))
    check(len(hist) == 1 and hist[0].read_bytes() == before and hist[0].with_suffix(".json").is_file(), f"前の絵を残していない: {hist}")
    check(ex[0]["refs"] == [str(hist[0])], f"土台が添えられていない: {ex[0]['refs']}")
    text = (out / "prompts" / "slide-02.txt").read_text(encoding="utf-8")
    for want in ("改稿（微修正）: 渡された現在のページ画像を土台にする", "添付の参照画像1は、いまのこのページの絵（土台）", "利用者の指示:\n見出しを大きく",
                 "新設しない"):
        check(want in text, f"微修正の指示に {want!r} が無い")
    b = json.loads((out / "build.json").read_text(encoding="utf-8"))
    row = b["slides"][1]
    check(row["status"] == "generated" and row["revision"]["mode"] == "minor" and row["revision"]["note"] == "見出しを大きく", row)
    check(all(r["status"] == "cached" for r in b["slides"] if r["no"] != 2), "ほかのページを描いた")
    # 台本を変えずに deck を流し直しても、直した絵を使い回す（鍵は台本から作った元のまま）
    shutil.rmtree(root / "calls", ignore_errors=True)
    rc, so, se, calls = run(["deck", "--script", str(p), "--out", str(out), "--no-pdf"], root)
    check(rc == 0 and not _execs(calls), f"直した絵を deck が描き直した: {[c['page'] for c in _execs(calls)]}")
    check(json.loads((out / "build.json").read_text(encoding="utf-8"))["slides"][1].get("revision"), "使い回しで直しの記録が消えた")
    # 作り直し: 今の絵は渡さない。添えた参考画像は渡す。文が無くても参考画像だけで直せる
    shutil.rmtree(root / "calls", ignore_errors=True)
    rc, so, se, calls = rv("--page", "2", "--remake", "--attach", str(root / "hint.png"))
    check(rc == 0, se)
    ex = _execs(calls)
    check(ex[0]["refs"] == [str((root / "hint.png").resolve())], f"作り直しで渡したもの: {ex[0]['refs']}")
    text = (out / "prompts" / "slide-02.txt").read_text(encoding="utf-8")
    check("改稿（作り直し）" in text and "参照画像1は、利用者が添えた参考画像" in text and "文章の指示は無い" in text and "土台）" not in text, text[-400:])
    # 台本が変わったほかのページは、revise では描かない（警告だけ）
    data["slides"][2]["headline"] = "変えた見出し"
    p = write_script(root, data)
    shutil.rmtree(root / "calls", ignore_errors=True)
    rc, so, se, calls = rv("--page", "2", "--note", "y")
    check(rc == 0 and [c["page"] for c in _execs(calls)] == [2] and "3枚目は台本が変わっている" in se, f"{[c['page'] for c in _execs(calls)]} {se}")
    # 誤り: 範囲の外・中身なし・枚数の上限
    for args, word in ((["--page", "9", "--note", "x"], "範囲の外"), (["--page", "2"], "--note"),
                       (["--page", "2", "--note", "x"] + sum((["--attach", str(root / "hint.png")] for _ in range(5)), []), "5 枚まで")):
        rc, _, err, _ = rv(*args)
        check(rc == 2 and word in err, f"{args[:2]} で止まらない／言い方が違う: {err}")


# --- スタイル（ADR 0007） --------------------------------------------------------------------

def t_styles_builtin_and_custom():
    import styles as ST
    check(ST.load(None)["name"] == "visual" and ST.load("standard")["name"] == "standard", "組み込みのスタイルの名前・既定")
    for old, new in (("pitch", "visual"), ("business", "standard")):
        try:
            ST.load(old)
            raise AssertionError(f"前の名前 {old} が通った")
        except S.ScriptError as e:
            check(new in str(e), f"{old} → {new} の案内が無い: {e}")
    root = tmpdir()
    (root / "weekly-style.yaml").write_text("name: weekly\ndescription: 週報\ntext: フラットイラスト。写真は使わない\n", encoding="utf-8")
    st = ST.load(str(root / "weekly-style.yaml"))
    p = write_script(root, {"slides": [{"type": "statement", "headline": "h"}]})
    deck = S.load(p)
    t = P.build(deck, deck["slides"][0], S.load_theme(None, root, {}), style=st)["text"]
    check("スタイル「weekly」（週報）" in t and "フラットイラスト" in t and "主役を1つに絞り" not in t, "カスタムスタイルが入っていない")
    bad = [("name: x\n", "text が無い"), ("name: x\ntext: y\nlayout: z\n", "知らないキー"), ("- a\n", "name・description・text"),
           ("name: x\ntext: y\n1: a\nz: b\n", "知らないキー")]
    for body, want in bad:
        (root / "bad.yaml").write_text(body, encoding="utf-8")
        rc, so, se, _ = run(["check", "--script", str(p), "--style", str(root / "bad.yaml")], root)
        check(rc == 2 and want in se, f"{body!r}: rc={rc} {se}")
    rc, so, se, _ = run(["check", "--script", str(p), "--style", "nazo"], root)
    check(rc == 2 and "standard・visual" in se, se)
    rc, so, se, _ = run(["check", "--script", str(p), "--style", str(root / "weekly-style.yaml")], root)
    check(rc == 0 and "スタイル weekly" in so, so)
    lf = root / "lay.yaml"
    lf.write_text("name: l\ncomposition: c\nstyle: s\n", encoding="utf-8")
    rc, so, se, _ = run(["check", "--script", str(p), "--layout", str(lf)], root)
    check(rc == 2 and "--style で渡す" in se, f"レイアウトのファイルの style を止めていない: {se}")


def t_templates_load():
    import layouts as Y
    import styles as ST
    st = ST.load(str(SKILL / "templates" / "custom-style.example.yaml"))
    lays = Y.load([SKILL / "templates" / "custom-layout.example.yaml"])
    check(st["name"] == "weekly" and "weekly-report" in lays and lays["weekly-report"]["parts"], "雛形が読めない")



def t_theme_primary_accent():
    """テーマの主役の色と差し色: 2色なら描く指示に両方と使い分けが入る・1色なら差し色の言葉は出ない・古い書き方は rc=2。"""
    import prompt as P
    root = tmpdir()
    one = S.load_theme(None, root, {})
    t1 = P.colors_text(one)
    check("主役の色 #006f83" in t1 and "差し色" not in t1 and "強調色は主役の色" in t1, t1)
    two = root / "two.yaml"
    two.write_text('background: "#ffffff"\nink: "#1a1a1a"\nprimary: "#f5a623"\naccent: "#399dd1"\n', encoding="utf-8")
    t2 = P.colors_text(S.load_theme(str(two), root, {}))
    for want in ("主役の色 #f5a623", "差し色 #399dd1", "一番広く目に入る色", "必ず主役の色にする", "自分で選んだ語も", "強調に主役の色は使わない"):
        check(want in t2, f"2色の描く指示に {want!r} が無い")
    old = root / "old.yaml"
    old.write_text('background: "#ffffff"\nink: "#1a1a1a"\naccent: "#f5a623"\n', encoding="utf-8")
    try:
        S.load_theme(str(old), root, {})
        check(False, "古い書き方（accent だけ）が通った")
    except S.ScriptError as e:
        check("primary に名前が変わった" in str(e), str(e))

# --- infographic（台本の全ページを1枚に凝縮する） ----------------------------------------

def t_infographic_one_page():
    root = tmpdir()
    out = root / "out"
    rc, so, se, calls = run(["deck", "--script", str(SAMPLE), "--out", str(out), "--style", "infographic"], root)
    check(rc == 0, se)
    check(len(_execs(calls)) == 1, f"5枚の台本で exec が {len(_execs(calls))} 回（1回のはず）")
    b = json.loads((out / "build.json").read_text(encoding="utf-8"))
    check(len(b["slides"]) == 1 and b["style"] == "infographic", b["slides"])
    text = (out / "prompts" / "slide-01.txt").read_text(encoding="utf-8")
    data = json.loads(SAMPLE.read_text(encoding="utf-8"))
    for want in ("1枚に凝縮する", '"区画": 5', "朝会は3つのことだけを話す", "まずは11月の4週間、試させてほしい", "（このページの中身）",
                 "並べ方の指定ではない", "スライドを並べたような見た目", "最初から設計する"):
        check(want in text, f"1枚の描く指示に {want!r} が無い")
    from pptx import Presentation
    prs = Presentation(str(out / "deck.pptx"))
    check(len(prs.slides) == 1, "PPTX が1枚でない")
    names = [x.name for x in prs.slides[0].shapes]
    check("フッタ" in names and "ページ番号" not in names, f"1枚のポスターの帯の部品: {names}")
    notes = prs.slides[0].notes_slide.notes_text_frame.text
    check("【1枚目】" in notes and "【5枚目】" in notes and data["slides"][4]["talk"][0] in notes, "ノートに全ページの話す台本が入っていない")
    insp = (out / "inspect.md").read_text(encoding="utf-8")
    check('"区画": 3' in insp, "inspect.md に区画が出ていない")


def t_infographic_layout_and_errors():
    import styles as ST
    check(ST.load("infographic")["one_page"] and not ST.load("visual")["one_page"], "one_page の印")
    root = tmpdir()
    lf = _layout_file(root, ref=False)
    data = json.loads(SAMPLE.read_text(encoding="utf-8"))
    data["slides"][0]["_layout"] = {"name": "weekly", "parts": {"week_badge": "W40"}}
    p = write_script(root, data)
    rc, so, se, _ = run(["check", "--script", str(p), "--style", "infographic", "--layout", str(lf), "--prompts", str(root / "pr")], root)
    check(rc == 0 and "W40" in (root / "pr" / "slide-01.txt").read_text(encoding="utf-8"), f"1枚目の _layout が効いていない: {se}")
    data["slides"][2]["_layout"] = "weekly2"
    rc, so, se, _ = run(["check", "--script", str(write_script(root, data)), "--style", "infographic", "--layout", str(lf)], root)
    check(rc == 2, f"2枚目以降の別のレイアウトを止めていない: {se}")
    (root / "s.yaml").write_text("name: x\ntext: y\none_page: 'yes'\n", encoding="utf-8")
    rc, so, se, _ = run(["check", "--script", str(SAMPLE), "--style", str(root / "s.yaml")], root)
    check(rc == 2 and "one_page は true か false" in se, se)
    (root / "s2.yaml").write_text("name: poster\ntext: y\none_page: true\n", encoding="utf-8")
    rc, so, se, _ = run(["check", "--script", str(SAMPLE), "--style", str(root / "s2.yaml")], root)
    check(rc == 0 and so.count("枚目") == 1, f"カスタムスタイルの one_page: true で1枚にならない: {so}")


def t_image_look_same_as_html_slide():
    """--image-look は html-slide と同じ3語・同じ言い回し。指定したときだけスタイルの節に1行足す。Codex の呼び方の写しも同じ。"""
    deck = S.load(SAMPLE)
    theme = S.load_theme(None, deck["dir"], {})
    pg = deck["slides"][2]
    plain_ = P.build(deck, pg, theme)["text"]
    check("画風:" not in plain_, "指定していないのに画風の行が入った")
    st = ST.load("standard")
    st["image_look"] = "handdrawn"
    t = P.build(deck, pg, theme, style=st)["text"]
    check(ST.IMAGE_LOOKS["handdrawn"] in t and "絵の部分" in t, "画風の行が無い")
    check(P.page_key(t, []) != P.page_key(plain_, []), "画風を変えても鍵が変わらない")
    root = tmpdir()
    (root / "s.yaml").write_text("name: x\ntext: y\nimage_look: oil\n", encoding="utf-8")
    rc, so, se, _ = run(["check", "--script", str(SAMPLE), "--style", str(root / "s.yaml")], root)
    check(rc == 2 and "image_look は" in se, se)
    rc, so, se, _ = run(["check", "--script", str(SAMPLE), "--image-look", "photo", "--prompts", str(root / "p")], root)
    check(rc == 0 and ST.IMAGE_LOOKS["photo"] in (root / "p" / "slide-01.txt").read_text(encoding="utf-8"), se)
    if not (HTML_SLIDE / "scripts" / "imagegen.py").is_file():
        print("html-slide が無いので、見比べは飛ばす")
        return
    sys.path.insert(0, str(HTML_SLIDE / "scripts"))
    import imagegen as HG
    check(ST.IMAGE_LOOKS == HG.IMAGE_LOOKS, "画風の言い回しが html-slide とずれた")
    mine = (SKILL / "scripts" / "codex_run.py").read_text(encoding="utf-8")
    theirs = (HTML_SLIDE / "scripts" / "codex_run.py").read_text(encoding="utf-8")
    norm = mine.replace("image-slide", "html-slide").replace("IMAGE_SLIDE", "HTML_SLIDE")
    check(norm == theirs, "html-slide の codex_run.py（写し）とずれた。直したら写し直す（名前の image-slide → html-slide だけが違う）")


def t_formats_and_thumbs():
    """出力（ADR 0013）: 既定で PPTX・PDF・ページの画像・サムネイル。--format で選ぶ（ページの画像はどれでも残る）・
    誤りは rc=2・ページを減らしたら古いページの画像を消す。outputs.py は html-slide と同じもの。"""
    root = tmpdir()
    home = fake_home(root)
    out = root / "out"
    p = write_script(root, json.loads(SAMPLE.read_text(encoding="utf-8")))
    (out / "slides").mkdir(parents=True)
    (out / "slides" / "slide-09.png").write_bytes(b"old")
    rc, so, se, _ = run(["deck", "--script", str(p), "--out", str(out)], root, home=home)
    check(rc == 0, se)
    for f in ("deck.pptx", "deck.pdf", "thumbs.png"):
        check((out / f).is_file(), f"{f} が無い")
    check(not (out / "slides" / "slide-09.png").exists(), "前の回の9枚目が残った")
    check(sorted(q.name for q in (out / "slides").glob("*.png")) == [f"slide-{i:02d}.png" for i in range(1, 6)], "ページの画像が5枚でない")
    b = json.loads((out / "build.json").read_text(encoding="utf-8"))
    check(b["outputs"]["format"] == ["pptx", "pdf", "png", "thumbs"] and len(b["outputs"]["png"]) == 5, b["outputs"])
    check("thumbs: " in so and "pdf: " in so, so)
    from PIL import Image
    check(Image.open(out / "thumbs.png").size[0] == 2072, "サムネイルの幅が 4列ぶんでない")
    for f in ("deck.pptx", "deck.pdf", "thumbs.png"):
        (out / f).unlink()
    rc, so, se, _ = run(["deck", "--script", str(p), "--out", str(out), "--format", "png"], root, home=home)
    check(rc == 0, se)
    check(not any((out / f).exists() for f in ("deck.pptx", "deck.pdf", "thumbs.png")), "png だけなのにほかの出力を作った")
    check(len(list((out / "slides").glob("slide-*.png"))) == 5, "ページの画像が無い")
    check("PPTX: 作っていない" in (out / "inspect.md").read_text(encoding="utf-8"), "inspect.md が PPTX のあるふりをした")
    rc, so, se, _ = run(["deck", "--script", str(p), "--out", str(out), "--format", "png,jpg"], root, home=home)
    check(rc == 2 and "--format は" in se, se)
    if (HTML_SLIDE / "scripts" / "outputs.py").is_file():
        check((SKILL / "scripts" / "outputs.py").read_bytes() == (HTML_SLIDE / "scripts" / "outputs.py").read_bytes(),
              "html-slide の outputs.py とずれた（同じファイル。直したら写す）")


TESTS = {k[2:]: v for k, v in globals().items() if k.startswith("t_")}

if __name__ == "__main__":
    name = sys.argv[1] if len(sys.argv) > 1 else ""
    if name not in TESTS:
        print("使えるテスト: " + ", ".join(TESTS))
        sys.exit(2)
    TESTS[name]()
    print("ok")
