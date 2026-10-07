"""絵を並べて PPTX と PDF にする。

PPTX の1枚 = 全面の絵（生成した絵を 16:9 にそろえたもの・帯の上の部品は焼かない）
           ＋ 帯の板（帯が空いていなかったときだけ・地の色）
           ＋ フッタの文字・ページ番号・ロゴの地・ロゴ（PowerPoint の部品。あとで直せる。ロゴの置き場はページごと・ADR 0008）
           ＋ ノート（台本の talk）
絵には「画面の文字」が無いので、見出しを絵の代替テキストに入れる（検索と読み上げのため）。
ページ番号はスライド番号のフィールドにする（並べ替えても番号が追従する）。
"""
from __future__ import annotations

import pathlib

import layout as L

EMU_PER_PX = 12192000 // L.W     # 13.333 インチ（12192000 EMU）を 1920px に（1px = 6350 EMU）
PT_PER_PX = EMU_PER_PX / 12700   # 1pt = 12700 EMU なので 1px = 0.5pt。帯の文字 22px は 11pt（確認用の絵と同じ大きさ）


def _emu(px) -> int:
    return int(round(px * EMU_PER_PX))


def _color(fore_color, hexstr):
    from pptx.dml.color import RGBColor
    fore_color.rgb = RGBColor.from_string(hexstr.lstrip("#"))


def _textbox(slide, box, text, *, font, color, align_right=False, align_center=False, field_slidenum=False):
    """1行の文字の枠。field_slidenum なら、文字の代わりにスライド番号のフィールド（<a:fld type="slidenum">）を置く。"""
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    x, y, w, h = box
    pad = 8
    tb = slide.shapes.add_textbox(_emu(x - pad), _emu(y - pad), _emu(w + 2 * pad), _emu(h + 2 * pad))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = _emu(pad)
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER if align_center else PP_ALIGN.RIGHT if align_right else PP_ALIGN.LEFT
    if field_slidenum:
        run = p._p.makeelement(qn("a:fld"), {"id": "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", "type": "slidenum"})
        p._p.append(run)
    else:
        run = p.add_run()._r
    rpr = run.makeelement(qn("a:rPr"), {"lang": "ja-JP", "sz": str(round(L.TEXT_PX * PT_PER_PX * 100))})
    fill = rpr.makeelement(qn("a:solidFill"), {})
    fill.append(fill.makeelement(qn("a:srgbClr"), {"val": color.lstrip("#").upper()}))
    rpr.append(fill)
    for tag in ("a:latin", "a:ea"):
        rpr.append(rpr.makeelement(qn(tag), {"typeface": font}))
    old = run.find(qn("a:rPr"))
    if old is not None:
        run.remove(old)
    run.insert(0, rpr)
    t = run.find(qn("a:t"))
    if t is None:
        t = run.makeelement(qn("a:t"), {})
        run.append(t)
    t.text = text
    return tb


def write_pptx(pages: list[dict], out: pathlib.Path, *, theme: dict, footer: str, logo: dict | None,
               font_file: str | None, title: str) -> None:
    """pages: [{no, base（16:9 にそろえた絵。透かしはここに焼いてある）, band, plate, chrome, notes, alt, number, logo（place_logo の返り値）,
    mid（下の行の中心・layout.row_mid）}]。logo は compose.load_logo の返り値。"""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Emu
    import compose as C

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(_emu(L.W)), Emu(_emu(L.H))
    prs.core_properties.title = title
    blank = prs.slide_layouts[6]
    for pg in pages:
        s = prs.slides.add_slide(blank)
        pic = s.shapes.add_picture(str(pg["base"]), 0, 0, prs.slide_width, prs.slide_height)
        pic.name = f"絵 {pg['no']}"
        pic._element.nvPicPr.cNvPr.set("descr", pg["alt"])
        if pg["plate"]:
            r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, _emu(L.BAND_TOP), prs.slide_width, _emu(L.BAND_H))
            r.name = "帯の板"
            r.fill.solid()
            _color(r.fill.fore_color, theme["background"])
            r.line.fill.background()
            r.shadow.inherit = False
        number = pg["number"]
        place = pg.get("logo") or {"off": True}
        on = bool(logo and not place.get("off"))
        foot = footer if pg["chrome"] else ""
        boxes = C.chrome_boxes(foot, number, place if on else None, font_file, pg.get("mid", L.LOGO_MID_Y))
        if foot:
            x, y, _, h = boxes["footer"]
            fw = L.footer_box(place["w"] if on and place["corner"] == "bottom-left" else 0, bool(number))[2]
            _textbox(s, (x, y, fw, h), foot, font=theme["font"], color=pg["ink"]["footer"]).name = "フッタ"
        if number:
            box = boxes["number"]
            _textbox(s, (L.W // 2 - 60, box[1], 120, box[3]), number, font=theme["font"], color=pg["ink"]["number"],
                     align_center=True, field_slidenum=True).name = "ページ番号"
        if on:
            x, y, w, h = place["x"], place["y"], place["w"], place["h"]
            if place["plate"]:
                px, py, pw, ph = place["plate_box"]
                shape = MSO_SHAPE.ROUNDED_RECTANGLE if place.get("sticker") else MSO_SHAPE.RECTANGLE
                r = s.shapes.add_shape(shape, _emu(px), _emu(py), _emu(pw), _emu(ph))
                if place.get("sticker"):
                    r.adjustments[0] = L.LOGO_STICKER_RADIUS / min(pw, ph)
                r.name = "ロゴの地"
                r.fill.solid()
                _color(r.fill.fore_color, place["plate"])
                r.line.fill.background()
                r.shadow.inherit = False
            # 縦横比は原本のまま（四捨五入した w・h ではなく、原寸 × 倍率で置く）
            sw, sh = logo["size"]
            s.shapes.add_picture(str(C.logo_file(logo, place)), _emu(x), _emu(y), _emu(sw * logo["scale"]),
                                 _emu(sh * logo["scale"])).name = "ロゴ"
        if pg["notes"]:
            s.notes_slide.notes_text_frame.text = pg["notes"]
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_suffix(".tmp.pptx")
    prs.save(str(tmp))
    tmp.replace(out)


def write_pdf(pngs: list[pathlib.Path], out: pathlib.Path) -> None:
    """確認用の合成画像（帯の部品を焼いたもの）を1冊の PDF にする。1920px を 13.333 インチに（144dpi）。
    Pillow は JPEG で詰めるので、絵の中の文字が荒れないよう画質を上げる。"""
    from PIL import Image
    ims = [Image.open(p).convert("RGB") for p in pngs]
    try:
        tmp = out.with_suffix(".tmp.pdf")
        ims[0].save(str(tmp), "PDF", resolution=144.0, save_all=True, append_images=ims[1:], quality=92)
        tmp.replace(out)
    finally:
        for im in ims:
            im.close()

