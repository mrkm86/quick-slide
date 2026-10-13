#!/usr/bin/env python3
"""出力（--format・PDF・ページの画像・サムネイル）のテスト（ADR 0012）。

    python3 test_outputs.py <ケース名>     # 1ケース
    python3 test_outputs.py all           # 全部

deck_formats だけ Chrome（node・puppeteer-core）を使う。PDF のページ数は pdfinfo があれば数える。
"""
from __future__ import annotations

import json
import pathlib
import shutil
import subprocess
import sys
import tempfile

SKILL = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(SKILL / "scripts"))
import html_slide as H  # noqa: E402
import outputs as OUT  # noqa: E402

SAMPLE = SKILL / "examples" / "sample-deck.json"
TMP = pathlib.Path(tempfile.mkdtemp(prefix="html-slide-outputs-test-"))


def case_formats():
    """--format: 省けば4つ・並びは決まった順・誤りは ValueError。"""
    assert OUT.formats_of(None) == ("pptx", "pdf", "png", "thumbs")
    assert OUT.formats_of("thumbs, PNG") == ("png", "thumbs")
    for bad in ("jpg", "png,jpg", ","):
        try:
            OUT.formats_of(bad)
        except ValueError as e:
            assert "--format は" in str(e)
        else:
            raise AssertionError(f"{bad!r} が通った")


def case_thumbs():
    """サムネイル: 16ページまでは thumbs.png・超えたら thumbs-01.png …・前の回の分は消す・見出しの選び方。"""
    from PIL import Image
    d = TMP / "thumbs"
    d.mkdir()
    pngs = []
    for i in range(18):
        q = d / f"slide-{i + 1:02d}.png"
        Image.new("RGB", (1920, 1080), (200, 200 - i * 5, 220)).save(q)
        pngs.append(q)
    got = OUT.write_thumbs(pngs[:10], ["見出し"] * 10, d, "試験")
    assert [p.name for p in got] == ["thumbs.png"]
    w, h = Image.open(got[0]).size
    assert w == OUT.MARGIN * 2 + OUT.COLS * OUT.THUMB_W + (OUT.COLS - 1) * OUT.GAP and h < w
    got = OUT.write_thumbs(pngs, ["とても長い見出しが二行に入りきらないときは字を縮めて、それでも入らなければ切る" * 2] * 18, d, "試験")
    assert [p.name for p in got] == ["thumbs-01.png", "thumbs-02.png"]
    assert not (d / "thumbs.png").exists(), "1枚だった前の回の thumbs.png が残った"
    OUT.write_thumbs(pngs[:3], ["x"] * 3, d, "試験")
    assert sorted(p.name for p in d.glob("thumbs*.png")) == ["thumbs.png"], "2枚だった前の回の分が残った"
    assert OUT.label_of({"type": "statement", "headline": "**強調**の\n見出し"}) == "強調の見出し"
    assert OUT.label_of({"type": "cover", "title": "表紙"}) == "表紙"
    assert OUT.label_of({"type": "big-number", "value": "50", "unit": "%"}) == "50%"
    assert OUT.label_of({"type": "agenda"}) == "agenda"
    OUT.prune_pages(d, 5)
    assert sorted(p.name for p in d.glob("slide-*.png")) == [f"slide-{i:02d}.png" for i in range(1, 6)]


def case_print_html():
    """PDF 用の写し: 不透明な PNG は JPEG に・透過のある絵は PNG のまま・相対パスは絶対の file: URL・data: は触らない。"""
    from PIL import Image
    d = TMP / "print"
    (d / "assets").mkdir(parents=True)
    Image.new("RGB", (3000, 2000), (10, 120, 200)).save(d / "assets" / "photo.png")
    Image.new("RGBA", (300, 40), (0, 0, 0, 0)).save(d / "assets" / "logo.png")
    Image.new("RGBA", (300, 40), (5, 5, 5, 255)).save(d / "assets" / "opaque-rgba.png")
    html = d / "deck.html"
    html.write_text('<section data-slide><img class="photo" src="assets/photo.png" /><img src="assets/photo.png" />'
                    '<img class="logo" src="assets/logo.png" /><img src="assets/opaque-rgba.png" />'
                    '<img src="data:image/png;base64,AAAA" /></section>', encoding="utf-8")
    t = d / "tmp"
    t.mkdir()
    out = H.print_html(html, t).read_text(encoding="utf-8")
    srcs = [s.split('"')[0] for s in out.split('src="')[1:]]
    assert srcs[0].startswith("file://") and srcs[0].endswith(".jpg"), srcs
    assert srcs[1] == srcs[0], "同じ絵を2回変換した"
    assert srcs[2] == (d / "assets" / "logo.png").resolve().as_uri(), "透過のあるロゴを JPEG にした"
    assert srcs[3].endswith(".jpg"), "全部不透明な RGBA の PNG は JPEG にしてよい"
    assert srcs[4] == "data:image/png;base64,AAAA"
    with Image.open(t / pathlib.Path(srcs[0]).name) as im:
        assert max(im.size) == H.PDF_LONG, im.size


def _pages(pdf: pathlib.Path) -> int | None:
    if not shutil.which("pdfinfo"):
        return None
    r = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True, encoding="utf-8")
    return int(next(ln.split()[-1] for ln in r.stdout.splitlines() if ln.startswith("Pages:")))


def case_deck_formats():
    """deck（Chrome）: 既定で4つ・--format png,thumbs,pdf は PPTX を作らない・減ったページの画像を消す・PDF は 13.333×7.5 インチ。"""
    work = TMP / "deck"
    work.mkdir()
    shutil.copytree(SAMPLE.parent / "assets", work / "assets")
    sc = json.loads(SAMPLE.read_text(encoding="utf-8"))
    n = len(sc["slides"])
    script = work / "deck.json"
    script.write_text(json.dumps(sc, ensure_ascii=False), encoding="utf-8")
    out = work / "out"
    (out / "slides").mkdir(parents=True)
    (out / "slides" / f"slide-{n + 1:02d}.png").write_bytes(b"old")
    rc = H.main(["deck", "--script", str(script), "--out", str(out), "--no-odp", "--no-fill", "--format", "png,thumbs,pdf"])
    assert rc == 0, rc
    assert not (out / "deck.pptx").exists(), "pptx を選んでいないのに作った"
    assert sorted(p.name for p in (out / "slides").glob("*.png")) == [f"slide-{i:02d}.png" for i in range(1, n + 1)]
    sheets = sorted(p.name for p in out.glob("thumbs*.png"))       # 見本は16ページを超えるので2枚
    assert sheets == (["thumbs.png"] if n <= 16 else [f"thumbs-{i:02d}.png" for i in range(1, (n + 15) // 16 + 1)]), sheets
    assert (out / "deck.pdf").is_file()
    assert _pages(out / "deck.pdf") in (None, n)
    info = json.loads((out / "deck-info.json").read_text(encoding="utf-8"))
    assert info["outputs"]["format"] == ["pdf", "png", "thumbs"] and len(info["outputs"]["png"]) == n
    if shutil.which("pdfinfo"):
        r = subprocess.run(["pdfinfo", str(out / "deck.pdf")], capture_output=True, text=True, encoding="utf-8")
        assert "960 x 540 pts" in r.stdout, "PDF のページが PPTX と同じ 13.333×7.5 インチでない"
    rc = H.main(["deck", "--script", str(script), "--out", str(out), "--no-odp", "--no-fill"])
    assert rc == 0 and (out / "deck.pptx").is_file(), rc
    assert json.loads((out / "deck-info.json").read_text(encoding="utf-8"))["outputs"]["format"] == ["pptx", "pdf", "png", "thumbs"]
    assert H.main(["deck", "--script", str(script), "--out", str(out), "--no-fill", "--format", "gif"]) == 2


def case_assemble():
    """--parts と assemble（ADR 0024）: 2つの回（テーマ違い）のページを混ぜて PPTX・PDF を組む。ページの中身は元のまま・
    読めない PDF は止める・指示の誤りは rc=2。"""
    import pdfpages as PDF
    from pptx import Presentation
    work = TMP / "assemble"
    work.mkdir()
    shutil.copytree(SAMPLE.parent / "assets", work / "assets")
    sc = json.loads(SAMPLE.read_text(encoding="utf-8"))
    sc["slides"] = sc["slides"][:4]
    script = work / "deck.json"
    script.write_text(json.dumps(sc, ensure_ascii=False), encoding="utf-8")
    outs = {}
    for theme in ("default", "forest"):
        out = work / theme
        rc = H.main(["deck", "--script", str(script), "--out", str(out), "--no-odp", "--no-fill", "--theme", theme,
                     "--format", "pptx,pdf", "--parts", str(out / "parts.json")])
        assert rc == 0, rc
        parts = json.loads((out / "parts.json").read_text(encoding="utf-8"))
        assert len(parts["slides"]) == 4 and PDF.page_count(out / "deck.pdf") == 4
        outs[theme] = (out, parts)
    (a, pa), (b, pb) = outs["default"], outs["forest"]
    picks = [(a, pa, 1), (b, pb, 2), (a, pa, 3)]
    man = work / "manifest.json"
    man.write_text(json.dumps({"title": "混ぜた資料", "size": pa["size"], "slides": [p["slides"][n - 1] for _, p, n in picks],
                               "pdf": [[str(d / "deck.pdf"), n] for d, _, n in picks]}, ensure_ascii=False), encoding="utf-8")
    mixed = work / "mixed"
    assert H.main(["assemble", "--manifest", str(man), "--out", str(mixed)]) == 0
    assert PDF.page_count(mixed / "deck.pdf") == 3
    for i, (d, _, n) in enumerate(picks, 1):
        assert PDF.page_streams(mixed / "deck.pdf", i) == PDF.page_streams(d / "deck.pdf", n), f"{i} ページ目が元と違う"
    assert PDF.page_streams(a / "deck.pdf", 2) != PDF.page_streams(b / "deck.pdf", 2), "テーマ違いなのに同じページ"

    def texts(path, n):
        sl = Presentation(str(path)).slides[n - 1]
        return [sh.text_frame.text for sh in sl.shapes if getattr(sh, "has_text_frame", False)], \
            sl.notes_slide.notes_text_frame.text if sl.has_notes_slide else ""
    for i, (d, _, n) in enumerate(picks, 1):
        assert texts(mixed / "deck.pptx", i) == texts(d / "deck.pptx", n), f"{i} 枚目の文字・ノートが元と違う"
    # PDF 1.5 の相互参照のストリームは読まない（混ぜずに止める）
    bad = work / "xref-stream.pdf"
    bad.write_bytes(b"%PDF-1.5\n1 0 obj\n<</Type /XRef /Size 1>>\nstream\n\nendstream\nendobj\nstartxref\n9\n%%EOF\n")
    try:
        PDF.assemble([(bad, 1)], work / "never.pdf")
    except PDF.PdfError:
        pass
    else:
        raise AssertionError("相互参照のストリームを読んだ")
    man.write_text(json.dumps({"slides": [], "pdf": [[str(bad), 1]]}), encoding="utf-8")
    assert H.main(["assemble", "--manifest", str(man), "--out", str(work / "bad")]) == 1
    man.write_text("{", encoding="utf-8")
    assert H.main(["assemble", "--manifest", str(man), "--out", str(work / "bad")]) == 2


CASES = {k[5:]: v for k, v in globals().items() if k.startswith("case_")}

if __name__ == "__main__":
    name = sys.argv[1] if len(sys.argv) > 1 else ""
    if name not in CASES and name != "all":
        print("使えるケース: all, " + ", ".join(CASES))
        sys.exit(2)
    try:
        for k, f in CASES.items():
            if name in ("all", k):
                f()
                print(f"ok   {k}")
    finally:
        shutil.rmtree(TMP, ignore_errors=True)
