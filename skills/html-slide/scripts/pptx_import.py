#!/usr/bin/env python3
"""既存の PPTX → 取り込みの台本（`html` 型のページの台本）。入口は `html_slide.py import`。

    python3 pptx_import.py --pptx <file.pptx> --out <フォルダ>

<フォルダ> に書くもの:
- `script.json`   … 台本。1ページ = `{"type": "html", "html": …, "background": …, "notes": …, "title": …}`
- `media/`        … ページで使う画像（切り抜き・形の切り抜き・回転・透明度は焼き込む）
- `theme.yaml`    … PPTX のテーマの色と字体（あとから型のページを足すとき、元のページと色をそろえる）
- `import-report.json` … 写せなかったもの・近似したもの（ページごと）

写し方の決まりは `docs/adr/0025-import-pptx.md`。座標は 1920×1080 に写す（16:9 でなければ中央に置き、余りは背景で埋める）。
終了コード: 0 = 写せた（報告があっても 0）/ 2 = PPTX として読めない・大きすぎる / 1 = 実行時エラー。
"""
from __future__ import annotations

import argparse
import hashlib
import html as _html
import io
import json
import math
import pathlib
import re
import sys
import zipfile

sys.dont_write_bytecode = True
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import ooxml as X  # noqa: E402
import ooxml_geom as GEO  # noqa: E402
from ooxml import A, NS, P, R  # noqa: E402

EMU_PER_PX = 6350
SLIDE_W, SLIDE_H = 1920, 1080
# 外から来たファイルの上限（展開した大きさ・数・1 つの XML）
MAX_TOTAL_BYTES = 2 * 1024 ** 3
MAX_ENTRIES = 20000
MAX_XML_BYTES = 64 * 1024 ** 2
MAX_RATIO = 200                       # 1 つの部品の圧縮率（zip 爆弾）
IMAGE_DENSITY = 2.0                   # 画像は表示の大きさの 2 倍まで（元より大きくはしない）
LINE_HEIGHT = 1.2                     # 行間 100% の行送り（字の大きさに対する倍率）

# 字体の名前は単引用符で囲む（style="…" の属性の中に書くため）
JA_FALLBACK = ["'Hiragino Sans'", "'Hiragino Kaku Gothic ProN'", "'Yu Gothic'", "'YuGothic'", "'Meiryo'", "'Noto Sans JP'", "sans-serif"]
JA_SERIF_FALLBACK = ["'Hiragino Mincho ProN'", "'Yu Mincho'", "'YuMincho'", "'Noto Serif JP'", "serif"]
FONT_ALIASES = {
    "游ゴシック": ["Yu Gothic", "YuGothic"], "游ゴシック Light": ["Yu Gothic Light", "YuGothic"],
    "游ゴシック Medium": ["Yu Gothic Medium", "YuGothic"], "Yu Gothic UI": ["Yu Gothic", "YuGothic"],
    "メイリオ": ["Meiryo"], "ＭＳ Ｐゴシック": ["MS PGothic"], "ＭＳ ゴシック": ["MS Gothic"],
    "ＭＳ 明朝": ["MS Mincho"], "ＭＳ Ｐ明朝": ["MS PMincho"], "游明朝": ["Yu Mincho", "YuMincho"],
    "BIZ UDPゴシック": ["BIZ UDPGothic"], "BIZ UDゴシック": ["BIZ UDGothic"], "Meiryo UI": ["Meiryo"],
    "UD デジタル 教科書体 N-B": ["UD Digi Kyokasho N-B"], "UD デジタル 教科書体 N-R": ["UD Digi Kyokasho N-R"],
}
SERIF_HINT = re.compile(r"明朝|Mincho|Serif|Times|Georgia|Garamond|Century", re.I)
CJK = re.compile(r"[　-ヿ㐀-鿿豈-﫿＀-￯]")
WINGDINGS = {"§": "▪", "Ø": "➢", "ü": "✔", "l": "●", "n": "■", "q": "❑", "v": "❖", "Ÿ": "•", "w": "◆", "o": "□",
             "u": "◆", "à": "→", "è": "➔", "ð": "⇨", "Ú": "▶", "p": "◻", "¡": "○"}


class ImportError_(Exception):
    """PPTX として読めない（rc=2）。"""


def esc(s: str) -> str:
    return _html.escape(s, quote=True)


def f2(v: float) -> str:
    s = f"{v:.2f}".rstrip("0").rstrip(".")
    return "0" if s in ("-0", "") else s


# --- 報告 ------------------------------------------------------------------------------

class Report:
    def __init__(self):
        self.items: list[dict] = []

    def add(self, slide: int, what: str, how: str):
        key = (slide, what, how)
        for it in self.items:
            if (it["slide"], it["what"], it["how"]) == key:
                it["count"] += 1
                return
        self.items.append({"slide": slide, "what": what, "how": how, "count": 1})

    def for_slide(self, no: int) -> list[dict]:
        return [it for it in self.items if it["slide"] == no]


# --- 座標 ------------------------------------------------------------------------------

class Affine:
    """2 次元のアフィン変換（a c e / b d f）。"""

    def __init__(self, a=1.0, b=0.0, c=0.0, d=1.0, e=0.0, f=0.0):
        self.m = (a, b, c, d, e, f)

    def __mul__(self, o: "Affine") -> "Affine":
        a1, b1, c1, d1, e1, f1 = self.m
        a2, b2, c2, d2, e2, f2_ = o.m
        return Affine(a1 * a2 + c1 * b2, b1 * a2 + d1 * b2, a1 * c2 + c1 * d2, b1 * c2 + d1 * d2,
                      a1 * e2 + c1 * f2_ + e1, b1 * e2 + d1 * f2_ + f1)

    def apply(self, x, y):
        a, b, c, d, e, f = self.m
        return a * x + c * y + e, b * x + d * y + f

    @staticmethod
    def translate(x, y):
        return Affine(1, 0, 0, 1, x, y)

    @staticmethod
    def scale(sx, sy):
        return Affine(sx, 0, 0, sy, 0, 0)

    @staticmethod
    def rotate(deg):
        r = math.radians(deg)
        return Affine(math.cos(r), math.sin(r), -math.sin(r), math.cos(r), 0, 0)


class Frame:
    """グループをたどった変換。子の EMU 座標 → スライドの EMU 座標。回転と反転は向きの足し算で持つ。"""

    def __init__(self, m: Affine | None = None, sx=1.0, sy=1.0, rot=0.0, fh=False, fv=False):
        self.m = m or Affine()
        self.sx, self.sy, self.rot, self.fh, self.fv = sx, sy, rot, fh, fv

    def place(self, x, y, w, h, rot, fh, fv):
        """子の箱 → (中心 x, 中心 y, 幅, 高さ, 回転, 左右反転, 上下反転)（スライドの EMU）。"""
        cx, cy = self.m.apply(x + w / 2, y + h / 2)
        flip = self.fh != self.fv
        r = (-rot if flip else rot) + self.rot
        return cx, cy, w * self.sx, h * self.sy, r % 360, fh != self.fh, fv != self.fv

    def group(self, off, ext, choff, chext, rot, fh, fv) -> "Frame":
        sx = ext[0] / chext[0] if chext[0] else 1.0
        sy = ext[1] / chext[1] if chext[1] else 1.0
        cx, cy = off[0] + ext[0] / 2, off[1] + ext[1] / 2
        local = Affine.translate(off[0], off[1]) * Affine.scale(sx, sy) * Affine.translate(-choff[0], -choff[1])
        flipm = Affine.translate(cx, cy) * Affine.scale(-1 if fh else 1, -1 if fv else 1) * Affine.translate(-cx, -cy)
        rotm = Affine.translate(cx, cy) * Affine.rotate(rot) * Affine.translate(-cx, -cy)
        m = self.m * rotm * flipm * local
        flip = self.fh != self.fv
        return Frame(m, self.sx * sx, self.sy * sy, ((-rot if flip else rot) + self.rot) % 360, self.fh != fh, self.fv != fv)


class Geo:
    """スライドの EMU → 1920×1080 の px。16:9 でなければ中央に置く。"""

    def __init__(self, cx: int, cy: int):
        wpx, hpx = cx / EMU_PER_PX, cy / EMU_PER_PX
        self.s = min(SLIDE_W / wpx, SLIDE_H / hpx)
        self.w, self.h = wpx * self.s, hpx * self.s
        self.ox, self.oy = (SLIDE_W - self.w) / 2, (SLIDE_H - self.h) / 2
        self.emu_w, self.emu_h = cx, cy
        self.wide = abs(self.w - SLIDE_W) < 1 and abs(self.h - SLIDE_H) < 1

    def x(self, emu):
        return emu / EMU_PER_PX * self.s + self.ox

    def y(self, emu):
        return emu / EMU_PER_PX * self.s + self.oy

    def len(self, emu):
        return emu / EMU_PER_PX * self.s

    def pt(self, pt):
        """pt → px（1px = 0.5pt。縮めたぶんも掛ける）。"""
        return pt * 2 * self.s


# --- 経路（SVG の d）を点に -----------------------------------------------------------------

def flatten(d: str) -> list[list[tuple[float, float]]]:
    """絶対座標の M L H V C Q A Z を折れ線に（画像の形の切り抜き・矢印の向きに使う）。"""
    toks = re.findall(r"[MLHVCQAZ]|-?[\d.]+(?:e-?\d+)?", d)
    subs: list[list[tuple[float, float]]] = []
    cur: list[tuple[float, float]] = []
    i, cmd, x, y = 0, None, 0.0, 0.0

    def num():
        nonlocal i
        v = float(toks[i])
        i += 1
        return v

    while i < len(toks):
        if re.fullmatch(r"[MLHVCQAZ]", toks[i]):
            cmd = toks[i]
            i += 1
        if cmd == "Z":
            if cur:
                subs.append(cur)
            cur = []
            continue
        try:
            if cmd == "M":
                if cur:
                    subs.append(cur)
                x, y = num(), num()
                cur = [(x, y)]
                cmd = "L"
            elif cmd == "L":
                x, y = num(), num()
                cur.append((x, y))
            elif cmd == "H":
                x = num()
                cur.append((x, y))
            elif cmd == "V":
                y = num()
                cur.append((x, y))
            elif cmd == "C":
                x1, y1, x2, y2, x3, y3 = (num() for _ in range(6))
                for k in range(1, 13):
                    t = k / 12
                    mt = 1 - t
                    cur.append((mt ** 3 * x + 3 * mt * mt * t * x1 + 3 * mt * t * t * x2 + t ** 3 * x3,
                                mt ** 3 * y + 3 * mt * mt * t * y1 + 3 * mt * t * t * y2 + t ** 3 * y3))
                x, y = x3, y3
            elif cmd == "Q":
                x1, y1, x2, y2 = (num() for _ in range(4))
                for k in range(1, 9):
                    t = k / 8
                    mt = 1 - t
                    cur.append((mt * mt * x + 2 * mt * t * x1 + t * t * x2, mt * mt * y + 2 * mt * t * y1 + t * t * y2))
                x, y = x2, y2
            elif cmd == "A":
                rx, ry, phi, fa, fs, x2, y2 = (num() for _ in range(7))
                cur.extend(_arc_points(x, y, rx, ry, phi, int(fa), int(fs), x2, y2))
                x, y = x2, y2
            else:
                i += 1
        except (IndexError, ValueError):
            break
    if cur:
        subs.append(cur)
    return subs


def _arc_points(x1, y1, rx, ry, phi, fa, fs, x2, y2, n=24):
    if rx == 0 or ry == 0:
        return [(x2, y2)]
    p = math.radians(phi)
    cp, sp = math.cos(p), math.sin(p)
    dx, dy = (x1 - x2) / 2, (y1 - y2) / 2
    x1p, y1p = cp * dx + sp * dy, -sp * dx + cp * dy
    rx, ry = abs(rx), abs(ry)
    lam = x1p ** 2 / rx ** 2 + y1p ** 2 / ry ** 2
    if lam > 1:
        rx, ry = rx * math.sqrt(lam), ry * math.sqrt(lam)
    num_ = rx * rx * ry * ry - rx * rx * y1p * y1p - ry * ry * x1p * x1p
    den = rx * rx * y1p * y1p + ry * ry * x1p * x1p
    co = math.sqrt(max(0.0, num_ / den)) if den else 0.0
    if fa == fs:
        co = -co
    cxp, cyp = co * rx * y1p / ry, -co * ry * x1p / rx
    cx, cy = cp * cxp - sp * cyp + (x1 + x2) / 2, sp * cxp + cp * cyp + (y1 + y2) / 2

    def ang(ux, uy, vx, vy):
        a = math.atan2(ux * vy - uy * vx, ux * vx + uy * vy)
        return a

    t1 = ang(1, 0, (x1p - cxp) / rx, (y1p - cyp) / ry)
    dt = ang((x1p - cxp) / rx, (y1p - cyp) / ry, (-x1p - cxp) / rx, (-y1p - cyp) / ry)
    if not fs and dt > 0:
        dt -= 2 * math.pi
    elif fs and dt < 0:
        dt += 2 * math.pi
    out = []
    for k in range(1, n + 1):
        t = t1 + dt * k / n
        out.append((cx + rx * math.cos(t) * cp - ry * math.sin(t) * sp, cy + rx * math.cos(t) * sp + ry * math.sin(t) * cp))
    return out


# --- 画像 ------------------------------------------------------------------------------

class Media:
    """ページで使う画像を media/ に書く。同じ中身・同じ加工は 1 つにまとめる。"""

    def __init__(self, out: pathlib.Path, report: Report):
        self.dir = out / "media"
        self.report = report
        self.done: dict[str, str] = {}
        self.svg_jobs: list[tuple[int, bytes, dict, str]] = []

    @staticmethod
    def _turned(size, rot):
        w, h = size
        if rot % 360:
            r = math.radians(rot)
            w, h = abs(size[0] * math.cos(r)) + abs(size[1] * math.sin(r)), abs(size[0] * math.sin(r)) + abs(size[1] * math.cos(r))
        return w, h

    def svg(self, no: int, blob: bytes, **ops) -> tuple[str, float, float]:
        """SVG の画像（PNG の控えが無いもの）。名前だけ先に決め、`finish` で Chrome に描かせてから加工する。"""
        key = hashlib.sha256(blob + repr(sorted(ops.items())).encode()).hexdigest()[:16]
        if key not in self.done:
            self.done[key] = f"media/s{no:02d}-{key}.png"
            self.svg_jobs.append((no, blob, ops, self.done[key]))
        rel = self.done[key]
        w, h = self._turned(ops.get("size", (100, 100)), ops.get("rot", 0.0))
        return rel, w, h

    def finish(self) -> None:
        """溜めた SVG を 1 回の Chrome でまとめて PNG にし、切り抜き・回転などを掛ける。描けなければ透明の画像にして報告する。"""
        if not self.svg_jobs:
            return
        import subprocess
        import tempfile
        from PIL import Image
        self.dir.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="qs-svg-") as tmp:
            jobs = []
            for i, (no, blob, ops, rel) in enumerate(self.svg_jobs):
                src = pathlib.Path(tmp) / f"{i}.svg"
                src.write_bytes(blob)
                w, h = ops.get("size", (100, 100))
                k = min(IMAGE_DENSITY, 4096 / max(w, h, 1))
                jobs.append({"svg": str(src), "png": str(pathlib.Path(tmp) / f"{i}.png"), "w": max(1, w * k), "h": max(1, h * k)})
            (pathlib.Path(tmp) / "jobs.json").write_text(json.dumps(jobs))
            ok = False
            try:
                import browser as B
                node, _ = B.node()
                if node and not B.missing():
                    r = subprocess.run([node, str(pathlib.Path(__file__).resolve().parent / "svg_raster.mjs"),
                                        str(pathlib.Path(tmp) / "jobs.json")], capture_output=True, text=True,
                                       timeout=600, env=B.child_env())
                    ok = r.returncode == 0
            except (OSError, subprocess.SubprocessError):
                ok = False
            for (no, blob, ops, rel), j in zip(self.svg_jobs, jobs):
                png = pathlib.Path(j["png"])
                dest = self.dir / pathlib.Path(rel).name
                if ok and png.is_file():
                    self.picture(no, png.read_bytes(), ".png", name=dest.name, **ops)
                    if dest.is_file():
                        continue
                self.report.add(no, "画像（SVG）", "描けないので透明にした（Chrome が要る）")
                Image.new("RGBA", (2, 2), (0, 0, 0, 0)).save(dest)

    def picture(self, no: int, blob: bytes, ext: str, *, crop=(0, 0, 0, 0), size=(100, 100), mask: str | None = None,
                rot: float = 0.0, fh=False, fv=False, alpha: float = 1.0, gray=False, duo=None, clr_change=None,
                tile=False, name: str | None = None) -> tuple[str, float, float] | None:
        """画像を加工して書く。(media/ からの相対パス, 書いた画像の表示の幅, 高さ px) を返す。読めなければ None。

        size は表示の大きさ（px）。crop は (左, 上, 右, 下) の割合（負は余白）。mask は形の d（size の座標）。
        回転した画像は、回した外接の箱の大きさになる（呼ぶ側は中心をそろえて置く）。"""
        from PIL import Image, ImageDraw, ImageOps
        key = hashlib.sha256(blob + repr((crop, size, mask, round(rot, 2), fh, fv, alpha, gray, duo, clr_change, tile)).encode()).hexdigest()[:16]
        if key in self.done and name is None:
            path = self.done[key]
            w, h = size
            if rot % 360:
                r = math.radians(rot)
                w, h = abs(size[0] * math.cos(r)) + abs(size[1] * math.sin(r)), abs(size[0] * math.sin(r)) + abs(size[1] * math.cos(r))
            return path, w, h
        try:
            im = Image.open(io.BytesIO(blob))
            im.load()
        except Exception:
            self.report.add(no, f"画像（{ext or '形式不明'}）", "読めないので空の枠にした（EMF・WMF などは写せない）")
            return None
        if getattr(im, "n_frames", 1) > 1:
            im.seek(0)
        jpeg_src = (im.format == "JPEG")
        im = ImageOps.exif_transpose(im)
        im = im.convert("RGBA")
        iw, ih = im.size
        l_, t_, r_, b_ = crop
        if any(crop):
            # 割合の切り抜き（負の値は余白を足す）
            x0, y0 = l_ * iw, t_ * ih
            x1, y1 = iw - r_ * iw, ih - b_ * ih
            if x1 - x0 < 1 or y1 - y0 < 1:
                return None
            canvas = Image.new("RGBA", (max(1, round(x1 - x0)), max(1, round(y1 - y0))), (0, 0, 0, 0))
            src = im.crop((max(0, round(x0)), max(0, round(y0)), min(iw, round(x1)), min(ih, round(y1))))
            canvas.paste(src, (max(0, -round(x0)), max(0, -round(y0))))
            im = canvas
        dw, dh = max(1.0, size[0]), max(1.0, size[1])
        # 画素の数: 表示の 2 倍まで。元の画素より大きくはしない（縦横比は表示に合わせる＝伸ばす）
        k = min(IMAGE_DENSITY, max(im.width / dw, im.height / dh))
        k = max(k, 1.0) if max(im.width / dw, im.height / dh) >= 1 else max(im.width / dw, im.height / dh)
        tw, th = max(1, round(dw * k)), max(1, round(dh * k))
        if tile:
            tiled = Image.new("RGBA", (tw, th))
            for ty in range(0, th, max(1, im.height)):
                for tx in range(0, tw, max(1, im.width)):
                    tiled.paste(im, (tx, ty))
            im = tiled
        else:
            im = im.resize((tw, th), Image.LANCZOS)
        if gray:
            g = ImageOps.grayscale(im.convert("RGB"))
            a = im.getchannel("A")
            im = Image.merge("RGBA", (g, g, g, a))
        if duo:
            c1, c2 = duo
            g = ImageOps.grayscale(im.convert("RGB"))
            a = im.getchannel("A")
            im = ImageOps.colorize(g, c1[:3], c2[:3]).convert("RGBA")
            im.putalpha(a)
        if clr_change:
            src_c, tol = clr_change, 24
            px = im.load()
            for yy in range(im.height):
                for xx in range(im.width):
                    r, g, b, a = px[xx, yy]
                    if abs(r - src_c[0]) <= tol and abs(g - src_c[1]) <= tol and abs(b - src_c[2]) <= tol:
                        px[xx, yy] = (r, g, b, 0)
        if fh:
            im = ImageOps.mirror(im)
        if fv:
            im = ImageOps.flip(im)
        if mask:
            m = Image.new("L", (tw, th), 0)
            dr = ImageDraw.Draw(m)
            sx, sy = tw / dw, th / dh
            for sub in flatten(mask):
                if len(sub) >= 3:
                    dr.polygon([(x * sx, y * sy) for x, y in sub], fill=255)
            a = im.getchannel("A")
            from PIL import ImageChops
            im.putalpha(ImageChops.multiply(a, m))
        if alpha < 0.999:
            a = im.getchannel("A").point(lambda v: round(v * alpha))
            im.putalpha(a)
        out_w, out_h = dw, dh
        if rot % 360:
            im = im.rotate(-rot, resample=Image.BICUBIC, expand=True)
            r = math.radians(rot)
            out_w = abs(dw * math.cos(r)) + abs(dh * math.sin(r))
            out_h = abs(dw * math.sin(r)) + abs(dh * math.cos(r))
        opaque = im.getchannel("A").getextrema()[0] >= 255
        self.dir.mkdir(parents=True, exist_ok=True)
        if opaque and jpeg_src and name is None:
            name = f"s{no:02d}-{key}.jpg"
            im.convert("RGB").save(self.dir / name, "JPEG", quality=90, optimize=True)
        else:
            name = name or f"s{no:02d}-{key}.png"
            if opaque:
                im = im.convert("RGB")
            im.save(self.dir / name, "PNG", optimize=True)
        rel = f"media/{name}"
        self.done.setdefault(key, rel)
        return rel, out_w, out_h


# --- 1 枚を読む ----------------------------------------------------------------------------

class SlideCtx:
    def __init__(self, deck: "Deck", no: int, slide, layout, master):
        self.deck, self.no = deck, no
        self.slide, self.layout, self.master = slide, layout, master
        self.sld = slide.part._element if slide is not None else None
        self.lay = layout.part._element
        self.mst = master.part._element
        self.theme = deck.theme_of(master)
        cmap = {}
        cm = X.q(self.mst, "p:clrMap")
        if cm is not None:
            cmap = dict(cm.attrib)
        for el in (self.lay, self.sld):
            if el is None:
                continue
            ov = X.q(el, "p:clrMapOvr/a:overrideClrMapping")
            if ov is not None:
                cmap.update(ov.attrib)
        self.colors = X.Colors(self.theme.scheme, cmap)
        self.geo = deck.geo
        self.report = deck.report
        self.media = deck.media
        self.ids = 0

    def rep(self, what, how):
        self.report.add(self.no, what, how)

    def part_of(self, owner: str):
        return {"slide": self.slide.part, "layout": self.layout.part, "master": self.master.part}[owner]


class Deck:
    def __init__(self, path: pathlib.Path, out: pathlib.Path):
        from pptx import Presentation
        self.path = path
        self.out = out
        self.report = Report()
        self.prs = Presentation(str(path))
        sz = X.q(self.prs.part._element, "p:sldSz")
        cx = X.ival(sz, "cx", 12192000)
        cy = X.ival(sz, "cy", 6858000)
        self.geo = Geo(cx, cy)
        self.media = Media(out, self.report)
        self._themes: dict = {}
        self.default_text = X.q(self.prs.part._element, "p:defaultTextStyle")
        self.table_styles = None
        for rel in self.prs.part.rels.values():
            if rel.reltype == X.RT_TABLE_STYLES and not rel.is_external:
                try:
                    self.table_styles = X.parse(rel.target_part.blob)
                except Exception:
                    self.table_styles = None

    def theme_of(self, master):
        key = id(master.part)
        if key not in self._themes:
            th = None
            for rel in master.part.rels.values():
                if rel.reltype == X.RT_THEME and not rel.is_external:
                    th = X.parse(rel.target_part.blob)
            self._themes[key] = X.Theme(th if th is not None else X.parse(b'<a:theme xmlns:a="%s"/>' % NS["a"].encode()))
        return self._themes[key]


def xfrm_of(sppr):
    x = X.q(sppr, "a:xfrm")
    if x is None:
        return None
    off, ext = X.q(x, "a:off"), X.q(x, "a:ext")
    if off is None or ext is None:
        return None
    return {"x": X.ival(off, "x", 0), "y": X.ival(off, "y", 0), "w": X.ival(ext, "cx", 0), "h": X.ival(ext, "cy", 0),
            "rot": X.ival(x, "rot", 0) / 60000, "fh": X.bval(x, "flipH"), "fv": X.bval(x, "flipV")}


# --- プレースホルダーの継承 -------------------------------------------------------------------

TITLE_TYPES = {"title", "ctrTitle"}


def ph_of(sp):
    nv = None
    for c in sp:
        if isinstance(c.tag, str) and X.local(c).startswith("nv"):
            nv = c
    if nv is None:
        return None
    ph = X.q(nv, "p:nvPr/p:ph")
    if ph is None:
        return None
    return {"type": ph.get("type") or "obj", "idx": ph.get("idx")}


def _tree_phs(root):
    out = []
    tree = X.q(root, "p:cSld/p:spTree")
    if tree is None:
        return out

    def walk(t):
        for c in X.unwrap_mc(t):
            n = X.local(c)
            if n == "grpSp":
                walk(c)
            elif n in ("sp", "pic", "graphicFrame"):
                ph = ph_of(c)
                if ph:
                    out.append((ph, c))
    walk(tree)
    return out


def match_ph(root, ph, by_idx=True):
    """レイアウト（かマスター）の中で、同じプレースホルダーを探す。idx が合うものを先に、無ければ種類で。"""
    cands = _tree_phs(root)
    if by_idx and ph.get("idx") is not None:
        for p, el in cands:
            if p.get("idx") == ph["idx"]:
                return el
    t = ph["type"]
    for p, el in cands:
        if p["type"] == t:
            return el
    if t in TITLE_TYPES:
        for p, el in cands:
            if p["type"] in TITLE_TYPES:
                return el
    if t in ("obj", "body", "subTitle"):
        for p, el in cands:
            if p["type"] in ("body", "obj"):
                return el
    return None


def master_ph(mst, ph):
    t = ph["type"]
    want = "title" if t in TITLE_TYPES else (t if t in ("dt", "ftr", "sldNum") else "body")
    for p, el in _tree_phs(mst):
        if p["type"] == want or (want == "body" and p["type"] == "obj"):
            return el
    return None


def tx_style_of(mst, ph):
    ts = X.q(mst, "p:txStyles")
    if ts is None:
        return None
    if ph is None:
        return X.q(ts, "p:otherStyle")
    t = ph["type"]
    if t in TITLE_TYPES:
        return X.q(ts, "p:titleStyle")
    if t in ("dt", "ftr", "sldNum", "hdr"):
        return X.q(ts, "p:otherStyle")
    return X.q(ts, "p:bodyStyle")


# --- 1 枚の HTML を組む -------------------------------------------------------------------------

class Builder:
    """1 枚ぶんの HTML を組む。`parts` は section の中身（奥から手前の順）。"""

    def __init__(self, ctx: SlideCtx):
        self.ctx = ctx
        self.parts: list[str] = []
        self.anim_of: dict[str, dict] = {}      # 図形の id → 段の属性（pptx_anim が埋める）
        self.texts_of: dict[str, list] = {}     # 図形の id → 段落ごとに分けた文字の要素（段落ごとの動き）

    # 置く ---
    def emit(self, name: str | None, pieces: list[str], an: dict | None = None):
        pieces = [p for p in pieces if p]
        if not pieces:
            return
        attrs = ""
        if name:
            attrs += f' data-name="{esc(name[:60])}"'
        if an:
            attrs += f' data-step="{an["step"]}" data-fx="{an["fx"]}"' + (f' data-delay="{an["delay"]}"' if an.get("delay") else "")
        self.parts.append(f'<div style="position:absolute;left:0;top:0;width:0;height:0"{attrs}>' + "".join(pieces) + "</div>")


def css_fill(fill: X.Fill | None) -> str | None:
    if fill is None or fill.kind in ("none", "group", "blip"):
        return None
    if fill.kind == "grad" and len(fill.stops) >= 2:
        stops = ", ".join(f"{X.css_color(c)} {f2(p * 100)}%" for p, c in fill.stops)
        return f"linear-gradient({f2(fill.angle + 90)}deg, {stops})"
    c = fill.main_color()
    return X.css_color(c) if c is not None else None


def svg_paint(fill: X.Fill | None, ctx: SlideCtx) -> str:
    """SVG の塗り（グラデーションは写せないので、真ん中の色で塗る）。"""
    if fill is None or fill.kind in ("none", "group", "blip"):
        return 'fill="none"'
    c = fill.main_color()
    if fill.kind == "grad" and fill.stops:
        c = grad_at(fill.stops, 0.5)
        ctx.rep("グラデーションの塗り（四角以外の形）", "真ん中の色の単色にした")
    if c is None:
        return 'fill="none"'
    if c[3] >= 0.999:
        return f'fill="{X.hex_of(c)}"'
    return f'fill="{X.hex_of(c)}" fill-opacity="{round(c[3], 3)}"'


def grad_at(stops, t: float):
    """グラデーションの t（0〜1）の位置の色。"""
    if t <= stops[0][0]:
        return stops[0][1]
    for (p0, c0), (p1, c1) in zip(stops, stops[1:]):
        if p0 <= t <= p1:
            k = (t - p0) / (p1 - p0) if p1 > p0 else 0
            return tuple(round(c0[i] + (c1[i] - c0[i]) * k) for i in range(3)) + (c0[3] + (c1[3] - c0[3]) * k,)
    return stops[-1][1]


def svg_stroke(line: X.Line, ctx: SlideCtx) -> str:
    if not line.visible:
        return ""
    w = max(0.5, ctx.geo.len(line.width))
    s = f' stroke="{X.hex_of(line.color)}" stroke-width="{f2(w)}"'
    if line.color[3] < 0.999:
        s += f' stroke-opacity="{round(line.color[3], 3)}"'
    if line.dash:
        if "dot" in line.dash.lower() and "dash" not in line.dash.lower():
            s += f' stroke-dasharray="{f2(w)},{f2(w)}"'
        else:
            s += f' stroke-dasharray="{f2(w * 4)},{f2(w * 3)}"'
    return s + ' stroke-linejoin="round"'


def arrow_head(kind: str, tip, back, w_line: float, color, size_w="med", size_l="med") -> str:
    """線の端の矢印。tip は先、back は手前の点（向き）。"""
    k = {"sm": 2.0, "med": 3.0, "lg": 5.0}
    lw = max(w_line, 1.0)
    hw, hl = lw * k.get(size_w, 3.0) / 2 * 1.4, lw * k.get(size_l, 3.0) * 1.1
    dx, dy = tip[0] - back[0], tip[1] - back[1]
    n = math.hypot(dx, dy) or 1.0
    ux, uy = dx / n, dy / n
    vx, vy = -uy, ux
    bx, by = tip[0] - ux * hl, tip[1] - uy * hl
    col = X.hex_of(color)
    op = f' fill-opacity="{round(color[3], 3)}"' if color[3] < 0.999 else ""
    if kind == "oval":
        return f'<ellipse cx="{f2(tip[0] - ux * hl / 2)}" cy="{f2(tip[1] - uy * hl / 2)}" rx="{f2(hl / 2)}" ry="{f2(hw)}" fill="{col}"{op}/>'
    if kind == "diamond":
        pts = [(tip[0], tip[1]), (bx + ux * hl / 2 + vx * hw, by + uy * hl / 2 + vy * hw), (bx, by), (bx + ux * hl / 2 - vx * hw, by + uy * hl / 2 - vy * hw)]
        return f'<polygon points="{" ".join(f"{f2(x)},{f2(y)}" for x, y in pts)}" fill="{col}"{op}/>'
    if kind == "arrow":
        pts = [(bx + vx * hw, by + vy * hw), tip, (bx - vx * hw, by - vy * hw)]
        return (f'<polyline points="{" ".join(f"{f2(x)},{f2(y)}" for x, y in pts)}" fill="none" stroke="{col}" '
                f'stroke-width="{f2(lw)}" stroke-linejoin="round" stroke-linecap="round"/>')
    pts = [tip, (bx + vx * hw, by + vy * hw), (bx - vx * hw, by - vy * hw)]
    if kind == "stealth":
        pts = [tip, (bx + vx * hw, by + vy * hw), (tip[0] - ux * hl * 0.6, tip[1] - uy * hl * 0.6), (bx - vx * hw, by - vy * hw)]
    return f'<polygon points="{" ".join(f"{f2(x)},{f2(y)}" for x, y in pts)}" fill="{col}"{op}/>'


class ShapeReader:
    """図形（sp・cxnSp・pic・grpSp・graphicFrame）を読んで HTML にする。"""

    def __init__(self, ctx: SlideCtx, b: Builder):
        self.ctx, self.b = ctx, b

    # 共通 ---
    def anim(self, sid, owner):
        """その図形の段（スライドの図形だけ。マスター・レイアウトの図形の id は別の数え方）。"""
        return self.b.anim_of.get(sid or "") if owner == "slide" and sid else None

    def para_anim(self, sid, owner):
        return self.b.texts_of.get(sid or "") if owner == "slide" and sid else None

    def walk(self, tree, frame: Frame, owner: str, skip_ph: bool):
        for el in X.unwrap_mc(tree):
            n = X.local(el)
            if n in ("nvGrpSpPr", "grpSpPr", "extLst"):
                continue
            try:
                if n == "grpSp":
                    self.group(el, frame, owner, skip_ph)
                elif n in ("sp", "cxnSp"):
                    if skip_ph and ph_of(el):
                        continue
                    self.shape(el, frame, owner)
                elif n == "pic":
                    if skip_ph and ph_of(el):
                        continue
                    self.picture(el, frame, owner)
                elif n == "graphicFrame":
                    if skip_ph and ph_of(el):
                        continue
                    self.graphic(el, frame, owner)
                elif n == "contentPart":
                    self.ctx.rep("手書き（インク）", "写さない")
            except Exception as e:  # 1 つの図形が読めなくても、ほかは写す
                self.ctx.rep("図形", f"読めないので飛ばした（{type(e).__name__}: {str(e)[:80]}）")

    def nv(self, el):
        for c in el:
            if isinstance(c.tag, str) and X.local(c).startswith("nv"):
                cnv = X.child(c, "cNvPr")
                if cnv is not None:
                    return cnv.get("id"), cnv.get("name"), X.bval(cnv, "hidden")
        return None, None, False

    def box(self, frame: Frame, xf):
        cx, cy, w, h, rot, fh, fv = frame.place(xf["x"], xf["y"], xf["w"], xf["h"], xf["rot"], xf["fh"], xf["fv"])
        g = self.ctx.geo
        pw, ph = g.len(w), g.len(h)
        return {"x": g.x(cx) - pw / 2, "y": g.y(cy) - ph / 2, "w": pw, "h": ph, "rot": rot, "fh": fh, "fv": fv}

    def outside(self, bx) -> bool:
        """元のスライドの外に丸ごと出ている（スライドショーでは見えない）。写さない。"""
        g = self.ctx.geo
        r = math.radians(bx["rot"])
        hw = (abs(bx["w"] * math.cos(r)) + abs(bx["h"] * math.sin(r))) / 2
        hh = (abs(bx["w"] * math.sin(r)) + abs(bx["h"] * math.cos(r))) / 2
        cx, cy = bx["x"] + bx["w"] / 2, bx["y"] + bx["h"] / 2
        out = cx + hw <= g.ox or cx - hw >= g.ox + g.w or cy + hh <= g.oy or cy - hh >= g.oy + g.h
        if out:
            self.ctx.rep("スライドの外の図形", "スライドショーで見えないので写さない")
        return out

    def group(self, el, frame: Frame, owner: str, skip_ph: bool):
        sid, name, hidden = self.nv(el)
        if hidden:
            return
        gp = X.child(el, "grpSpPr")
        x = X.q(gp, "a:xfrm")
        if x is None:
            sub = frame
        else:
            off, ext = X.q(x, "a:off"), X.q(x, "a:ext")
            choff, chext = X.q(x, "a:chOff"), X.q(x, "a:chExt")
            o = (X.ival(off, "x", 0), X.ival(off, "y", 0))
            e = (X.ival(ext, "cx", 0), X.ival(ext, "cy", 0))
            co = (X.ival(choff, "x", o[0]), X.ival(choff, "y", o[1]))
            ce = (X.ival(chext, "cx", e[0]), X.ival(chext, "cy", e[1]))
            sub = frame.group(o, e, co, ce, X.ival(x, "rot", 0) / 60000, X.bval(x, "flipH"), X.bval(x, "flipV"))
        an = self.anim(sid, owner)
        if an:
            # グループに動きがあれば、中身を 1 つの単位にまとめる（中の図形の動きは外側の段で出る）
            inner = Builder(self.ctx)
            ShapeReader(self.ctx, inner).walk(el, sub, owner, skip_ph)
            self.b.emit(name, inner.parts, an)
        else:
            self.walk(el, sub, owner, skip_ph)

    # 図形 ---
    def inherit(self, el, owner: str):
        """プレースホルダーなら (レイアウトの同じもの, マスターの同じもの, 種類)。"""
        ph = ph_of(el)
        if not ph or owner != "slide":
            return None, None, ph
        lay = match_ph(self.ctx.lay, ph)
        mst = master_ph(self.ctx.mst, ph) if lay is None else (master_ph(self.ctx.mst, ph_of(lay) or ph))
        return lay, mst, ph

    def shape(self, el, frame: Frame, owner: str):
        ctx = self.ctx
        sid, name, hidden = self.nv(el)
        if hidden:
            return
        sppr = X.child(el, "spPr")
        style = X.child(el, "style")
        lay, mst, ph = self.inherit(el, owner)
        xf = xfrm_of(sppr) or xfrm_of(X.child(lay, "spPr")) or xfrm_of(X.child(mst, "spPr"))
        if xf is None:
            return
        bx = self.box(frame, xf)
        if self.outside(bx):
            return
        geom_owner = sppr
        for cand in (sppr, X.child(lay, "spPr"), X.child(mst, "spPr")):
            if cand is not None and (X.q(cand, "a:prstGeom") is not None or X.q(cand, "a:custGeom") is not None):
                geom_owner = cand
                break
        fill = X.first_fill(sppr, ctx.colors)
        if fill is None:
            for inh in (lay, mst):
                if inh is not None:
                    fill = X.first_fill(X.child(inh, "spPr"), ctx.colors)
                    if fill is not None:
                        break
        if fill is None and style is not None:
            idx, refc = X.style_ref(style, "fillRef", ctx.colors)
            fill = X.ref_fill(ctx.theme, idx, refc, ctx.colors)
        if fill is not None and fill.kind == "patt":
            ctx.rep("模様の塗り", "前景と背景を混ぜた単色にした")
        line_src = sppr
        if X.q(sppr, "a:ln") is None:
            for inh in (lay, mst):
                if inh is not None and X.q(X.child(inh, "spPr"), "a:ln") is not None:
                    line_src = X.child(inh, "spPr")
                    break
        line = X.line_of(line_src, style, ctx.theme, ctx.colors)
        pieces = self.geometry(geom_owner, bx, fill, line, sid, el)
        tx = X.child(el, "txBody")
        an, pan = self.anim(sid, owner), self.para_anim(sid, owner)
        text = self.text(el, tx, bx, lay, mst, ph, style, owner, geom_owner, pan) if tx is not None else ""
        if pan:
            # 段落ごとに出る: 図形（地）は図形の段で、文字は段落ごとの段で（入れ子にしない）
            self.b.emit(name, pieces, an)
            self.b.emit((name or "") + "/文字", [text])
        else:
            self.b.emit(name, pieces + [text], an)

    def geometry(self, sppr, bx, fill, line, sid, el) -> list[str]:
        ctx = self.ctx
        prst = X.q(sppr, "a:prstGeom")
        cust = X.q(sppr, "a:custGeom")
        name = prst.get("prst") if prst is not None else ("custom" if cust is not None else "rect")
        adj = GEO.adjusts(prst) if prst is not None else {}
        w, h = bx["w"], bx["h"]
        has_fill = fill is not None and fill.kind not in ("none", "group")
        if not has_fill and not line.visible:
            return []
        if fill is not None and fill.kind == "blip":
            d = None
            if name != "rect" or bx["rot"]:
                d = GEO.preset(name, w, h, adj) if prst is not None else (GEO.custom(cust, w, h) or (None,))[0]
            out = []
            img = self.blip_img(fill.blip, bx, mask=d, el=el)
            if img:
                out.append(img)
            if line.visible:
                out.append(self.path_svg(GEO.preset(name, w, h, adj) or GEO.preset("rect", w, h, {}), bx, None, line))
            return out
        rot = bx["rot"] % 360
        straight = rot in (0, 180) or (rot in (90, 270) and abs(w - h) < 0.5)
        if name in ("rect", "roundRect", "flowChartProcess") and straight and not (fill and fill.radial):
            return [self.box_div(bx, fill, line, name, adj, el)]
        if name in ("ellipse", "flowChartConnector") and straight and not (fill and fill.kind == "grad"):
            return [self.path_svg(None, bx, fill, line, ellipse=True)]
        if cust is not None:
            got = GEO.custom(cust, w, h)
            if got is None:
                ctx.rep("自由な形", "読めないので四角にした")
                d, filled = GEO.preset("rect", w, h, {}), True
            else:
                d, filled = got
            return [self.path_svg(d, bx, fill if filled else None, line)]
        d = GEO.preset(name, w, h, adj)
        if d is None:
            ctx.rep(f"図形の形「{name}」", "写せる形に無いので四角にした")
            d = GEO.preset("rect", w, h, {})
        open_path = name in GEO.PATH_ONLY
        return [self.path_svg(d, bx, None if open_path else fill, line, arrows=open_path)]

    def box_div(self, bx, fill, line, name, adj, el) -> str:
        ctx = self.ctx
        x, y, w, h = bx["x"], bx["y"], bx["w"], bx["h"]
        rot = bx["rot"] % 360
        if rot in (90, 270):
            x, y, w, h = x + w / 2 - h / 2, y + h / 2 - w / 2, h, w
        st = []
        bg = css_fill(fill)
        lw = ctx.geo.len(line.width) if line.visible else 0
        if lw:
            x, y, w, h = x - lw / 2, y - lw / 2, w + lw, h + lw
        st.append(f"position:absolute;left:{f2(x)}px;top:{f2(y)}px;width:{f2(w)}px;height:{f2(h)}px")
        if bg:
            st.append(("background:" if bg.startswith("linear") else "background-color:") + bg)
        if lw:
            kind = "solid"
            if line.dash:
                kind = "dotted" if "dot" in line.dash.lower() and "dash" not in line.dash.lower() else "dashed"
            st.append(f"border:{f2(max(0.5, lw))}px {kind} {X.css_color(line.color)}")
        if name == "roundRect":
            r = min(bx["w"], bx["h"]) * adj.get("adj", 16667) / 100000
            st.append(f"border-radius:{f2(r + lw / 2)}px")
        sh = self.shadow(el)
        if sh:
            st.append(f"box-shadow:{sh}")
        return f'<div style="{";".join(st)}"></div>'

    def shadow(self, el) -> str | None:
        sppr = X.child(el, "spPr")
        sh = X.q(sppr, "a:effectLst/a:outerShdw")
        if sh is None:
            return None
        c = self.ctx.colors.find(sh)
        if c is None:
            return None
        dist = self.ctx.geo.len(X.ival(sh, "dist", 0))
        ang = math.radians(X.ival(sh, "dir", 0) / 60000)
        blur = self.ctx.geo.len(X.ival(sh, "blurRad", 0))
        return f"{f2(dist * math.cos(ang))}px {f2(dist * math.sin(ang))}px {f2(blur)}px {X.css_color(c)}"

    def path_svg(self, d, bx, fill, line, ellipse=False, arrows=False) -> str:
        ctx = self.ctx
        w, h = bx["w"], bx["h"]
        tf = []
        if bx["rot"]:
            tf.append(f"rotate({f2(bx['rot'])} {f2(w / 2)} {f2(h / 2)})")
        if bx["fh"]:
            tf.append(f"translate({f2(w)} 0) scale(-1 1)")
        if bx["fv"]:
            tf.append(f"translate(0 {f2(h)}) scale(1 -1)")
        g_open = f'<g transform="{" ".join(tf)}">' if tf else "<g>"
        paint = svg_paint(fill, ctx)
        stroke = svg_stroke(line, ctx)
        if ellipse:
            body = f'<ellipse cx="{f2(w / 2)}" cy="{f2(h / 2)}" rx="{f2(w / 2)}" ry="{f2(h / 2)}" {paint}{stroke}/>'
        else:
            body = f'<path d="{d}" {paint}{stroke} fill-rule="evenodd"/>'
        heads = ""
        if arrows and line.visible and (line.head or line.tail):
            subs = flatten(d)
            lw = max(0.5, ctx.geo.len(line.width))
            if subs and len(subs[0]) >= 2:
                s = subs[0]
                if line.head:
                    heads += arrow_head(line.head, s[0], s[1], lw, line.color)
                last = subs[-1]
                if line.tail and len(last) >= 2:
                    heads += arrow_head(line.tail, last[-1], last[-2], lw, line.color)
        return (f'<svg style="position:absolute;left:{f2(bx["x"])}px;top:{f2(bx["y"])}px;width:{f2(max(w, 1))}px;'
                f'height:{f2(max(h, 1))}px;overflow:visible" viewBox="0 0 {f2(max(w, 1))} {f2(max(h, 1))}">'
                f"{g_open}{body}{heads}</g></svg>")

    # 画像 ---
    def blip_blob(self, blip_fill, el):
        blip = X.q(blip_fill, "a:blip")
        if blip is None:
            return None, None, None
        rid = blip.get(R + "embed")
        part = self.part_for(el)
        if not rid:
            # PNG の控えを持たない SVG の画像（asvg:svgBlip だけ）
            sv = next((e for e in blip.iter() if X.local(e) == "svgBlip"), None)
            rid = sv.get(R + "embed") if sv is not None else None
        if not rid:
            if blip.get(R + "link"):
                self.ctx.rep("外部の画像（リンク）", "取りに行かない。空の枠にした")
            return None, None, blip
        try:
            ip = part.related_part(rid)
        except KeyError:
            return None, None, blip
        ext = pathlib.Path(str(ip.partname)).suffix.lower()
        return ip.blob, ext, blip

    def owner_of(self, el):
        root = el.getroottree().getroot()
        if root is self.ctx.lay:
            return "layout"
        if root is self.ctx.mst:
            return "master"
        return "slide"

    def part_for(self, el):
        """その要素の関係（画像などの rId）を持つ部品。SmartArt の図は図の部品。"""
        root = el.getroottree().getroot()
        if root in (self.ctx.lay, self.ctx.mst, self.ctx.sld):
            return self.ctx.part_of(self.owner_of(el))
        dsp = getattr(self, "dsp_part", None)
        return dsp if dsp is not None else self.ctx.part_of("slide")

    def blip_img(self, blip_fill, bx, mask=None, el=None, name=None) -> str | None:
        ctx = self.ctx
        blob, ext, blip = self.blip_blob(blip_fill, el)
        if blob is None:
            if blip is not None and not blip.get(R + "link"):
                ctx.rep("画像", "中身が無いので空の枠にした")
            return self.missing_img(bx)
        if ext in (".emf", ".wmf", ".emz", ".wmz"):
            conv = convert_metafile(blob, ext)
            if conv is None:
                ctx.rep(f"画像（{ext[1:].upper()}）", "写せないので空の枠にした（PNG にして貼り直すと写る）")
                return self.missing_img(bx)
            blob, ext = conv, ".png"
        src = X.q(blip_fill, "a:srcRect")
        crop = tuple(X.ival(src, k, 0) / 100000 for k in ("l", "t", "r", "b")) if src is not None else (0, 0, 0, 0)
        alpha = 1.0
        gray = False
        duo = None
        clr = None
        for e in blip:
            n = X.local(e)
            if n == "alphaModFix":
                alpha = X.ival(e, "amt", 100000) / 100000
            elif n == "grayscl":
                gray = True
            elif n == "duotone":
                cs = [ctx.colors.resolve(c) for c in e if isinstance(c.tag, str) and X.local(c) in X.COLOR_TAGS]
                if len(cs) == 2:
                    duo = (cs[0], cs[1])
            elif n == "clrChange":
                fc = ctx.colors.find(X.q(e, "a:clrFrom"))
                if fc is not None:
                    clr = fc[:3]
            elif n in ("lum", "biLevel", "alphaBiLevel", "tint", "hsl"):
                ctx.rep("画像の色の調整", "写さない（元の色のまま）")
        tile = X.q(blip_fill, "a:tile") is not None
        ops = dict(crop=crop, size=(bx["w"], bx["h"]), mask=mask, rot=bx["rot"], fh=bx["fh"], fv=bx["fv"], alpha=alpha,
                   gray=gray, duo=duo, clr_change=clr, tile=tile)
        if ext == ".svg":
            got = ctx.media.svg(ctx.no, blob, **ops)
        else:
            got = ctx.media.picture(ctx.no, blob, ext, **ops)
        if got is None:
            return self.missing_img(bx)
        rel, ow, oh = got
        cx, cy = bx["x"] + bx["w"] / 2, bx["y"] + bx["h"] / 2
        alt = ""
        if el is not None:
            for c in el:
                if isinstance(c.tag, str) and X.local(c).startswith("nv"):
                    cnv = X.child(c, "cNvPr")
                    if cnv is not None:
                        alt = cnv.get("descr") or ""
        return (f'<img src="{esc(rel)}" alt="{esc(alt[:200])}" style="position:absolute;left:{f2(cx - ow / 2)}px;'
                f'top:{f2(cy - oh / 2)}px;width:{f2(ow)}px;height:{f2(oh)}px;object-fit:cover" />')

    def missing_img(self, bx) -> str:
        return (f'<div style="position:absolute;left:{f2(bx["x"])}px;top:{f2(bx["y"])}px;width:{f2(bx["w"])}px;'
                f'height:{f2(bx["h"])}px;border:2px dashed #b0b0b0;background-color:rgba(176,176,176,0.12)"></div>')

    def picture(self, el, frame: Frame, owner: str):
        ctx = self.ctx
        sid, name, hidden = self.nv(el)
        if hidden:
            return
        sppr = X.child(el, "spPr")
        lay, mst, ph = self.inherit(el, owner)
        xf = xfrm_of(sppr) or xfrm_of(X.child(lay, "spPr")) or xfrm_of(X.child(mst, "spPr"))
        if xf is None:
            return
        bx = self.box(frame, xf)
        if self.outside(bx):
            return
        nvpr = X.q(el, "p:nvPicPr/p:nvPr")
        if nvpr is not None and any(X.local(e) in ("videoFile", "audioFile", "quickTimeFile", "media") for e in nvpr.iter()):
            ctx.rep("動画・音", "表紙の画像だけを置いた（再生はしない）")
        prst = X.q(sppr, "a:prstGeom")
        cust = X.q(sppr, "a:custGeom")
        mask = None
        if prst is not None and prst.get("prst") not in (None, "rect"):
            mask = GEO.preset(prst.get("prst"), bx["w"], bx["h"], GEO.adjusts(prst))
            if mask is None:
                ctx.rep(f"画像の形「{prst.get('prst')}」", "四角のまま置いた")
        elif cust is not None:
            got = GEO.custom(cust, bx["w"], bx["h"])
            mask = got[0] if got else None
        bf = X.child(el, "blipFill")
        pieces = [self.blip_img(bf, bx, mask=mask, el=el)]
        line = X.line_of(sppr, X.child(el, "style"), ctx.theme, ctx.colors)
        if line.visible:
            d = mask or GEO.preset("rect", bx["w"], bx["h"], {})
            pieces.append(self.path_svg(d, bx, None, line))
        self.b.emit(name, pieces, self.anim(sid, owner))

    # 表・グラフ・SmartArt・OLE ---
    def graphic(self, el, frame: Frame, owner: str):
        ctx = self.ctx
        sid, name, hidden = self.nv(el)
        if hidden:
            return
        xf = X.q(el, "p:xfrm")
        if xf is None:
            return
        off, ext = X.q(xf, "a:off"), X.q(xf, "a:ext")
        xd = {"x": X.ival(off, "x", 0), "y": X.ival(off, "y", 0), "w": X.ival(ext, "cx", 0), "h": X.ival(ext, "cy", 0),
              "rot": 0.0, "fh": False, "fv": False}
        gd = X.q(el, "a:graphic/a:graphicData")
        if gd is None or self.outside(self.box(frame, xd)):
            return
        uri = gd.get("uri") or ""
        if uri.endswith("/table"):
            tbl = X.q(gd, "a:tbl")
            if tbl is not None:
                import pptx_table
                self.b.emit(name, pptx_table.render(self, tbl, self.box(frame, xd)), self.anim(sid, owner))
            return
        if uri.endswith("/chart"):
            import pptx_chart
            ch = X.q(gd, "c:chart")
            rid = ch.get(R + "id") if ch is not None else None
            part = self.part_for(el)
            try:
                cpart = part.related_part(rid) if rid else None
            except KeyError:
                cpart = None
            pieces = pptx_chart.render(self, cpart, self.box(frame, xd)) if cpart is not None else []
            self.b.emit(name, pieces, self.anim(sid, owner))
            return
        if uri.endswith("/diagram"):
            self.diagram(el, gd, frame, xd, owner, self.anim(sid, owner), name)
            return
        # OLE など: 代わりの画像（p:pic）があれば置く
        pics = [p for p in gd.iter(P + "pic")]
        if pics:
            ctx.rep("埋め込みのオブジェクト（OLE）", "代わりの画像を置いた（中身は編集できない）")
            sub = Builder(ctx)
            sr = ShapeReader(ctx, sub)
            for p in pics[:1]:
                # 代わりの画像の座標はフレームに合わせる
                spPr = X.child(p, "spPr")
                if xfrm_of(spPr) is None:
                    sr.picture_at(p, self.box(frame, xd))
                else:
                    sr.picture(p, frame, owner)
            self.b.emit(name, sub.parts, self.anim(sid, owner))
            return
        ctx.rep("埋め込みのオブジェクト", "代わりの画像が無いので写さない")

    def picture_at(self, el, bx):
        bf = X.child(el, "blipFill")
        img = self.blip_img(bf, bx, el=el)
        if img:
            self.b.parts.append(img)

    def diagram(self, el, gd, frame: Frame, xd, owner, an, name):
        ctx = self.ctx
        part = self.part_for(el)
        drawing = None
        for rel in part.rels.values():
            if rel.reltype == X.RT_DIAGRAM_DRAWING and not rel.is_external:
                # 1 枚に SmartArt が複数あるときは、データの部品から図の部品をたどる
                drawing = drawing or rel
        rels = X.q(gd, "dgm:relIds")
        dm_rid = rels.get(R + "dm") if rels is not None else None
        target = None
        if dm_rid:
            try:
                dm = part.related_part(dm_rid)
                dmx = X.parse(dm.blob)
                ext = dmx.find(".//dsp:dataModelExt", {"dsp": X.NS["dsp"]})
                if ext is None:
                    ext = next((e for e in dmx.iter() if X.local(e) == "dataModelExt"), None)
                rid = ext.get("relId") if ext is not None else None
                if rid:
                    target = part.related_part(rid)
            except (KeyError, StopIteration):
                target = None
        if target is None and drawing is not None:
            target = drawing.target_part
        if target is None:
            ctx.rep("SmartArt", "描いた結果が入っていないので写さない（PowerPoint で開いて保存し直すと写る）")
            return
        root = X.parse(target.blob)
        tree = next((e for e in root.iter() if X.local(e) == "spTree"), None)
        if tree is None:
            return
        sub = frame.group((xd["x"], xd["y"]), (xd["w"], xd["h"]), (0, 0), (xd["w"], xd["h"]), 0, False, False)
        inner = Builder(ctx)
        sr = ShapeReader(ctx, inner)
        sr.dsp_part = target
        for sp in tree:
            if isinstance(sp.tag, str) and X.local(sp) == "sp":
                sr.dsp_shape(sp, sub)
        self.b.emit(name, inner.parts, an)

    def dsp_shape(self, sp, frame: Frame):
        """SmartArt の描いた結果の図形（dsp:sp）。文字の箱は dsp:txXfrm。"""
        ctx = self.ctx
        sppr = X.child(sp, "spPr")
        xf = xfrm_of(sppr)
        if xf is None:
            return
        bx = self.box(frame, xf)
        style = X.child(sp, "style")
        fill = X.first_fill(sppr, ctx.colors)
        if fill is None and style is not None:
            idx, refc = X.style_ref(style, "fillRef", ctx.colors)
            fill = X.ref_fill(ctx.theme, idx, refc, ctx.colors)
        line = X.line_of(sppr, style, ctx.theme, ctx.colors)
        pieces = self.geometry(sppr, bx, fill if (fill is None or fill.kind != "blip") else None, line, None, sp)
        tx = X.child(sp, "txBody")
        if tx is not None:
            txf = X.child(sp, "txXfrm")
            tb = bx
            if txf is not None:
                off, ext = X.child(txf, "off"), X.child(txf, "ext")
                if off is not None and ext is not None:
                    tb = self.box(frame, {"x": X.ival(off, "x", 0), "y": X.ival(off, "y", 0), "w": X.ival(ext, "cx", 0),
                                          "h": X.ival(ext, "cy", 0), "rot": X.ival(txf, "rot", 0) / 60000, "fh": False, "fv": False})
            pieces.append(self.text(sp, tx, tb, None, None, None, style, "dsp", None))
        self.b.parts.extend(p for p in pieces if p)

    # 文字 ---
    def text(self, el, tx, bx, lay, mst, ph, style, owner, geom_owner, para_hook=None) -> str:
        import pptx_text
        ctx = self.ctx
        lists = [X.child(tx, "lstStyle")]
        bodies = [X.child(tx, "bodyPr")]
        for inh in (lay, mst):
            if inh is not None:
                itx = X.child(inh, "txBody")
                lists.append(X.child(itx, "lstStyle"))
                bodies.append(X.child(itx, "bodyPr"))
        if ph is not None:
            lists.append(tx_style_of(ctx.mst, ph))
            lists.append(ctx.deck.default_text)
        else:
            lists.append(ctx.deck.default_text)
            if owner != "slide":
                lists.append(tx_style_of(ctx.mst, None))
        font_ref = X.child(style, "fontRef") if style is not None else None
        chain = X.TextChain(lists, ctx.theme, ctx.colors, font_ref)
        bp = X.body_props(bodies)
        prst = X.q(geom_owner, "a:prstGeom") if geom_owner is not None else None
        return pptx_text.render(ctx, tx, chain, bp, bx, prst.get("prst") if prst is not None else "rect", para_hook=para_hook,
                                adj=GEO.adjusts(prst) if prst is not None else None)


def convert_metafile(blob: bytes, ext: str) -> bytes | None:
    """EMF・WMF を PNG に（LibreOffice があるときだけ。無ければ None）。"""
    import shutil
    import subprocess
    import tempfile
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return None
    with tempfile.TemporaryDirectory(prefix="qs-meta-") as tmp:
        src = pathlib.Path(tmp) / f"m{ext}"
        src.write_bytes(blob)
        try:
            # 別の LibreOffice が開いていてもぶつからないよう、使い捨てのプロファイルで動かす
            subprocess.run([exe, f"-env:UserInstallation={pathlib.Path(tmp, 'profile').as_uri()}", "--headless",
                            "--convert-to", "png", "--outdir", tmp, str(src)],
                           capture_output=True, timeout=60, check=False)
        except (OSError, subprocess.SubprocessError):
            return None
        out = pathlib.Path(tmp) / "m.png"
        return out.read_bytes() if out.is_file() else None


# --- 背景 ------------------------------------------------------------------------------------

def background(ctx: SlideCtx) -> tuple[str | None, str | None]:
    """(section の背景の CSS, 一番奥に置く画像の HTML)。スライド → レイアウト → マスターの順に探す。"""
    for owner, root in (("slide", ctx.sld), ("layout", ctx.lay), ("master", ctx.mst)):
        bg = X.q(root, "p:cSld/p:bg")
        if bg is None:
            continue
        bgpr, ref = X.q(bg, "p:bgPr"), X.q(bg, "p:bgRef")
        fill = None
        if bgpr is not None:
            fill = X.first_fill(bgpr, ctx.colors)
        elif ref is not None:
            fill = X.ref_fill(ctx.theme, ref.get("idx"), ctx.colors.find(ref), ctx.colors)
        if fill is None:
            continue
        if fill.kind == "blip":
            g = ctx.geo
            # 16:9 を覆う大きさに広げる（余りに見えないように。中身は中央のまま）
            k = max(SLIDE_W / g.w, SLIDE_H / g.h)
            bw, bh = g.w * k, g.h * k
            bx = {"x": (SLIDE_W - bw) / 2, "y": (SLIDE_H - bh) / 2, "w": bw, "h": bh, "rot": 0, "fh": False, "fv": False}
            sr = ShapeReader(ctx, Builder(ctx))
            img = sr.blip_img(fill.blip, bx, el=bg)
            return "#ffffff", (img.replace("<img ", '<img data-name="背景" ', 1) if img else None)
        css = css_fill(fill)
        if fill.kind == "grad" and len(fill.stops) > 2:
            first, last = fill.stops[0], fill.stops[-1]
            css = f"linear-gradient({f2(fill.angle + 90)}deg, {X.css_color(first[1])} 0%, {X.css_color(last[1])} 100%)"
            ctx.rep("背景のグラデーション", "3 色以上なので、最初と最後の 2 色にした")
        if fill.kind == "grad" and fill.radial:
            ctx.rep("背景の放射状のグラデーション", "直線のグラデーションにした")
        return css or "#ffffff", None
    return X.css_color(ctx.colors.scheme_rgb("bg1") + (1.0,)), None


# --- 1 枚 -------------------------------------------------------------------------------------

def slide_title(sld) -> str:
    """ページの見出し（一覧に出す）。題のプレースホルダーの文字、無ければ一番大きな字の段落。"""
    def text_of(p):
        return "".join(t.text or "" for t in p.iter(A + "t")).strip()

    for ph, el in _tree_phs(sld):
        if ph["type"] in TITLE_TYPES:
            s = " ".join(x for x in (text_of(p) for p in el.iter(A + "p")) if x)
            if s:
                return s[:120]
    best, size = "", -1
    for p in sld.iter(A + "p"):
        t = text_of(p)
        if not t:
            continue
        szs = [X.ival(r, "sz", 1800) for r in p.iter(A + "rPr")] or [1800]
        if max(szs) > size:
            best, size = t, max(szs)
    return best[:120]


def notes_of(slide) -> str:
    if not slide.has_notes_slide:
        return ""
    ns = slide.notes_slide
    for ph in ns.placeholders:
        try:
            if ph.placeholder_format.type is not None and "BODY" in str(ph.placeholder_format.type):
                return ph.text_frame.text.strip()
        except Exception:
            continue
    tf = ns.notes_text_frame
    return tf.text.strip() if tf is not None else ""


def transition_of(sld) -> str:
    for el in sld.iter(P + "transition"):
        if any(isinstance(c.tag, str) and X.local(c) not in ("extLst", "sndAc") for c in el):
            return "fade"
        return "none"
    return "none"


def read_slide(deck: Deck, no: int, slide) -> dict:
    import pptx_anim
    layout = slide.slide_layout
    master = layout.slide_master
    ctx = SlideCtx(deck, no, slide, layout, master)
    b = Builder(ctx)
    anims = pptx_anim.read(ctx, ctx.sld)
    b.anim_of = anims.shapes
    b.texts_of = anims.paragraphs
    bg_css, bg_img = background(ctx)
    if bg_img:
        b.parts.append(bg_img)
    sr = ShapeReader(ctx, b)
    show_master = X.bval(ctx.sld, "showMasterSp", True)
    if show_master:
        if X.bval(ctx.lay, "showMasterSp", True):
            sr.walk(X.q(ctx.mst, "p:cSld/p:spTree"), Frame(), "master", skip_ph=True)
        sr.walk(X.q(ctx.lay, "p:cSld/p:spTree"), Frame(), "layout", skip_ph=True)
    sr.walk(X.q(ctx.sld, "p:cSld/p:spTree"), Frame(), "slide", skip_ph=False)
    spec = {"type": "html", "title": slide_title(ctx.sld), "background": bg_css or "#ffffff",
            "html": "\n".join(b.parts), "notes": notes_of(slide), "transition": transition_of(ctx.sld),
            "_import": {"slide": no, "layout": layout.name or ""}}
    for w in anims.notes:
        ctx.rep("アニメーション", w)
    if not X.bval(ctx.sld, "show", True):
        spec["_hidden"] = True
        ctx.rep("非表示のスライド", "取り込んだ（非表示の印を付けた）")
    return spec


def check_zip(path: pathlib.Path) -> None:
    try:
        zf = zipfile.ZipFile(path)
    except (zipfile.BadZipFile, OSError) as e:
        raise ImportError_(f"PPTX として読めない（zip でない）: {e}")
    with zf:
        infos = zf.infolist()
        if len(infos) > MAX_ENTRIES:
            raise ImportError_(f"部品の数が多すぎる（{len(infos)}）")
        total = 0
        names = set()
        for i in infos:
            names.add(i.filename)
            total += i.file_size
            if i.filename.startswith("/") or ".." in pathlib.PurePosixPath(i.filename).parts:
                raise ImportError_(f"部品の名前がおかしい: {i.filename[:80]}")
            if i.filename.endswith(".xml") or i.filename.endswith(".rels"):
                if i.file_size > MAX_XML_BYTES:
                    raise ImportError_(f"XML の部品が大きすぎる: {i.filename[:80]}")
            if i.compress_size and i.file_size / i.compress_size > MAX_RATIO and i.file_size > 10 * 1024 ** 2:
                raise ImportError_(f"展開すると大きくなりすぎる部品がある: {i.filename[:80]}")
        if total > MAX_TOTAL_BYTES:
            raise ImportError_("展開した大きさが大きすぎる")
        if "[Content_Types].xml" not in names or not any(n.startswith("ppt/presentation") for n in names):
            raise ImportError_("PowerPoint のファイル（.pptx）ではない")
        if any(n.startswith("ppt/vbaProject") for n in names):
            pass   # マクロは読まない（中身に触れない）


def theme_yaml(deck: Deck) -> tuple[str, dict]:
    """PPTX のテーマ（最初のマスター）→ html-slide のテーマの YAML。"""
    master = deck.prs.slide_masters[0]
    th = deck.theme_of(master)
    cm = X.q(master.part._element, "p:clrMap")
    colors = X.Colors(th.scheme, dict(cm.attrib) if cm is not None else None)
    bg = X.hex_of(colors.scheme_rgb("bg1"))
    ink = X.hex_of(colors.scheme_rgb("tx1"))
    primary = X.hex_of(colors.scheme_rgb("accent1"))
    accent = X.hex_of(colors.scheme_rgb("accent2"))
    font = th.fonts.get("mn-ea") or th.fonts.get("mn-lt") or ""
    d = {"background": bg, "ink": ink, "primary": primary, "accent": accent}
    lines = ["# 取り込んだ PPTX のテーマ（最初のスライドマスター）から写した色と字体。型のページを足すときに使う",
             *(f'{k}: "{v}"' for k, v in d.items())]
    if font and _font_known(font):
        d["font"] = font
        lines.append(f'font: "{font}"')
    elif font:
        # 型のページは字体のファイルで文字の幅を測って組む。この機械で字体のファイルが分からなければ、既定の字体で組む
        lines.append(f"# 元の字体: {font}（字体のファイルが分からないので、型のページは既定の字体で組む）")
        d["font_original"] = font
    return "\n".join(lines) + "\n", d


def _font_known(name: str) -> bool:
    try:
        import deck as D
        D.resolve_font({"font": name}, D.load_config())
        return True
    except Exception:
        return False


def import_pptx(src: pathlib.Path, out: pathlib.Path) -> dict:
    src = src.expanduser().resolve()
    if not src.is_file():
        raise ImportError_(f"ファイルが無い: {src}")
    if src.suffix.lower() not in (".pptx", ".pptm", ".ppsx", ".potx"):
        raise ImportError_(f"PowerPoint のファイル（.pptx）を渡す: {src.name}")
    check_zip(src)
    out = out.expanduser().resolve()
    out.mkdir(parents=True, exist_ok=True)
    try:
        deck = Deck(src, out)
    except ImportError_:
        raise
    except Exception as e:
        raise ImportError_(f"PPTX として読めない: {type(e).__name__}: {str(e)[:120]}")
    if src.suffix.lower() == ".pptm":
        deck.report.add(0, "マクロ", "読まない（中身に触れていない）")
    if not deck.geo.wide:
        deck.report.add(0, "スライドの大きさ", f"16:9 でない（{round(deck.geo.emu_w / 914400, 2)}×{round(deck.geo.emu_h / 914400, 2)} インチ）。"
                        "中央に置き、余りは背景で埋めた")
    slides = []
    for no, slide in enumerate(deck.prs.slides, 1):
        try:
            slides.append(read_slide(deck, no, slide))
        except Exception as e:
            deck.report.add(no, "スライド", f"読めないので空のページにした（{type(e).__name__}: {str(e)[:80]}）")
            slides.append({"type": "html", "title": "", "background": "#ffffff", "html": "", "notes": "",
                           "transition": "none", "_import": {"slide": no}})
    if not slides:
        raise ImportError_("スライドが 1 枚も無い")
    deck.media.finish()
    title = (deck.prs.core_properties.title or "").strip() or src.stem
    ty, tdict = theme_yaml(deck)
    (out / "theme.yaml").write_text(ty, encoding="utf-8")
    script = {"title": title[:120], "slides": slides}
    if _theme_ok(out / "theme.yaml"):
        script["theme"] = "theme.yaml"
    else:
        deck.report.add(0, "テーマの色", "文字と地の色の差が足りず、html-slide のテーマにできない（型のページを足すときは既定の色で組む）")
    (out / "script.json").write_text(json.dumps(script, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    rep = {"source": src.name, "slides": len(slides), "size": {"w_in": round(deck.geo.emu_w / 914400, 3),
                                                                 "h_in": round(deck.geo.emu_h / 914400, 3), "wide": deck.geo.wide},
           "theme": tdict, "items": deck.report.items}
    (out / "import-report.json").write_text(json.dumps(rep, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    return rep


def _theme_ok(path: pathlib.Path) -> bool:
    try:
        import deck as D
        th = D.load_theme(path)
        pal = D.palette(th)
        return not D.contrast_problems(pal)
    except Exception:
        return False


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="既存の PPTX → 取り込みの台本（html 型のページ）")
    ap.add_argument("--pptx", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args(argv)
    try:
        rep = import_pptx(pathlib.Path(a.pptx), pathlib.Path(a.out))
    except ImportError_ as e:
        print(f"NG: {e}", file=sys.stderr)
        return 2
    print(f"取り込み: {rep['slides']} 枚 → {pathlib.Path(a.out).expanduser() / 'script.json'}")
    for it in rep["items"]:
        where = f"{it['slide']}枚目" if it["slide"] else "全体"
        print(f"報告: {where}: {it['what']} — {it['how']}" + (f"（{it['count']} か所）" if it["count"] > 1 else ""), file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
