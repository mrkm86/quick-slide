"""部品 JSON（extract.mjs・from_video.py の出力）→ 編集できて動く PPTX（試作の build_pptx.py を一般化したもの）。

座標は 1920×1080 の px をそのまま 16:9 のスライドに写す（1px = 6350 EMU = 0.5pt）。
- 部品は PowerPoint の図形・文字・画像。同じ単位（unit）が重なり順で続く区間を1つのグループにする
- 動き: step 0 は画面に入ったら自動（最初の段を onBegin ＋ afterEffect）、step N は N 回目のクリック
  （同じ段は withEffect で、data-delay の小さい順）。段の番号は飛んでいても詰めて数える
- 出方（fx）: rise＝浮上 / fade＝フェード / pop＝ズーム（少し行き過ぎて戻る）/ wheel＝ホイール / wipe-*＝ワイプ
- data-keep は図形名 `!!名前`（モーフが画面をまたいで対応させる）
- data-keep="watermark"（透かし）は、選べない・動かせない・書き換えられない部品にする（a:spLocks。消すには［選択］ウィンドウから）
- 切り替え: morph（mc:AlternateContent で、モーフ未対応のアプリはフェード）/ fade / none
- 文字: 1行は折り返さず真ん中に置く。複数行は要素の幅で折り返し、行送りを pt の固定値にする（余白0・上揃え）。
  文節で折る（word-break: auto-phrase など）文字だけは、ブラウザが折り返した所に改行（a:br）を入れて固定する

write() の返り値（report）に、スライドごとの段・効果・ノートの［クリック］の数と使った fx を入れる（verify.py が照合に使う）。
"""
from __future__ import annotations

import base64
import copy
import io
import math
import re
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

PX = 6350
W, H = 1920, 1080
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
FXS = ("rise", "fade", "pop", "wheel", "wipe-right", "wipe-left", "wipe-up", "wipe-down")
CLICK = "［クリック］"
TEXT_PAD = 4            # 1行の文字の箱の左右の遊び（字幅のわずかな差で切れないように）
PIN_SLACK = 0.04        # 改行の位置を固定した文字の箱の幅の遊び（割合）
WRAP_SLACK = 0          # 複数行の箱の幅の遊び（px）。広げると Chrome より1字多く入る行が出る（447.3px に 32px×14字）ので足さない


def E(v):
    return Emu(int(round(v * PX)))


# --- 塗り・線・影 ---------------------------------------------------------------------

def _alpha(clr_parent, a):
    if clr_parent is not None and a is not None and a < 0.999:
        clr = clr_parent.find(qn("a:srgbClr"))
        etree.SubElement(clr, qn("a:alpha")).set("val", str(int(round(a * 100000))))


def _grad_xml(stops, css_angle):
    """CSS の linear-gradient（deg・0 が上向き）→ a:gradFill（ang は右向きが 0、時計回り）。"""
    gs = "".join(
        f'<a:gs pos="{int(round(s["pos"] * 100000))}"><a:srgbClr val="{s["color"]["hex"]}">'
        + (f'<a:alpha val="{int(round(s["color"]["a"] * 100000))}"/>' if s["color"].get("a", 1) < 0.999 else "")
        + "</a:srgbClr></a:gs>" for s in stops)
    ang = int(round(((css_angle - 90) % 360) * 60000))
    return f'<a:gradFill xmlns:a="{NS_A}" rotWithShape="1"><a:gsLst>{gs}</a:gsLst><a:lin ang="{ang}" scaled="0"/></a:gradFill>'


def style_fill(shape, fill, grad=None):
    if grad and grad.get("stops"):
        shape.fill.background()
        spPr = shape._element.spPr
        nf = spPr.find(qn("a:noFill"))
        g = etree.fromstring(_grad_xml(grad["stops"], grad.get("angle", 180)))
        if nf is not None:
            nf.addprevious(g)
            spPr.remove(nf)
        else:
            spPr.append(g)
    elif fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill["hex"])
        _alpha(shape.fill._xPr.find(qn("a:solidFill")), fill.get("a"))
    else:
        shape.fill.background()


def style_line(shape, stroke, sw, dash=None, cap=None, join=None):
    if stroke and sw and sw > 0:
        shape.line.color.rgb = RGBColor.from_string(stroke["hex"])
        shape.line.width = E(sw)
        ln = shape.line._get_or_add_ln()
        _alpha(ln.find(qn("a:solidFill")), stroke.get("a"))
        if dash:
            etree.SubElement(ln, qn("a:prstDash")).set("val", "sysDot" if dash == "dot" else "dash")
        if cap == "round":
            ln.set("cap", "rnd")
        elif cap == "square":
            ln.set("cap", "sq")
        if join == "round" or (join is None and cap == "round"):
            etree.SubElement(ln, qn("a:round"))
        elif join == "bevel":
            etree.SubElement(ln, qn("a:bevel"))
    else:
        shape.line.fill.background()


def add_shadow(shape, sh):
    """CSS の box-shadow（dx, dy, blur, color）→ a:outerShdw。True なら試作の既定（薄い下向きの影）。"""
    if sh is True:
        sh = {"dx": 0, "dy": 6, "blur": 16, "color": {"hex": "3c466e", "a": 0.08}}
    spPr = shape._element.spPr
    eff = spPr.find(qn("a:effectLst"))
    if eff is None:
        eff = etree.SubElement(spPr, qn("a:effectLst"))
    dist = math.hypot(sh["dx"], sh["dy"])
    d = int(round((math.degrees(math.atan2(sh["dy"], sh["dx"])) % 360) * 60000)) if dist else 5400000
    o = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(int(sh["blur"] * PX)), dist=str(int(dist * PX)),
                         dir=str(d), algn="ctr", rotWithShape="0")
    clr = etree.SubElement(o, qn("a:srgbClr"), val=sh["color"]["hex"].upper())
    etree.SubElement(clr, qn("a:alpha"), val=str(int(round(sh["color"].get("a", 1) * 100000))))


def _no_style(s):
    # 既定の図形スタイル（青い線・影・白い文字）を消す
    st = s._element.find(qn("p:style"))
    if st is not None:
        s._element.remove(st)


# --- 部品 -------------------------------------------------------------------------------

def add_rect(shapes, p, name):
    x, y, w, h = p["rect"]["x"], p["rect"]["y"], p["rect"]["w"], p["rect"]["h"]
    sw = p.get("sw") or 0
    if p.get("inset") and p.get("stroke") and sw:      # CSS の枠線は箱の内側。PowerPoint の線は辺の中心
        x, y, w, h = x + sw / 2, y + sw / 2, w - sw, h - sw
    if w <= 0 or h <= 0:
        return None
    rx = p.get("rx") or 0
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rx > 0.5 else MSO_SHAPE.RECTANGLE
    s = shapes.add_shape(kind, E(x), E(y), E(w), E(h))
    if rx > 0.5:
        s.adjustments[0] = min(0.5, max(0, rx - (sw / 2 if p.get("inset") else 0)) / min(w, h))
    style_fill(s, p.get("fill"), p.get("grad"))
    style_line(s, p.get("stroke"), sw, p.get("dash"), p.get("cap"), p.get("join"))
    if p.get("shadow"):
        add_shadow(s, p["shadow"])
    s.name = name
    _no_style(s)
    return s


def add_ellipse(shapes, p, name):
    r = p["rect"]
    if r["w"] <= 0 or r["h"] <= 0:
        return None
    s = shapes.add_shape(MSO_SHAPE.OVAL, E(r["x"]), E(r["y"]), E(r["w"]), E(r["h"]))
    style_fill(s, p.get("fill"))
    style_line(s, p.get("stroke"), p.get("sw"), p.get("dash"))
    s.name = name
    _no_style(s)
    return s


def add_path(shapes, p, name):
    subs = [sp for sp in p["subs"] if len(sp["pts"]) >= 2]
    if not subs:
        return None
    first = subs[0]["pts"][0]
    fb = shapes.build_freeform(E(first[0]), E(first[1]), scale=1.0)
    for i, sp in enumerate(subs):
        pts = sp["pts"]
        if i > 0:
            fb.move_to(E(pts[0][0]), E(pts[0][1]))
        fb.add_line_segments([(E(x), E(y)) for x, y in pts[1:]], close=sp["closed"])
    s = fb.convert_to_shape()
    fill = p.get("fill")
    stroke = p.get("stroke") if (p.get("sw") or 0) > 0 else None
    if fill and not all(sp["closed"] for sp in subs):
        # SVG の塗りは開いたサブパスも閉じて塗る。線は開いたまま。塗り用（全部閉じる・線なし）と線用（塗りなし）の2本に分ける
        path = s._element.find(".//" + qn("a:pathLst") + "/" + qn("a:path"))
        filled = copy.deepcopy(path)
        kids = list(filled)
        for i, k in enumerate(kids):
            nxt = kids[i + 1] if i + 1 < len(kids) else None
            if k.tag != qn("a:close") and (nxt is None or nxt.tag == qn("a:moveTo")):
                k.addnext(etree.Element(qn("a:close")))
        if stroke:
            filled.set("stroke", "0")
            path.set("fill", "none")
            path.addprevious(filled)
        else:
            path.addprevious(filled)
            path.getparent().remove(path)
    style_fill(s, fill)
    style_line(s, stroke, p.get("sw"), p.get("dash"), p.get("cap"), p.get("join"))
    s.name = name
    _no_style(s)
    return s


def _run(para, x):
    if x.get("color") and x["color"].get("a", 1) <= 0.002:
        return                       # 透明な文字は書かない（色を省くと PowerPoint の既定の色で見える）
    run = para.add_run()
    run.text = x["text"]
    f = run.font
    f.size = Pt(round(x["size"] * 0.5, 2))
    f.bold = bool(x.get("bold"))
    if x.get("italic"):
        f.italic = True
    if x.get("underline"):
        f.underline = True
    font = x.get("font") or "Yu Gothic"
    f.name = font
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        etree.SubElement(rPr, qn(tag)).set("typeface", font)
    rPr.set("lang", "ja-JP")
    if x.get("strike"):
        rPr.set("strike", "sngStrike")
    if x.get("spc"):
        rPr.set("spc", str(int(round(x["spc"] * 0.5 * 100))))
    if x.get("color"):
        f.color.rgb = RGBColor.from_string(x["color"]["hex"])
        _alpha(rPr.find(qn("a:solidFill")), x["color"].get("a"))


def add_text(shapes, p, name, line_breaks=False):
    r = p["rect"]
    runs = [x for x in p["runs"] if x.get("br") or x.get("text")]
    if not any(x.get("text", "").strip() for x in runs):
        return None
    multi = (p.get("lines") or 1) > 1
    pin = multi and (line_breaks or p.get("pin"))
    if pin:
        # 改行の位置を固定した文字は、字幅のわずかな差で余計に折れないよう、揃えの反対側に遊びを足す
        s = r["w"] * PIN_SLACK
        x = r["x"] - {"ctr": s / 2, "r": s}.get(p.get("align"), 0)
        tb = shapes.add_textbox(E(x), E(r["y"]), E(r["w"] + s), E(r["h"]))
    elif multi:
        tb = shapes.add_textbox(E(r["x"]), E(r["y"]), E(r["w"] + WRAP_SLACK), E(r["h"]))
    else:
        tb = shapes.add_textbox(E(r["x"] - TEXT_PAD), E(r["y"]), E(r["w"] + 2 * TEXT_PAD), E(r["h"]))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = multi
    tf.vertical_anchor = MSO_ANCHOR.TOP if multi else MSO_ANCHOR.MIDDLE
    bp = tf._txBody.find(qn("a:bodyPr"))
    for c in list(bp):
        bp.remove(c)
    etree.SubElement(bp, qn("a:noAutofit"))
    para = tf.paragraphs[0]
    para.alignment = {"ctr": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT, "just": PP_ALIGN.JUSTIFY}.get(p.get("align"), PP_ALIGN.LEFT) \
        if multi else PP_ALIGN.CENTER
    pPr = para._p.get_or_add_pPr()
    # ぶら下げ（行末の「。」を枠の外に出す）は Chrome がしないので切る。PowerPoint の既定はオン
    pPr.set("hangingPunct", "0")
    if multi and p.get("lh"):
        ln = etree.SubElement(pPr, qn("a:lnSpc"))
        etree.SubElement(ln, qn("a:spcPts")).set("val", str(int(round(p["lh"] * 0.5 * 100))))
    if pin:
        # 改行を固定するときは、折り返した所の空白（単語の切れ目）を行の終わりと頭から落とす
        runs = [dict(x) for x in runs]
        for i, x in enumerate(runs):
            if x.get("br"):
                continue
            prev = next((y for y in reversed(runs[:i]) if not y.get("br")), None)
            nxt = next((y for y in runs[i + 1:] if not y.get("br")), None)
            if nxt is not None and nxt.get("line") != x.get("line"):
                x["text"] = x["text"].rstrip(" ")
            if prev is not None and prev.get("line") != x.get("line"):
                x["text"] = x["text"].lstrip(" ")
        runs = [x for x in runs if x.get("br") or x["text"]]
    line, forced = None, False
    for x in runs:
        if x.get("br"):
            para.add_line_break()
            forced = True
            continue
        if pin and line is not None and x.get("line", line) != line and not forced:
            para.add_line_break()
        line, forced = x.get("line", line), False
        _run(para, x)
    if p.get("rot"):
        tb.rotation = p["rot"]
    tb.name = name
    return tb


def _image_stream(src: str):
    """パス・data URI → (stream, 拡張子)。webp は PNG にする（python-pptx が読めない）。"""
    if src.startswith("data:"):
        head, _, data = src.partition(",")
        raw = base64.b64decode(data) if ";base64" in head else data.encode()
        ext = re.search(r"image/(\w+)", head).group(1)
    else:
        raw = Path(src).read_bytes()
        ext = Path(src).suffix.lower().lstrip(".")
    if ext in ("webp",):
        from PIL import Image
        buf = io.BytesIO()
        Image.open(io.BytesIO(raw)).save(buf, "PNG")
        raw = buf.getvalue()
    return io.BytesIO(raw)


def add_img(shapes, p, name):
    r = p["rect"]
    if r["w"] <= 0 or r["h"] <= 0:
        return None
    pic = shapes.add_picture(_image_stream(p["src"]), E(r["x"]), E(r["y"]), E(r["w"]), E(r["h"]))
    c = p.get("crop")
    if c:
        pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = c["l"], c["t"], c["r"], c["b"]
    if p.get("radius"):
        pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
        prst = pic._element.spPr.find(qn("a:prstGeom"))
        av = prst.find(qn("a:avLst"))
        if av is None:
            av = etree.SubElement(prst, qn("a:avLst"))
        etree.SubElement(av, qn("a:gd"), name="adj", fmla=f"val {int(min(0.5, p['radius'] / min(r['w'], r['h'])) * 100000)}")
    if p.get("shadow"):
        add_shadow(pic, p["shadow"])
    if p.get("op") is not None and p["op"] < 0.999:
        blip = pic._element.find(".//" + qn("a:blip"))
        etree.SubElement(blip, qn("a:alphaModFix")).set("amt", str(int(round(p["op"] * 100000))))
    if p.get("alt"):
        pic._element.nvPicPr.cNvPr.set("descr", p["alt"])     # 代替テキスト（読み上げ・PowerPoint の「代替テキスト」欄）
    pic.name = name
    return pic


def add_prim(shapes, p, name, line_breaks=False):
    t = p["type"]
    if t == "rect":
        return add_rect(shapes, p, name)
    if t == "ellipse":
        return add_ellipse(shapes, p, name)
    if t == "path":
        return add_path(shapes, p, name)
    if t == "text":
        return add_text(shapes, p, name, line_breaks)
    if t == "img":
        return add_img(shapes, p, name)
    raise ValueError(f"部品の種類が無い: {t}")


# --- アニメーション ----------------------------------------------------------------------

# ワイプ: fx は「描く向き」（wipe-right = 左から右へ現れる = PowerPoint の「左から」）。
# presetSubtype は PowerPoint の「方向」、filter は現れ始める辺（LibreOffice の読み込み表と同じ解釈・2026-10-03 に照合）
WIPE = {"right": (8, "left"), "left": (2, "right"), "down": (1, "up"), "up": (4, "down")}


class Timing:
    """<p:timing> を組む。groups = [("auto"|"click", [(delay_ms, spid, fx), ...]), ...]（表示順）。"""

    def __init__(self):
        self.n = 2

    def _id(self):
        self.n += 1
        return str(self.n)

    @staticmethod
    def _set_visible(spid, cid):
        return (f'<p:set><p:cBhvr><p:cTn id="{cid}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
                f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
                f'<p:to><p:strVal val="visible"/></p:to></p:set>')

    def _anim(self, spid, attr, vals, dur):
        tav = "".join(f'<p:tav tm="{tm}"><p:val><p:strVal val="{v}"/></p:val></p:tav>' for tm, v in vals)
        return (f'<p:anim calcmode="lin" valueType="num"><p:cBhvr additive="base"><p:cTn id="{self._id()}" dur="{dur}" fill="hold"/>'
                f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>{attr}</p:attrName></p:attrNameLst></p:cBhvr>'
                f'<p:tavLst>{tav}</p:tavLst></p:anim>')

    def _filter(self, spid, flt, dur):
        return (f'<p:animEffect transition="in" filter="{flt}"><p:cBhvr><p:cTn id="{self._id()}" dur="{dur}"/>'
                f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>')

    def effect(self, spid, fx, delay, node_type):
        cid = self._id()
        body = self._set_visible(spid, self._id())
        if fx == "rise":
            pid, sub = 42, 0
            body += self._filter(spid, "fade", 500) + self._anim(spid, "ppt_x", [(0, "#ppt_x"), (100000, "#ppt_x")], 500) \
                + self._anim(spid, "ppt_y", [(0, "#ppt_y+0.022"), (100000, "#ppt_y")], 500)
        elif fx == "pop":             # GSAP の back.out(2) に寄せて、少し行き過ぎてから戻る
            pid, sub = 53, 16
            body += self._filter(spid, "fade", 300) \
                + self._anim(spid, "ppt_w", [(0, "#ppt_w*0.82"), (65000, "#ppt_w*1.04"), (100000, "#ppt_w")], 400) \
                + self._anim(spid, "ppt_h", [(0, "#ppt_h*0.82"), (65000, "#ppt_h*1.04"), (100000, "#ppt_h")], 400)
        elif fx == "wheel":           # 手書きの丸: 時計回りに描く（ホイール・1スポーク）
            pid, sub = 21, 1
            body += self._filter(spid, "wheel(1)", 700)
        elif fx.startswith("wipe-"):
            pid = 22
            sub, edge = WIPE[fx[5:]]
            body += self._filter(spid, f"wipe({edge})", 600 if fx[5:] in ("right", "left") else 500)
        else:                         # fade（知らない fx もフェード）
            pid, sub = 10, 0
            body += self._filter(spid, "fade", 500)
        return (f'<p:par><p:cTn id="{cid}" presetID="{pid}" presetClass="entr" presetSubtype="{sub}" fill="hold" '
                f'nodeType="{node_type}"><p:stCondLst><p:cond delay="{int(delay)}"/></p:stCondLst>'
                f'<p:childTnLst>{body}</p:childTnLst></p:cTn></p:par>')

    def build(self, groups):
        outer = []
        for mode, effs in groups:
            if not effs:
                continue
            cond = ('<p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond>'
                    if mode == "auto" else '<p:cond delay="indefinite"/>')
            oid, iid = self._id(), self._id()
            body = "".join(
                self.effect(spid, fx, delay, ("afterEffect" if mode == "auto" else "clickEffect") if i == 0 else "withEffect")
                for i, (delay, spid, fx) in enumerate(effs))
            outer.append(f'<p:par><p:cTn id="{oid}" fill="hold"><p:stCondLst>{cond}</p:stCondLst><p:childTnLst>'
                         f'<p:par><p:cTn id="{iid}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
                         f'<p:childTnLst>{body}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
        if not outer:
            return None
        xml = (f'<p:timing xmlns:p="{NS_P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
               f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
               + "".join(outer) +
               '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
               '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
               '</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')
        return etree.fromstring(xml)


MORPH = (f'<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:p="{NS_P}">'
         '<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" '
         'xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" Requires="p159">'
         '<p:transition spd="slow" p14:dur="800"><p159:morph option="byObject"/></p:transition></mc:Choice>'
         '<mc:Fallback><p:transition spd="med"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>')
FADE = f'<p:transition xmlns:p="{NS_P}" spd="med"><p:fade/></p:transition>'


def set_background(slide, bg):
    if not bg:
        return
    if bg["type"] == "grad":
        fill = _grad_xml(bg["stops"], bg.get("angle", 180))
    else:
        fill = f'<a:solidFill xmlns:a="{NS_A}"><a:srgbClr val="{bg["color"]["hex"]}"/></a:solidFill>'
    el = etree.fromstring(f'<p:bg xmlns:p="{NS_P}" xmlns:a="{NS_A}"><p:bgPr>{fill}<a:effectLst/></p:bgPr></p:bg>')
    slide._element.find(qn("p:cSld")).insert(0, el)


# --- 組み立て ---------------------------------------------------------------------------

def _segments(prims):
    """重なり順を保ったまま、同じ単位が続く区間に分ける。"""
    segs = []
    for p in prims:
        u = p.get("unit")
        if segs and segs[-1][0] == u:
            segs[-1][1].append(p)
        else:
            segs.append((u, [p]))
    return segs


LOCKED_KEEPS = {"watermark"}


def lock_shape(sh) -> None:
    """図形を、クリックで選べず・動かせず・大きさも文字も変えられないようにする（透かし。PowerPoint の［ロック］より強い noSelect つき）。"""
    el = sh._element
    nv = el.find(qn("p:nvSpPr"))
    if nv is None:                                   # 文字の枠（p:sp）でなければ何もしない
        return
    c = nv.find(qn("p:cNvSpPr"))
    for old in c.findall(qn("a:spLocks")):
        c.remove(old)
    locks = etree.SubElement(c, qn("a:spLocks"))
    for k in ("noGrp", "noSelect", "noRot", "noMove", "noResize", "noTextEdit"):
        locks.set(k, "1")
    c.insert(0, locks)


def build_slide(prs, s, si, line_breaks=False):
    """1枚を足して、そのスライドの report を返す。"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, s.get("background"))
    units = s.get("units") or {}
    sid = s.get("id") or f"s{si + 1}"
    spids = {}
    for k, (u, ps) in enumerate(_segments(s.get("prims") or [])):
        if u is None:
            for j, p in enumerate(ps):
                add_prim(slide.shapes, p, p.get("name") or f"{sid}/static-{k}-{j}", line_breaks)
            continue
        info = units.get(u) or {}
        name = f"!!{info['keep']}" if info.get("keep") else (info.get("name") or u)
        if len(ps) == 1:
            sh = add_prim(slide.shapes, ps[0], name, line_breaks)
        else:
            grp = slide.shapes.add_group_shape()
            made = [add_prim(grp.shapes, p, f"{name}/{j}", line_breaks) for j, p in enumerate(ps)]
            if not any(made):
                grp._element.getparent().remove(grp._element)
                sh = None
            else:
                grp.name = name
                sh = grp
        if sh is not None:
            spids.setdefault(u, []).append(sh.shape_id)
            if info.get("keep") in LOCKED_KEEPS:
                lock_shape(sh)

    # 段 → 効果
    by_step = {}
    for u, ids in spids.items():
        info = units.get(u) or {}
        st = info.get("step")
        if st is None:
            continue
        fx = info.get("fx") or "rise"
        for spid in ids:
            by_step.setdefault(int(st), []).append((int(info.get("delay") or 0), spid, fx))
    groups = []
    for st in sorted(by_step):
        effs = sorted(by_step[st], key=lambda e: e[0])          # 同じ遅れなら重なり順（sorted は安定）
        groups.append(("auto" if st == 0 else "click", effs))
    tim = Timing().build(groups)
    tr = s.get("transition") or ("fade" if si == 0 else "morph")
    el = slide._element
    if tr == "morph":
        el.append(etree.fromstring(MORPH))
    elif tr == "fade":
        el.append(etree.fromstring(FADE))
    if tim is not None:
        el.append(tim)
    notes = s.get("notes") or ""
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    fx_count = {}
    for _, effs in groups:
        for _, _, fx in effs:
            fx_count[fx] = fx_count.get(fx, 0) + 1
    steps = sum(1 for m, _ in groups if m == "click")
    said = notes.count(CLICK)
    return {"no": si + 1, "id": sid, "transition": tr, "steps": steps,
            "effects": sum(len(e) for _, e in groups), "auto": sum(len(e) for m, e in groups if m == "auto"),
            "notes_clicks": said, "fx": fx_count, "shapes": len(slide.shapes), "ok": steps == said}


def write(parts: dict, out_path, *, line_breaks: bool = False) -> dict:
    """部品 JSON → PPTX を out_path に保存し、report を返す。

    line_breaks=True なら、複数行の文字すべてで、ブラウザが折り返した所に改行（a:br）を入れる（PowerPoint でも同じ所で切れる）。
    既定は、文節で折る・行をならす指定（部品の pin）の文字だけに入れる。ほかは PowerPoint が箱の幅で折り返すので、
    直したときに自然に流れ直す。

    report = {"renderer": "html-slide", "pptx", "source", "mode", "slides": [{no, id, transition, steps, effects, auto,
              notes_clicks, fx: {fx: 数}, shapes, ok}], "totals": {slides, steps, effects, notes_clicks}, "mismatch": [no…]}
    """
    out_path = Path(out_path)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W * PX), Emu(H * PX)
    slides = [build_slide(prs, s, si, line_breaks) for si, s in enumerate(parts.get("slides") or [])]
    prs.core_properties.title = parts.get("title") or ""
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return {
        "renderer": "html-slide", "pptx": str(out_path.resolve()), "source": parts.get("source"), "mode": parts.get("mode"),
        "slides": slides,
        "totals": {"slides": len(slides), "steps": sum(s["steps"] for s in slides),
                   "effects": sum(s["effects"] for s in slides), "notes_clicks": sum(s["notes_clicks"] for s in slides)},
        "mismatch": [s["no"] for s in slides if not s["ok"]],
    }


if __name__ == "__main__":
    import json
    import sys
    if len(sys.argv) != 3:
        sys.exit("使い方: python3 pptx_writer.py <parts.json> <out.pptx>")
    rep = write(json.loads(Path(sys.argv[1]).read_text()), sys.argv[2])
    for s in rep["slides"]:
        print(f"slide{s['no']:>2} {s['id']:<8} steps={s['steps']} effects={s['effects']:>2} notes={s['notes_clicks']} "
              f"{s['transition']:<5} {'OK' if s['ok'] else 'NG'}")
    print(json.dumps(rep["totals"], ensure_ascii=False))
