#!/usr/bin/env python3
"""html-slide の入口。

    catalog     画面の型と置き方の一覧画像（どれを使うかを見て選ぶ）
    images      台本の画像の枠の一覧（何を描くか・縦横比・置き場・出どころ）
    fill        prompt だけの枠（生成の枠）のうち、まだ絵が無いものを Codex で描いて埋める（deck も既定でこれを先にやる）
    revise      生成の枠の絵1枚を指示で直す（微修正＝今の絵を土台に・作り直し＝渡さない）。直したら deck を組み直す
    check       台本（--script）か HTML（--html）を検査する。PPTX は作らない
    deck        台本 → HTML → PPTX・PDF・ページの画像・サムネイル（--format で選ぶ）→ 照合 → inspect.md
    assemble    ページごとの部品（PPTX の部品 JSON・PDF の何ページ目か）を並べて、PPTX・PDF を 1 本ずつ書く（ADR 0024）
    render      決まりに沿った HTML → PPTX
    import      既存の PPTX → 取り込みの台本（html 型のページ）・画像・テーマ・報告（ADR 0025）。deck に渡せば組める
    from-video  html-video が焼いた解説動画（<出力>/<preset>）→ PPTX
    verify      PPTX のクリックとノートの照合（--odp で LibreOffice の読み取りも）

終了コード: 0=成功 / 1=実行時エラー（描けなかった枠がある・deck はその枠を仮にして組む）/
          2=台本・HTML の誤り（決まり違反・はみ出し・照合のずれ）/ 4=依存不足・Codex の認証
"""
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.dont_write_bytecode = True   # 置いた場所（配る物の中・ほかの環境からのリンク先）に __pycache__ を残さない
sys.path.insert(0, str(HERE))

import browser as B  # noqa: E402
import cache as K  # noqa: E402
import codex_run as X  # noqa: E402
import deck  # noqa: E402
import imagegen as G  # noqa: E402
import outputs as OUT  # noqa: E402
import pdfpages as PDF  # noqa: E402
import places as PL  # noqa: E402
import render as R  # noqa: E402
import verify as V  # noqa: E402

SNAPSHOT = HERE / "snapshot.mjs"


def out_root() -> Path:
    return PL.out_root("html-slide", _cfg())


def say_violations(vs):
    for v in vs:
        print(R.format_violation(v), file=sys.stderr)


# --- 検品用の写し -------------------------------------------------------------------

def _texts(shapes):
    """図形を重なり順にたどり、文字を持つものの文字を返す（グループの中も）。"""
    out = []
    for sh in shapes:
        if sh.shape_type == 6:  # group
            out += _texts(sh.shapes)
        elif getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
            out.append(sh.text_frame.text.strip().replace("\n", " / ").replace("\v", " / "))
    return out


def write_inspect(pptx: Path, report: dict, vres: dict, info: dict | None, script: Path | None) -> Path:
    """INSPECTION.md の検品役が読む写し。スライドごとに 型・画面の文字・ノート・段と効果の数・機械の検査。"""
    from pptx import Presentation
    prs = Presentation(str(pptx))
    types = {s["no"]: s["type"] for s in (info or {}).get("slides", [])}
    styled = {s["no"]: s.get("styled") or {} for s in (info or {}).get("slides", [])}
    rs = {s["no"]: s for s in report.get("slides", [])}
    vs = {s["no"]: s for s in vres.get("slides", [])}
    L = [f"# 検品用の写し — {info.get('title') if info else pptx.stem}", ""]
    if script:
        L.append(f"- 台本: `{script}`")
    L += [f"- PPTX: `{pptx}`", f"- モード: {(info or {}).get('mode', '—')}",
          f"- スタイル: {(info or {}).get('style', '—')}", f"- 枚数: {len(prs.slides)}", ""]
    errs = R.errors(report.get("violations") or [])
    warns = [v for v in (report.get("violations") or []) if v not in errs]
    L += ["## 機械の検査", "",
          f"- check（HTML の決まり・はみ出し）: {'違反なし' if not errs else f'違反 {len(errs)} 件'}"
          + (f"（警告 {len(warns)} 件）" if warns else ""),
          f"- verify（クリックの段 = ノートの［クリック］）: {'全スライド一致' if vres.get('ok') else 'ずれあり'}"]
    o = vres.get("odp") or {}
    if o.get("skipped"):
        L.append(f"- LibreOffice の読み取り: 飛ばした（{o['skipped']}）")
    elif o:
        L.append(f"- LibreOffice の読み取り（効果の数）: {'全スライド一致' if not o.get('diff') else 'ずれ ' + str(o['diff'])}")
    for v in errs + warns:
        L.append(f"  - {R.format_violation(v)}")
    for w in (info or {}).get("warnings") or []:
        L.append(f"- 警告: {w}")
    L.append("")
    for i, slide in enumerate(prs.slides, 1):
        r, v = rs.get(i, {}), vs.get(i, {})
        L.append(f"## {i}枚目 — {types.get(i, '—')}（段 {r.get('steps', v.get('clicks', '—'))}・効果 {r.get('effects', v.get('effects', '—'))}）")
        if styled.get(i):
            L.append("")
            L.append("スタイルで足した既定: " + "・".join(f"{k}={v}" for k, v in styled[i].items()))
        L.append("")
        L.append("画面の文字（重なり順）:")
        for t in _texts(slide.shapes):
            L.append(f"- {t}")
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        L += ["", "ノート:", "", f"> {notes or '（なし）'}", ""]
    p = pptx.parent / "inspect.md"
    p.write_text("\n".join(L) + "\n", encoding="utf-8")
    return p


# --- 画像の枠を Codex で埋める（ADR 0011） ------------------------------------------------

DEFAULT_PARALLEL = 8      # image-slide と同じ
DEFAULT_TIMEOUT = 600     # 1枚の上限（秒）


def log(msg: str) -> None:
    print(msg, file=sys.stderr, flush=True)


def _cfg() -> dict:
    return deck.load_config()


def _sidecar(path: Path, slot: dict, extra: dict) -> None:
    """生成の絵の横に、何から描いたかを書く（gen-<鍵>.json。呼ぶ側の DB が拾える）。"""
    meta = {"key": slot["key"], "slide": slot["slide"], "slot": slot["slot"], "image_look": slot["image_look"],
            "prompt": slot["prompt"], "text": slot["gen_text"], "shape": slot["shape"],
            "characters": slot.get("characters") or [], "at": time.strftime("%Y-%m-%dT%H:%M:%S"), **extra}
    path.with_suffix(".json").write_text(json.dumps(meta, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")


def draw_slots(jobs: list[dict], parallel: int | None, retries: int) -> tuple[int, list[dict]]:
    """jobs（{slot, text, refs, dest, extra}）を並列に描く。(rc, 描けなかった job)。認証はここで1度だけ（rc=4）。"""
    cfg = _cfg()
    try:
        cli = X.find_cli(cfg)
        auth = X.check_auth(cli, log=log)
    except X.CodexError as e:
        log(f"NG: {e.detail}")
        return 4, jobs
    # Runner は「N枚目」と数えるが、ここの N は枠の通し番号（下の一覧の番号）。読み違えないよう「N.」に直して出す
    runner = X.Runner(parallel or int(cfg.get("max_parallel") or DEFAULT_PARALLEL), 1 + max(0, retries),
                      log=lambda m: log(re.sub(r"(\d+)枚目", r"\1.", m)))
    timeout = int(cfg.get("timeout") or DEFAULT_TIMEOUT)
    for i, j in enumerate(jobs, 1):
        j["no"] = i                      # Runner は no で数える（1枚の中に枠が複数あるので、通し番号にする）
        log(f"  {i}. {j['slot']['slide']}枚目の {j['slot']['slot']}（{j['slot']['image_look']}）: {j['slot']['prompt'][:40]}")

    def work(j, attempt, last):
        dest = Path(j["dest"])
        cand = dest.with_name(dest.stem + ".try.png")
        got = X.generate(cli, auth["snapshot"], j["text"], j["refs"], cand, dest.parent / "logs",
                         timeout=timeout, should_stop=runner.stop.is_set, shape=j["slot"]["shape"])
        os.replace(cand, dest)
        _sidecar(dest, j["slot"], {"seconds": got["seconds"], "attempts": attempt, **j.get("extra", {})})
        return got

    results = runner.run(jobs, work)
    K.trim("html-slide", deck.load_config(), Path(jobs[0]["dest"]).parent, log=log)
    failed = [j for j, _, err in results if err is not None]
    if runner.stop_reason is not None and runner.stop_reason.kind == "auth":
        log(f"NG: {runner.stop_reason.detail}")
        return 4, failed
    for j, _, err in results:
        if err is not None:
            log(f"警告: {j['slot']['slide']}枚目の {j['slot']['slot']} を描けなかった（{err.kind}: {str(err.detail)[:200]}）")
    return (1 if failed else 0), failed


def fill(script: Path, a) -> tuple[int, int]:
    """まだ絵が無い生成の枠を描く。(rc, 描いた数)。持ち込みの絵が無ければ描く前に rc=2。"""
    try:
        slots = deck.image_slots(script, theme=a.theme, style=a.style, image_look=a.image_look)
    except deck.DeckError as e:
        log(f"NG: {e}")
        return 2, 0
    lost = [x for x in slots if not x["exists"] and x["origin"] != "generated"]
    if lost:
        for x in lost:
            log(f"NG: {x['slide']}枚目の {x['slot']}: 持ち込みの絵が無い: {x['path']}")
        return 2, 0
    todo, seen = [], set()
    for x in slots:
        if x["origin"] == "generated" and not x["exists"] and x["key"] not in seen:
            seen.add(x["key"])           # 同じ指示の枠は1枚だけ描いて使い回す
            todo.append({"slot": x, "text": x["gen_text"], "refs": [Path(c["path"]) for c in x["characters"]],
                         "dest": x["gen_path"]})
    gen = sum(1 for x in slots if x["origin"] == "generated")
    log(f"画像の枠 {len(slots)}（生成 {gen}・描く {len(todo)}）")
    if not todo:
        return 0, 0
    K.area("html-slide", deck.load_config(), script.parent)    # 描いた絵の置き場（キャッシュの区画・ADR 0014）
    rc, failed = draw_slots(todo, getattr(a, "parallel", None), getattr(a, "retries", 1))
    return rc, len(todo) - len(failed)


def cmd_fill(a) -> int:
    rc, n = fill(Path(a.script).expanduser().resolve(), a)
    if rc == 0:
        print(f"描いた {n} 枚")
    return rc


def cmd_revise(a) -> int:
    """生成の枠の絵1枚を直す（image-slide の revise と同じ考え方・ADR 0011）。直したら deck を組み直す。

    鍵は台本から作った元のまま（同じ置き場に上書きする）。前の絵は区画の history/ に写す。台本を変えずに deck を流し直しても、
    直した絵が使われる。prompt・画風・キャラを変えると鍵が変わり、描き直す（直しは消える）。
    """
    script = Path(a.script).expanduser().resolve()
    if lacking_browser():              # 直したあと deck を組み直すので、描く前に確かめる
        return 4
    try:
        slots = deck.image_slots(script, theme=a.theme, style=a.style, image_look=a.image_look)
    except deck.DeckError as e:
        log(f"NG: {e}")
        return 2
    mine = [x for x in slots if x["slide"] == a.page]
    if not mine:
        log(f"NG: {a.page}枚目に画像の枠が無い")
        return 2
    if a.slot:
        mine = [x for x in mine if x["slot"] == a.slot]
    elif len(mine) > 1:
        log(f"NG: {a.page}枚目には枠が {len(mine)} つある。--slot で選ぶ（{', '.join(x['slot'] for x in mine)}）")
        return 2
    if not mine:
        log(f"NG: {a.page}枚目に枠 {a.slot} が無い")
        return 2
    x = mine[0]
    if x["origin"] != "generated":
        log(f"NG: {a.page}枚目の {x['slot']} は持ち込みの絵（{x['path']}）。直すのは Codex で描いた絵だけ。"
            "差し替えるなら台本の path を書き換える")
        return 2
    how = "remake" if a.remake else "minor"
    attach = [Path(p).expanduser().resolve() for p in (a.attach or [])]
    for p in attach:
        if not p.is_file() or p.suffix.lower() not in deck.T.IMAGE_EXTS:
            log(f"NG: 参考画像（--attach）が無いか、画像ではない: {p}")
            return 2
    note = (a.note or "").strip()
    if not note and not attach:
        log("NG: 直す中身が無い。--note（文）か --attach（参考画像）のどちらかを渡す")
        return 2
    dest = Path(x["gen_path"])
    if how == "minor" and not dest.is_file():
        log(f"NG: {a.page}枚目の {x['slot']} はまだ描いていないので、微修正の土台が無い。先に deck か fill を流すか、--remake で描く")
        return 2
    refs = [Path(c["path"]) for c in x["characters"]]
    n = len(refs) + (1 if how == "minor" else 0) + len(attach)
    if n > G.MAX_REFS:
        log(f"NG: 参照画像が {n} 枚（キャラ {len(refs)}・土台 {1 if how == 'minor' else 0}・--attach {len(attach)}）。"
            f"1枚に添えられるのは {G.MAX_REFS} 枚まで（Codex の画像ツールの上限）")
        return 2
    K.area("html-slide", deck.load_config(), script.parent)
    hist = None
    if dest.is_file():
        (dest.parent / "history").mkdir(exist_ok=True)
        hist = dest.parent / "history" / f"{dest.stem}-{time.strftime('%Y%m%d-%H%M%S')}.png"
        shutil.copy2(dest, hist)
        if dest.with_suffix(".json").is_file():
            shutil.copy2(dest.with_suffix(".json"), hist.with_suffix(".json"))
    base_no = None
    if how == "minor":
        refs.append(hist)
        base_no = len(refs)
    span = (len(refs) + 1, len(refs) + len(attach)) if attach else None
    refs += attach
    text = x["gen_text"] + "\n\n" + G.revision_part(how, note, base_no, span)
    what = {"minor": "微修正（今の絵を土台に渡す）", "remake": "作り直し（今の絵は渡さない）"}[how]
    log(f"{a.page}枚目の {x['slot']} を{what}" + (f"。前の絵は {hist}" if hist else ""))
    rec = {"revision": {"mode": how, "note": note, "attach": [str(p) for p in attach],
                        "previous": str(hist) if hist else None, "at": time.strftime("%Y-%m-%dT%H:%M:%S")}}
    rc, _ = draw_slots([{"slot": x, "text": text, "refs": refs, "dest": str(dest), "extra": rec}], 1, a.retries)
    if rc:
        return rc
    return cmd_deck(a)


# --- 入口 --------------------------------------------------------------------------

def lacking_browser() -> int:
    """Chrome で HTML を開く道具（Node・puppeteer-core・Chrome）が足りなければ、描く前に rc=4 で止める（ADR 0016）。
    先に Codex で絵を描いてから足りないと分かると、サブスクの枠を無駄に使う。"""
    lack = B.missing()
    if not lack:
        return 0
    print(f"NG: {'・'.join(lack)}。{B.SETUP_HINT}", file=sys.stderr)
    return 4


def cmd_check(a) -> int:
    if lacking_browser():
        return 4
    if a.script:
        with tempfile.TemporaryDirectory(prefix="html-slide-check-") as tmp:
            try:
                html, _ = deck.build_html(a.script, tmp, theme=a.theme, style=a.style, image_look=a.image_look,
                                          placeholder="generated")      # Codex で描く枠は、描く前でも検査を通す
            except deck.DeckError as e:
                print(f"NG: {e}", file=sys.stderr)
                return 2
            vs = R.check(html)
    else:
        vs = R.check(a.html)
    say_violations(vs)
    if R.errors(vs):
        return 2
    print("OK")
    return 0


def _node(args: list[str], what: str) -> bool:
    node, _ = B.node()
    if not node:
        print(f"NG: {what}には node が要る（見つからない）。{B.SETUP_HINT}", file=sys.stderr)
        return False
    r = subprocess.run([node, str(SNAPSHOT), *args], capture_output=True, text=True, timeout=600, env=B.child_env())
    if r.returncode != 0:
        print(f"NG: {what}を作れなかった: {r.stderr.strip()[:300]}", file=sys.stderr)
        return False
    return True


def snapshot(html: Path, out: Path, count: int) -> bool:
    """ページの画像（out/slide-NN.png・1920×1080）を撮る。前の回のページ数を超える古い画像は消す。"""
    out.mkdir(parents=True, exist_ok=True)
    OUT.prune_pages(out, count)
    return _node([str(html), str(out)], "ページの画像")


PDF_LONG = 2560     # PDF に入れる写真の長い辺の上限（画面では 1920px の枠に収まる）


def print_html(html: Path, tmp: Path) -> Path:
    """PDF 用の HTML の写し。不透明な PNG は JPEG にして渡す（写真を PNG のまま入れると 10枚で 30MB になる）。
    透過のある絵（ロゴ・切り抜き）は PNG のまま。写しは別のフォルダに置くので、絵はすべて絶対の file: URL にする"""
    from html import unescape
    from PIL import Image
    seen: dict[Path, str] = {}

    def swap(m):
        src = unescape(m.group(2))
        if re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*:", src):      # data:・file: などはそのまま
            return m.group(0)
        p = (html.parent / src).resolve()
        if p not in seen:
            seen[p] = p.as_uri()
            if p.suffix.lower() == ".png" and p.is_file():
                with Image.open(p) as im:
                    clear = im.mode in ("RGBA", "LA", "PA") or "transparency" in im.info
                    if not clear or im.convert("RGBA").getchannel("A").getextrema()[0] == 255:
                        rgb = im.convert("RGB")
                        rgb.thumbnail((PDF_LONG, PDF_LONG), Image.LANCZOS)
                        q = tmp / f"img-{len(seen):03d}.jpg"
                        rgb.save(q, quality=88)
                        seen[p] = q.as_uri()
        return m.group(1) + seen[p] + m.group(3)

    text = re.sub(r'(<img\b[^>]*?\bsrc=")([^"]+)(")', swap, html.read_text(encoding="utf-8"))
    out = tmp / "deck.print.html"
    out.write_text(text, encoding="utf-8")
    return out


def write_pdf(html: Path, pdf: Path) -> bool:
    """deck.html を Chrome で印刷して PDF にする（文字は選べる・字体は埋め込み・1ページ 13.333×7.5 インチで PPTX と同じ）。"""
    with tempfile.TemporaryDirectory(prefix="html-slide-pdf-") as tmp:
        part = pdf.with_suffix(".tmp.pdf")
        ok = _node([str(print_html(html, Path(tmp))), "-", "--pdf", str(part)], "PDF ")
        if ok:
            os.replace(part, pdf)
        else:
            part.unlink(missing_ok=True)
        return ok


def finish(pptx: Path, report: dict, info, script, odp: bool) -> int:
    vres = V.verify(pptx, odp=odp and bool(shutil.which("soffice")), report=report.get("report_path"))
    insp = write_inspect(pptx, report, vres, info, script)
    t = report.get("totals", {})
    print(f"slides={len(report.get('slides', []))} clicks={t.get('steps', '—')} effects={t.get('effects', '—')} "
          f"verify={'OK' if vres['ok'] else 'NG'}")
    print(f"inspect: {insp}")
    print(pptx)
    return 0 if vres["ok"] else 2


def cmd_deck(a) -> int:
    script = Path(a.script).expanduser().resolve()
    out = Path(a.out).expanduser() if a.out else out_root() / script.stem
    try:
        fmts = OUT.formats_of(a.format)
    except ValueError as e:
        print(f"NG: {e}", file=sys.stderr)
        return 2
    if a.no_png:            # 前の口（--format で png・thumbs を外すのと同じ）
        fmts = tuple(f for f in fmts if f not in ("png", "thumbs"))
    if lacking_browser():   # PPTX も PDF も画像も Chrome で組む。絵を描く前に確かめる
        return 4
    out.mkdir(parents=True, exist_ok=True)
    made = {out.resolve() / n for n in ("deck.html", "deck.pptx", "deck.pdf", "deck-info.json", "build.json", "inspect.md")}
    if script in made:
        print(f"NG: 台本 {script} が出力のファイル名と重なる（--out を別のフォルダにする）", file=sys.stderr)
        return 2
    fill_rc = 0
    if not (a.no_fill or a.placeholder):          # 既定で、まだ絵の無い生成の枠を Codex で描いて埋める（ADR 0011）
        fill_rc, _ = fill(script, a)
        if fill_rc in (2, 4):
            return fill_rc
    # 描けなかった枠（fill_rc=1）と、--no-fill の生成の枠は仮の枠にする。--placeholder は持ち込みの無い絵も仮にする
    ph = True if a.placeholder else ("generated" if (a.no_fill or fill_rc) else False)
    try:
        html, info = deck.build_html(script, out, theme=a.theme, mode=a.mode, placeholder=ph, style=a.style,
                                     logo=a.logo, logo_position=a.logo_position, watermark=a.watermark,
                                     image_look=a.image_look)
    except deck.DeckError as e:
        print(f"NG: {e}", file=sys.stderr)
        return 2
    for w in info.get("warnings") or []:
        print(f"警告: {w}", file=sys.stderr)
    if info.get("font_note"):
        print(f"注: {info['font_note']}", file=sys.stderr)
    report = None
    if "pptx" in fmts:
        try:
            report = R.render(html, out / "deck.pptx", report_path=out / "build.json",
                              parts_path=Path(a.parts).expanduser() if a.parts else None)
        except R.RenderError as e:
            (out / "deck-info.json").write_text(json.dumps(info, ensure_ascii=False, indent=1) + "\n")
            say_violations(e.violations)
            print(f"NG: {e}", file=sys.stderr)
            return e.rc
        say_violations(report.get("violations") or [])
    # ページの画像・PDF・サムネイル（ADR 0012）。作れなかったものがあれば rc=1
    made, out_rc = {}, 0
    slides = json.loads(script.read_text())["slides"]
    if "png" in fmts or "thumbs" in fmts:
        if snapshot(html, out / "slides", len(slides)):
            pngs = [out / "slides" / f"slide-{i:02d}.png" for i in range(1, len(slides) + 1)]
            made["png"] = [str(p) for p in pngs]
            if "thumbs" in fmts:
                title = info.get("title") or script.stem
                made["thumbs"] = [str(p) for p in OUT.write_thumbs(pngs, [OUT.label_of(s) for s in slides], out, title)]
        else:
            out_rc = 1
    if "pdf" in fmts:
        if write_pdf(html, out / "deck.pdf"):
            made["pdf"] = str(out / "deck.pdf")
        else:
            out_rc = 1
    if "pptx" in fmts:
        made["pptx"] = str(out / "deck.pptx")
    info["outputs"] = {"format": list(fmts), **made}
    (out / "deck-info.json").write_text(json.dumps(info, ensure_ascii=False, indent=1) + "\n")
    if "pdf" in made:
        print(f"pdf: {made['pdf']}")
    if "png" in made:
        print(f"png: {out / 'slides'}（{len(made['png'])}枚）")
    for t in made.get("thumbs", []):
        print(f"thumbs: {t}")
    rc = finish(out / "deck.pptx", report, info, script, odp=not a.no_odp) if report is not None else 0
    if fill_rc:
        log("NG: 描けなかった画像の枠がある（仮の枠で組んだ）。もう一度 deck を流すと、その枠だけ描き直す")
    return rc or out_rc or fill_rc


def cmd_assemble(a) -> int:
    """ページごとの部品を並べて PPTX と PDF を書く（版の違うページを混ぜた書き出し。ADR 0024）。

    `--manifest` の JSON: {"title", "size", "slides": [部品 JSON の slides の 1 枚（画像の src は絶対パスか data URI）],
    "pdf": [[PDF のパス, 何ページ目か（1 から）], …]}。slides があれば `<out>/deck.pptx`、pdf があれば `<out>/deck.pdf`。
    PDF が読めない形（Chrome の印刷でない）なら PDF は書かずに rc=1（PPTX は書く）。
    """
    try:
        man = json.loads(Path(a.manifest).expanduser().read_text(encoding="utf-8"))
        slides, picks = man.get("slides") or [], man.get("pdf") or []
        if not isinstance(slides, list) or not isinstance(picks, list):
            raise ValueError("slides と pdf は配列")
        picks = [(str(p), int(n)) for p, n in picks]
    except (OSError, ValueError, TypeError) as e:
        print(f"NG: 組み立ての指示が読めない: {e}", file=sys.stderr)
        return 2
    out = Path(a.out).expanduser()
    out.mkdir(parents=True, exist_ok=True)
    rc = 0
    if slides:
        import pptx_writer
        parts = {"version": 1, "mode": "contract", "title": man.get("title") or "", "size": man.get("size"),
                 "slides": slides}
        tmp = out / "deck.pptx.tmp"
        rep = pptx_writer.write(parts, tmp)
        os.replace(tmp, out / "deck.pptx")
        print(f"pptx: {out / 'deck.pptx'}（{len(rep['slides'])}枚）")
    if picks:
        try:
            n = PDF.assemble(picks, out / "deck.pdf")
            print(f"pdf: {out / 'deck.pdf'}（{n}ページ）")
        except (OSError, PDF.PdfError) as e:
            print(f"NG: PDF を組み立てられなかった: {e}", file=sys.stderr)
            rc = 1
    return rc


def cmd_render(a) -> int:
    html = Path(a.html).expanduser().resolve()
    out = Path(a.out).expanduser() if a.out else html.parent / (html.stem + ".pptx")
    try:
        report = R.render(html, out)
    except R.RenderError as e:
        say_violations(e.violations)
        print(f"NG: {e}", file=sys.stderr)
        return e.rc
    say_violations(report.get("violations") or [])
    return finish(Path(out).resolve(), report, None, None, odp=not a.no_odp)


def cmd_import(a) -> int:
    import pptx_import as PI
    out = Path(a.out).expanduser() if a.out else out_root() / Path(a.pptx).stem
    try:
        rep = PI.import_pptx(Path(a.pptx), out)
    except PI.ImportError_ as e:
        print(f"NG: {e}", file=sys.stderr)
        return 2
    for it in rep["items"]:
        where = f"{it['slide']}枚目" if it["slide"] else "全体"
        print(f"報告: {where}: {it['what']} — {it['how']}" + (f"（{it['count']} か所）" if it["count"] > 1 else ""), file=sys.stderr)
    print(f"取り込み: {rep['slides']} 枚 → {out / 'script.json'}（報告は {out / 'import-report.json'}）")
    return 0


def cmd_from_video(a) -> int:
    vdir = Path(a.video_dir).expanduser().resolve()
    project = vdir / "project"
    script = Path(a.script).expanduser().resolve() if a.script else vdir.parent / "script.json"
    build = vdir / "build.json"
    for p, what in ((project / "index.html", "project/index.html"), (script, "台本（script.json）"), (build, "build.json")):
        if not p.exists():
            print(f"NG: {what} が無い: {p}", file=sys.stderr)
            return 2
    out = Path(a.out).expanduser() if a.out else out_root() / f"video-{vdir.parent.name}"
    out.mkdir(parents=True, exist_ok=True)
    try:
        report = R.render(project, out / "deck.pptx", mode="hyperframes", script=script, build=build,
                          report_path=out / "build.json")
    except R.RenderError as e:
        say_violations(e.violations)
        print(f"NG: {e}", file=sys.stderr)
        return e.rc
    info = {"title": json.loads(script.read_text()).get("title", vdir.parent.name), "mode": "talk",
            "slides": [{"no": i, "type": sc.get("layout")} for i, sc in enumerate(json.loads(script.read_text())["scenes"], 1)]}
    return finish(out / "deck.pptx", report, info, script, odp=not a.no_odp)


def cmd_images(a) -> int:
    script = Path(a.script).expanduser().resolve()
    if a.json and Path(a.json).expanduser().resolve() == script:
        print(f"NG: --json の書き出し先が台本と同じ: {script}", file=sys.stderr)
        return 2
    try:
        slots = deck.image_slots(script, theme=a.theme, style=a.style, image_look=a.image_look)
    except deck.DeckError as e:
        print(f"NG: {e}", file=sys.stderr)
        return 2
    if a.missing:
        slots = [x for x in slots if not x["exists"]]
    if a.json:
        Path(a.json).expanduser().write_text(json.dumps(slots, ensure_ascii=False, indent=1) + "\n")
    for x in slots:
        mark = "あり" if x["exists"] else "なし"
        origin = {"path": "持ち込み", "placed": "置いた絵", "generated": "生成"}[x["origin"]]
        print(f"{x['slide']:>2}枚目 {x['type']:<12} {x['slot']:<10} {origin} {mark}  {x['size']}（枠 {x['aspect']}・{x['crop']}）  {x['path']}")
        if not x["exists"] and x.get("prompt_full"):
            print(f"      描くもの: {x.get('gen_text') or x['prompt_full']}")
    print(f"枠 {len(slots)} 件（無いもの {sum(1 for x in slots if not x['exists'])} 件）")
    return 0


def cmd_catalog(a) -> int:
    import catalog
    out = Path(a.out).expanduser() if a.out else out_root() / "catalog"
    try:
        sheets = catalog.build_catalog(out, only=a.only, theme=a.theme)
    except deck.DeckError as e:
        print(f"NG: {e}", file=sys.stderr)
        return 2
    except RuntimeError as e:
        print(f"NG: {e}", file=sys.stderr)
        return 1
    for p in sheets:
        print(p)
    return 0


def cmd_verify(a) -> int:
    return V.main([a.pptx] + (["--odp"] if a.odp else []) + (["--report", a.report] if a.report else []))


LOOK_HELP = ("Codex で描く画像の枠の画風: photo（写実の写真風）／illustration（平らなイラスト）／handdrawn（手描きのスケッチ）。"
             "省けばスタイルが決める（standard は illustration・visual は photo）。変えると、生成の枠だけ描き直す（持ち込みの絵は触らない）")


def deck_args(d) -> None:
    """deck と revise に共通の引数（revise は直したあと deck を組み直す）。"""
    d.add_argument("--script", required=True)
    d.add_argument("--out", help="出力フォルダ（省略時は config の out:/<台本の名前>）")
    d.add_argument("--theme", help="default / forest / warm / テーマの YAML")
    d.add_argument("--style", help="standard（既定・説明・報告・社内資料向け）／visual（ビジュアル寄り・プレゼン・ピッチ向け）／カスタムスタイルの YAML（ADR 0005）")
    d.add_argument("--mode", choices=("talk", "handout"),
                   help="talk（既定・話す資料。クリックで出る）／handout（配る資料。動きも切り替えも付けない）。台本の mode より優先")
    d.add_argument("--logo", help="ロゴの画像（PNG・JPEG）。高さ 44px・幅 329px の枠へ縦横比を保って収める（引き伸ばさない）")
    d.add_argument("--logo-position", choices=list(deck.LOGO_POSITIONS), default="bottom-right",
                   help="ロゴを置く角（既定 bottom-right）。ふさがったページは、右下 → 左下・左下 → 右下・右上 → 右下へ逃がし、それもだめなら置かない")
    d.add_argument("--watermark", help="透かしの文字（例: example.com・社外秘）。全ページの中央下の端に、固定の薄いグレーで置く。"
                                       "PPTX では一番上の、選べない・動かせない部品になる")
    d.add_argument("--image-look", choices=list(G.IMAGE_LOOKS), help=LOOK_HELP)
    d.add_argument("--no-fill", action="store_true", help="Codex で描かない。まだ絵の無い生成の枠は仮の枠にする（急ぐ下書き・Codex が使えないとき）")
    d.add_argument("--placeholder", action="store_true", help="描かず、持ち込みの無い絵も含めて仮の枠にする（下書き用。検品では落ちる）")
    d.add_argument("--parallel", type=int, help=f"同時に描く本数（既定 {DEFAULT_PARALLEL}）")
    d.add_argument("--retries", type=int, default=1, help="描けなかったときのやり直しの回数（既定 1）")
    d.add_argument("--format", help="作る出力を pptx,pdf,png,thumbs からカンマで選ぶ（既定は4つとも）。"
                                    "png はページの画像（slides/slide-NN.png）、thumbs は全ページを並べた一覧（thumbs.png）。例: --format png")
    d.add_argument("--no-png", action="store_true", help=argparse.SUPPRESS)   # 前の口（--format から png・thumbs を外す）
    d.add_argument("--no-odp", action="store_true", help="LibreOffice での照合を飛ばす")
    d.add_argument("--parts", help="PPTX の部品 JSON（extract.mjs の出力。形は references/parts-format.md）を書き出す先。"
                                   "ページごとに残して、版の違うページを混ぜた書き出しを assemble で組むのに使う（ADR 0024）")


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="html-slide — 編集できて動く PPTX を作る")
    sub = ap.add_subparsers(dest="cmd", required=True)
    style_help = "standard（既定・説明・報告・社内資料向け）／visual（ビジュアル寄り・プレゼン・ピッチ向け）／カスタムスタイルの YAML（ADR 0005）"
    ca = sub.add_parser("catalog", help="画面の型と置き方の一覧画像")
    ca.add_argument("--out", help="出力フォルダ（省略時は config の out:/catalog）")
    ca.add_argument("--only", help="image（画像を使うものだけ）/ text / 型の名前")
    ca.add_argument("--theme")
    im = sub.add_parser("images", help="台本の画像の枠の一覧（出どころ: 持ち込み・置いた絵・生成）")
    im.add_argument("--script", required=True)
    im.add_argument("--theme")
    im.add_argument("--style", help=style_help)
    im.add_argument("--image-look", choices=list(G.IMAGE_LOOKS), help=LOOK_HELP)
    im.add_argument("--missing", action="store_true", help="無い枠だけ")
    im.add_argument("--json", help="一覧を JSON で書き出す先")
    c = sub.add_parser("check", help="台本か HTML を検査する（Codex で描く枠は、まだ絵が無くても通す）")
    g = c.add_mutually_exclusive_group(required=True)
    g.add_argument("--script")
    g.add_argument("--html")
    c.add_argument("--theme")
    c.add_argument("--style", help=style_help)
    c.add_argument("--image-look", choices=list(G.IMAGE_LOOKS), help=LOOK_HELP)
    fl = sub.add_parser("fill", help="まだ絵の無い生成の枠を Codex で描いて埋める（PPTX は作らない）")
    fl.add_argument("--script", required=True)
    fl.add_argument("--theme")
    fl.add_argument("--style", help=style_help)
    fl.add_argument("--image-look", choices=list(G.IMAGE_LOOKS), help=LOOK_HELP)
    fl.add_argument("--parallel", type=int, help=f"同時に描く本数（既定 {DEFAULT_PARALLEL}）")
    fl.add_argument("--retries", type=int, default=1, help="描けなかったときのやり直しの回数（既定 1）")
    d = sub.add_parser("deck", help="台本 → PPTX（まだ絵の無い生成の枠は、先に Codex で描いて埋める）")
    deck_args(d)
    rv = sub.add_parser("revise", help="生成の枠の絵1枚を指示で直し、deck を組み直す")
    deck_args(rv)
    rv.add_argument("--page", type=int, required=True, help="何枚目か")
    rv.add_argument("--slot", help="枠の名前（image・member1 など。1枚に枠が複数あるときだけ）")
    rv.add_argument("--remake", action="store_true",
                    help="作り直し（今の絵を渡さずに描き直す。テイストを変えるとき）。付けなければ微修正（今の絵を土台に、指示だけ直す）")
    rv.add_argument("--note", help="直す指示（文）")
    rv.add_argument("--attach", action="append", help="参考画像（何枚でも。キャラ・土台と合わせて5枚まで）")
    asm = sub.add_parser("assemble", help="ページごとの部品を並べて PPTX・PDF を書く（版の違うページを混ぜた書き出し）")
    asm.add_argument("--manifest", required=True, help="組み立ての指示（JSON。形は cmd_assemble の説明）")
    asm.add_argument("--out", required=True, help="deck.pptx・deck.pdf を書くフォルダ")
    r = sub.add_parser("render", help="決まりに沿った HTML → PPTX")
    r.add_argument("--html", required=True)
    r.add_argument("--out")
    r.add_argument("--no-odp", action="store_true")
    ip = sub.add_parser("import", help="既存の PPTX → 取り込みの台本（deck --script <out>/script.json で組む）")
    ip.add_argument("--pptx", required=True, help="取り込む PPTX")
    ip.add_argument("--out", help="台本・画像・報告を書くフォルダ（省略時は config の out:/<PPTX の名前>）")
    f = sub.add_parser("from-video", help="html-video の出力 → PPTX")
    f.add_argument("--video-dir", required=True, help="html-video の <出力>/<preset>（project/ と build.json がある所）")
    f.add_argument("--script", help="台本（省略時は <video-dir>/../script.json）")
    f.add_argument("--out")
    f.add_argument("--no-odp", action="store_true")
    v = sub.add_parser("verify", help="PPTX の照合")
    v.add_argument("pptx")
    v.add_argument("--odp", action="store_true")
    v.add_argument("--report")
    a = ap.parse_args(argv)
    try:
        return {"images": cmd_images, "catalog": cmd_catalog, "check": cmd_check, "deck": cmd_deck, "render": cmd_render,
                "from-video": cmd_from_video, "verify": cmd_verify, "fill": cmd_fill, "revise": cmd_revise,
                "assemble": cmd_assemble, "import": cmd_import}[a.cmd](a)
    except R.RenderError as e:
        print(f"NG: {e}", file=sys.stderr)
        return e.rc


if __name__ == "__main__":
    sys.exit(main())
