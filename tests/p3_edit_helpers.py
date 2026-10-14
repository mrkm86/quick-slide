"""P3（直す）: 10 ページの 1 冊のうち 7・8 枚目だけを直し、ほかのページは台本の版でも 1 文字も変わらず、
描き直されもしないこと。画面のページと描いた回の PPTX の並び・枚数が一致すること。

- 台本係は偽物（`tests/fakes/writer_common.py`。指示の語で振り分けを決める）、**エンジンは本物**、絵を描く codex は
  エンジンの偽物（呼ばれた回とページを記録する）
- ピッチ（image-slide）は描いたページを codex の呼び出しの記録から数える（描く指示に「このページは N枚目」がある）
- プレゼン（html-slide）は Node と Chrome が無ければ skip。生成の枠の絵（区画の `gen-<鍵>.png`）が替わったかで数える
"""

from __future__ import annotations

import json
import time
import unittest
from pathlib import Path

from helpers import FAKES, fake_tools, intake_payload, runnable, start_server, wait_until_idle, writer_calls

from quick_slide.scriptdoc import dumps, without_pins
from quick_slide.settings import Settings, WorkerSettings

from test_engines_e2e import html_slide_ready

CONFIG = Settings(worker=WorkerSettings(parallel=2))


def pptx_notes(path: Path) -> list[str]:
    from pptx import Presentation

    prs = Presentation(str(path))
    out = []
    for slide in prs.slides:
        out.append(slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else "")
    return out


class EditBase(unittest.TestCase):
    mode = "pitch"
    runtime = "codex"

    def setUp(self) -> None:
        varied = runnable(FAKES / "fake_codex_varied.py")   # 描くたびに違う絵（「絵だけ直した」を中身で見分ける）
        self.tools, self.log = fake_tools(self, extra={"FAKE_WRITER_PAGES": "10", "FAKE_CODEX_SLEEP": "0.05",
                                                       "IMAGE_SLIDE_CODEX": varied, "HTML_SLIDE_CODEX": varied})
        self.fx = start_server(self, tools=self.tools, settings=CONFIG)

    # --- 道具 ---
    def make_deck(self, **attrs) -> str:
        r = self.fx.post("/api/intakes", intake_payload(runtime=self.runtime, mode=self.mode, **attrs))
        self.assertEqual(r.status, 201, r.body)
        deck = r.json()["deck_id"]
        run = wait_until_idle(self.fx, deck, seconds=300)
        self.assertEqual(run["state"], "done", run["message"])
        return deck

    def view(self, deck: str) -> dict:
        return self.fx.get(f"/api/decks/{deck}").json()

    def wait(self, deck: str) -> dict:
        run = wait_until_idle(self.fx, deck, seconds=300)
        return run

    def head(self, deck: str) -> tuple[dict, bytes]:
        with self.fx.server.store_lock:
            row = self.fx.store.head_script(deck)
            path = self.fx.store.script_path(row)
        return json.loads(path.read_text(encoding="utf-8")), path.read_bytes()

    def codex_calls(self) -> set[str]:
        d = self.log / "codex-calls"
        return {p.name for p in d.glob("*.json")} if d.is_dir() else set()

    def drawn_pages(self, before: set[str]) -> list[int]:
        out = []
        for name in sorted(self.codex_calls() - before):
            rec = json.loads((self.log / "codex-calls" / name).read_text(encoding="utf-8"))
            if rec.get("cmd") == "exec" and "t1" in rec:
                out.append(rec["page"])
        return sorted(out)

    def instruct(self, deck: str, positions: list[int], text: str, **extra) -> dict:
        pages = self.view(deck)["pages"]
        ids = [pages[n - 1]["slide_id"] for n in positions]
        r = self.fx.post(f"/api/decks/{deck}/instructions", {"scope": "pages", "slide_ids": ids, "text": text, **extra})
        self.assertEqual(r.status, 200, r.body)
        return r.json()

    def assert_screen_matches_pptx(self, deck: str) -> None:
        view = self.view(deck)
        visible = [p for p in view["pages"] if not p["hidden"]]
        with self.fx.server.store_lock:
            run = self.fx.store.shown_run(deck)
        notes = pptx_notes(self.fx.root.run_dir(deck, run["id"]) / "deck.pptx")
        data, _ = self.head(deck)
        talk = {s["id"]: s["talk"][0] for s in data["slides"]}
        self.assertEqual(len(notes), len(visible), "PPTX の枚数は画面の（非表示を除く）ページ数")
        self.assertEqual([n.replace("［クリック］", "").strip() for n in notes], [talk[p["key"]] for p in visible],
                         "PPTX の並びは画面の並び")
        for p in visible:
            self.assertEqual(p["display"]["key"], "ready", p)
        self.assertFalse(view["live"]["working"])

    def duplicate_and_rebuild(self, deck: str) -> str:
        """複製して、同じ姿か（ページ・絵・書き出し）確かめ、複製を組み直す。複製の id を返す。"""
        src_view = self.view(deck)
        r = self.fx.post(f"/api/decks/{deck}/duplicate", {"title": "写し"})
        self.assertEqual(r.status, 201, r.body)
        new = r.json()["deck"]["id"]
        view = self.view(new)
        self.assertEqual([(p["key"], p["title"], p["hidden"]) for p in view["pages"]],
                         [(p["key"], p["title"], p["hidden"]) for p in src_view["pages"]], "複製はいまの姿のページを持つ")
        for p in view["pages"]:
            if not p["hidden"]:
                self.assertEqual(p["display"]["key"], "ready", p)
                self.assertEqual(len(p["versions"]), 1, "版の履歴は写さない")
        with self.fx.server.store_lock:
            a, b = self.fx.store.shown_run(deck), self.fx.store.shown_run(new)
        self.assertEqual((self.fx.root.run_dir(new, b["id"]) / "deck.pptx").read_bytes(),
                         (self.fx.root.run_dir(deck, a["id"]) / "deck.pptx").read_bytes())
        self.assert_screen_matches_pptx(new)
        calls = self.codex_calls()
        r = self.fx.post(f"/api/decks/{new}/rebuild", {})
        self.assertEqual(r.status, 200, r.body)
        run = self.wait(new)
        self.assertEqual(run["state"], "done", run["message"])
        self.assertEqual(self.drawn_pages(calls), [], "複製を組み直しても絵は描き直さない（控えから引き継ぐ）")
        self.assert_screen_matches_pptx(new)
        self.assertEqual([p["key"] for p in self.view(deck)["pages"]], [p["key"] for p in src_view["pages"]], "元は変わらない")
        return new

    def page_bytes(self, data_bytes: bytes, data: dict) -> dict[str, str]:
        """台本の版のファイルの中の、ページごとの文字（同じ書き方で書いたときの断片）。"""
        text = data_bytes.decode("utf-8")
        out = {}
        for s in data["slides"]:
            piece = json.dumps(s, ensure_ascii=False, indent=1).replace("\n", "\n  ")
            self.assertIn(piece, text)
            out[s["id"]] = piece
        return out


class PresentationEditBase(EditBase):
    mode = "presentation"
    runtime = "claude_code"

    @classmethod
    def setUpClass(cls) -> None:
        why = html_slide_ready()
        if why:
            raise unittest.SkipTest(why)


    def gen_files(self) -> dict[str, float]:
        root = Path(self.tools.extra_env["HTML_SLIDE_CACHE"])
        return {str(p): p.stat().st_ino for p in root.rglob("gen-*.png") if "history" not in p.parts}
