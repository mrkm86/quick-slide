"""取り込んだページ（PPTX から）の画像と図解（docs/adr/0015-imported-images-and-diagrams.md）。

- 画像: 空いた所に足す（渡した画像か生成した画像）・差し替える（同じ枠のまま中身だけ）・主役にして組み直す（型のページ）
- 図解: 空いた所に足す（SVG の図）・主役にして組み直す（図解の型）
- 「差し替える」はページに画像があるときだけ画面が出す。組み直すは送る前に確かめる一言を持つ
本物のエンジン（html-slide）と偽物の台本係・偽物の Codex で通す。
"""

from __future__ import annotations

import io
import unittest

from helpers import fake_tools, png_bytes, start_server, temp_dir, wait_until_idle, writer_calls

from quick_slide import edit_kinds, importing
from quick_slide.settings import Settings, WorkerSettings

from test_engines_e2e import html_slide_ready


def two_pages(tmp) -> bytes:
    """1 枚目は題と箇条書き、2 枚目は写真 1 枚と題。どちらも見えるページ。"""
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Emu, Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "取り込みの見出し"
    s.placeholders[1].text = "一つ目の要点"
    photo = tmp / "p.png"
    Image.new("RGB", (320, 200), (30, 120, 200)).save(photo)
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "写真のページ"
    s2.shapes.add_picture(str(photo), Inches(1), Inches(2), width=Inches(4))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


IMG = ('<img data-name="背景" src="imported/media/bg.png" style="position:absolute;left:0px;top:0px;width:1920px;height:1080px" />'
       '<div><img src="imported/media/a.png" alt="" style="position:absolute;left:100px;top:200px;width:300px;height:150px;'
       'object-fit:cover" /></div>')


class UnitTests(unittest.TestCase):
    def test_images_are_listed_without_the_background_and_replaced_in_place(self) -> None:
        self.assertEqual(importing.page_images(IMG), [{"n": 1, "src": "imported/media/a.png", "x": 100, "y": 200, "w": 300, "h": 150}])
        out = importing.replace_image(IMG, 1, "overlays/x.png")
        self.assertIn('<img src="overlays/x.png" alt="" style="position:absolute;left:100px;top:200px;width:300px;height:150px;'
                      'object-fit:cover" />', out)
        self.assertIn('src="imported/media/bg.png"', out, "背景は替えない")

    def test_imported_kinds_ways_and_confirmation(self) -> None:
        kinds = {k["kind"]: k for k in edit_kinds.offered_for_imported("presentation", can_generate=True)}
        self.assertEqual([o["value"] for o in kinds["image"]["options"]], ["add", "replace", "main"])
        self.assertEqual([o["value"] for o in kinds["diagram"]["options"]], ["add", "main"])
        self.assertTrue(next(o for o in kinds["image"]["options"] if o["value"] == "replace")["needs_picture"])
        self.assertIn("組み直します", next(o for o in kinds["diagram"]["options"] if o["value"] == "main")["confirm"])
        self.assertFalse(any("attachment_reason" in o for o in kinds["image"]["options"]), "生成できれば画像は要らない")
        no_codex = {k["kind"]: k for k in edit_kinds.offered_for_imported("presentation", can_generate=False)}
        self.assertTrue(all("attachment_reason" in o for o in no_codex["image"]["options"] if o["value"] != "main"))
        self.assertEqual(edit_kinds.missing("image", ["add"], text="", attachments=0, imported=True, can_generate=False),
                         edit_kinds.IMPORTED_ATTACHMENT_REQUIRED)
        self.assertIsNone(edit_kinds.missing("image", ["main"], text="", attachments=0, imported=True, can_generate=False))
        self.assertEqual(edit_kinds.parse("presentation", "image", None, can_generate=True, imported=True), ("image", ["add"]))
        with self.assertRaises(ValueError):
            edit_kinds.parse("presentation", "image", ["generate"], can_generate=True, imported=True)


class EndToEndTests(unittest.TestCase):
    def setUp(self) -> None:
        why = html_slide_ready()
        if why:
            self.skipTest(why)
        self.tools, self.log = fake_tools(self)
        self.fx = start_server(self, tools=self.tools, settings=Settings(worker=WorkerSettings(parallel=1)))
        r = self.fx.request("POST", "/api/imports?runtime=cli&name=deck.pptx", raw_body=two_pages(temp_dir(self)),
                            session=True)
        self.assertEqual(r.status, 201, r.body)
        self.deck = r.json()["deck_id"]
        self.assertEqual(wait_until_idle(self.fx, self.deck, seconds=300)["state"], "done")
        self.view = self.fx.get(f"/api/decks/{self.deck}").json()

    def send(self, n: int, kind: str, option: str, *, picture: bool = False, text: str = "") -> dict:
        body = {"scope": "page", "slide_id": self.view["pages"][n - 1]["slide_id"], "text": text, "kind": kind,
                "options": [option]}
        if picture:
            up = self.fx.request("POST", "/api/settings/images", raw_body=png_bytes(), session=True)
            self.assertEqual(up.status, 201, up.body)
            body["attachments"] = [{"asset_id": up.json()["asset_id"]}]
        r = self.fx.post(f"/api/decks/{self.deck}/instructions", body)
        self.assertEqual(r.status, 200, r.body)
        run = wait_until_idle(self.fx, self.deck, seconds=300)
        self.assertEqual(run["state"], "done", run["message"])
        with self.fx.server.store_lock:
            return self.fx.store.load_script(self.fx.store.head_script(self.deck))["slides"][n - 1]

    def test_counts_of_pictures_per_page(self) -> None:
        counts = self.view["imported_images"]
        self.assertEqual([counts[p["slide_id"]] for p in self.view["pages"]], [0, 1], "差し替えは画像のあるページだけ")

    def test_add_a_diagram_a_given_picture_and_a_generated_picture(self) -> None:
        page = self.send(1, "diagram", "add")
        self.assertIn("<svg", page["illustrations"][-1]["image"]["svg"])
        page = self.send(1, "image", "add", picture=True)
        item = page["illustrations"][-1]
        self.assertTrue(item["image"]["path"].startswith("attachments/"))
        self.assertEqual(item["fit"], "cover")
        page = self.send(1, "image", "add")
        gen = page["illustrations"][-1]["image"]["path"]
        self.assertTrue(gen.startswith("overlays/gen-"), gen)
        self.assertTrue((self.fx.root.scripts_dir(self.deck) / gen).is_file(), "生成した画像は台本のフォルダに置く")
        self.assertEqual(list((self.fx.root.scripts_dir(self.deck) / "overlays").glob("gen-*/")), [], "仮の台本は残さない")
        boxes = [(i["x"], i["y"], i["x"] + i["w"], i["y"] + i["h"]) for i in page["illustrations"]]
        for a in range(len(boxes)):
            for b in range(a + 1, len(boxes)):
                (ax0, ay0, ax1, ay1), (bx0, by0, bx1, by1) = boxes[a], boxes[b]
                self.assertTrue(ax1 <= bx0 or bx1 <= ax0 or ay1 <= by0 or by1 <= ay0, f"重ねたものどうしが重なった: {boxes}")
        calls = [c for c in writer_calls(self.log) if c.get("what") == "overlay"]
        self.assertTrue(any("重ねた絵" in ln for ln in calls[-1]["prompt_lines"]), "前に重ねた絵を台本係に渡す")

    def test_replace_the_picture_in_place(self) -> None:
        before = len(writer_calls(self.log))
        page = self.send(2, "image", "replace", picture=True)
        self.assertEqual(len(writer_calls(self.log)), before, "画像が 1 つで、渡した画像があれば台本係を呼ばない")
        srcs = [i["src"] for i in importing.page_images(page["html"])]
        self.assertEqual(len(srcs), 1)
        self.assertTrue(srcs[0].startswith("attachments/"), srcs)

    def test_rebuild_around_a_picture_and_around_a_diagram(self) -> None:
        page = self.send(1, "image", "main", picture=True)
        self.assertEqual(page["type"], "photo")
        self.assertTrue(page["image"]["path"].startswith("attachments/"))
        self.assertEqual(page["id"], self.view["pages"][0]["key"], "ページの id は保つ")
        page = self.send(2, "diagram", "main")
        self.assertEqual(page["type"], "flow")
        view = self.fx.get(f"/api/decks/{self.deck}").json()
        self.assertEqual(view["imported_images"], {}, "組み直したページは取り込んだページでなくなる")


if __name__ == "__main__":
    unittest.main()
