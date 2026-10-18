"""既存の PPTX の取り込み（`POST /api/imports`・docs/adr/0011-import-pptx.md）。

- 口の断り: セッションが無い・PPTX でない名前・zip でない中身・大きすぎる（読む前に 413）。断った本文の一時ファイルは残さない
- 通し（本物の html-slide。Node と Chrome が要る。無ければ skip）: 取り込み → 台本の版（書き手は import・ページの id・
  画像は imported/media/）→ 描いた回（ロゴ・フッタ・ページ番号は足さない）→ スライドの画面のページ・発表の画面の HTML
- 語彙: 取り込めるのはプレゼンだけ（`deck_kinds[].importable`）・ボタンの語
- 複製: 取り込んだ画像（imported/）も写す
- 非表示のページの下見の絵（docs/adr/0012-hidden-page-previews.md）: 描いて画面に出す・枚数と PPTX には入れない・複製でも写す
"""

from __future__ import annotations

import io
import json
import sqlite3
import unittest
from unittest import mock

from helpers import fake_tools, start_server, temp_dir, wait_until_idle
from test_engines_e2e import html_slide_ready

from quick_slide import importing
from quick_slide.settings import Settings, WorkerSettings


def pptx_bytes(tmp) -> bytes:
    """2 枚の PPTX（題と箇条書き・写真）。2 枚目は非表示。"""
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Emu, Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "取り込みの見出し"
    s.placeholders[1].text = "一つ目の要点"
    s.notes_slide.notes_text_frame.text = "話すこと"
    photo = tmp / "p.png"
    Image.new("RGB", (320, 200), (30, 120, 200)).save(photo)
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    s2.shapes.add_picture(str(photo), Inches(1), Inches(1), width=Inches(4))
    s2._element.set("show", "0")
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


class ImportRouteRefusalTests(unittest.TestCase):
    def setUp(self) -> None:
        self.fx = start_server(self)

    def uploads(self) -> list:
        d = self.fx.root.uploads_dir
        return sorted(p.name for p in d.iterdir()) if d.is_dir() else []

    def test_needs_session_before_reading_the_body(self) -> None:
        r = self.fx.request("POST", "/api/imports?name=a.pptx", raw_body=b"PK\x03\x04" + b"0" * 100)
        self.assertEqual(r.status, 403, r.body)
        self.assertEqual(self.uploads(), [])

    def test_refuses_wrong_name_and_wrong_bytes(self) -> None:
        r = self.fx.request("POST", "/api/imports?name=notes.txt", raw_body=b"PK\x03\x04abc", session=True)
        self.assertEqual((r.status, r.json()["error"]["code"]), (400, "import_refused"))
        r = self.fx.request("POST", "/api/imports?name=deck.pptx", raw_body=b"not a zip", session=True)
        self.assertEqual((r.status, r.json()["error"]["code"]), (400, "import_refused"))
        self.assertEqual(self.uploads(), [], "断った本文の一時ファイルを残さない")
        with self.fx.server.store_lock:
            self.assertEqual(self.fx.store.conn.execute("SELECT COUNT(*) FROM decks WHERE deleted_at IS NULL").fetchone()[0],
                             len(self.fx.store.conn.execute("SELECT id FROM decks WHERE deleted_at IS NULL").fetchall()))
            self.assertEqual(self.fx.store.conn.execute("SELECT COUNT(*) FROM intakes WHERE origin_kind='file'").fetchone()[0], 0)

    def test_screen_sends_the_name_in_the_query(self) -> None:
        """画面は名前を問い合わせ（encodeURIComponent）で送る。日本語も「%」もそのまま、拡張子を外してデッキの名前になる。"""
        from urllib.parse import quote
        name = "【提案】会議を 50%20短く.pptx"   # 「%20」は二度戻すと空白に崩れる
        with mock.patch.object(type(self.fx.server), "kind_refusals", return_value={}):
            r = self.fx.request("POST", f"/api/imports?runtime=browser&name={quote(name)}",
                                raw_body=pptx_bytes(temp_dir(self)), session=True)
        self.assertEqual(r.status, 201, r.body)
        self.assertEqual(r.json()["deck_title"], "【提案】会議を 50%20短く")

    def test_too_large_is_refused_before_reading(self) -> None:
        r = self.fx.request("POST", "/api/imports?name=big.pptx", raw_body=b"PK", session=True,
                            headers={"Content-Length": str(importing.MAX_PPTX_BYTES + 1)})
        self.assertEqual(r.status, 413, r.body)
        self.assertEqual(self.uploads(), [])

    def test_other_routes_keep_the_small_body_limit(self) -> None:
        r = self.fx.request("POST", "/api/intakes", raw_body=b"x", session=True,
                            headers={"Content-Length": str(17 * 1024 * 1024)})
        self.assertEqual(r.status, 413)

    def test_vocabulary_says_only_presentation_imports(self) -> None:
        v = self.fx.get("/api/vocabulary").json()
        self.assertEqual({k["mode"]: k["importable"] for k in v["deck_kinds"]},
                         {"presentation": True, "pitch": False, "motion": False})
        self.assertEqual(v["import"]["button"], "プレゼン資料を取り込む")
        self.assertEqual(v["create_limits"]["import_bytes"], importing.MAX_PPTX_BYTES)


class ImportEndToEndTests(unittest.TestCase):
    def test_import_makes_a_presentation_deck(self) -> None:
        why = html_slide_ready()
        if why:
            self.skipTest(why)
        tools, _ = fake_tools(self)
        fx = start_server(self, tools=tools, settings=Settings(worker=WorkerSettings(parallel=1)))
        data = pptx_bytes(temp_dir(self))
        r = fx.request("POST", "/api/imports?runtime=cli", raw_body=data, session=True,
                       headers={"X-Quick-Slide-Filename": "%E5%8F%96%E3%82%8A%E8%BE%BC%E3%81%BF.pptx"})
        self.assertEqual(r.status, 201, r.body)
        body = r.json()
        self.assertEqual((body["mode"], body["deck_title"], body["import"]["bytes"]), ("presentation", "取り込み", len(data)))
        self.assertIn("token=", body["url"])
        run = wait_until_idle(fx, body["deck_id"], seconds=300)
        self.assertEqual(run["state"], "done", run["message"])
        self.assertTrue(run["args"]["imported"])
        look = run["args"]["look"]
        self.assertEqual((look["logo"], look["footer"], look["page_numbers"]), (None, "", False))
        root = fx.root
        self.assertEqual(root.import_source_path(body["deck_id"]).read_bytes(), data, "元の PPTX を残す")
        with fx.server.store_lock:
            intake = fx.store.intake(body["intake_id"])
            script = fx.store.script(run["script_id"])
            raw = fx.store.load_script(script)
        self.assertEqual((intake["origin_kind"], intake["origin_runtime"], intake["manuscript"]), ("file", "cli", ""))
        self.assertEqual(intake["payload"]["import"]["bytes"], len(data))
        self.assertEqual(script["written_by"], "import")
        self.assertEqual([s["id"] for s in raw["slides"]], ["p-001", "p-002"])
        self.assertTrue(raw["slides"][1]["_hidden"], "元の非表示のスライドは非表示のまま")
        img = [p for p in raw["slides"][1]["html"].split('"') if p.startswith("imported/media/")]
        self.assertTrue(img and (root.scripts_dir(body["deck_id"]) / img[0]).is_file(), img)
        self.assertEqual(raw["slides"][0]["notes"], "話すこと")
        view = fx.get(f"/api/decks/{body['deck_id']}").json()
        self.assertEqual([p["title"] for p in view["pages"]][0], "取り込みの見出し")
        self.assertIn("一つ目の要点", json.dumps(view["pages"][0], ensure_ascii=False))
        self.assertNotIn("<div", view["pages"][0].get("body") or "", "本文は HTML の見える文字だけ")
        out = root.run_dir(body["deck_id"], run["id"])
        self.assertTrue((out / "deck.pptx").is_file())
        # 非表示のページも下見の絵を持ち（画面は薄く出す）、発表・書き出し・枚数には入らない（docs/adr/0012-hidden-page-previews.md）
        hidden = view["pages"][1]
        self.assertTrue(hidden["hidden"])
        self.assertTrue(hidden["image_url"], "非表示のページの下見の絵が無い")
        self.assertEqual(hidden["display"]["key"], "ready")
        pic = fx.get(hidden["image_url"])
        self.assertEqual((pic.status, pic.headers["content-type"]), (200, "image/png"))
        card = fx.get("/api/decks?limit=5").json()["items"][0]
        self.assertEqual(card["page_count"], 1, "下見の絵は枚数に数えない")
        import zipfile
        with zipfile.ZipFile(out / "deck.pptx") as z:
            self.assertEqual(len([n for n in z.namelist() if n.startswith("ppt/slides/slide")]), 1, "PPTX に非表示のページを入れない")
        # 取り込んだページの直しは 1 ページずつ（docs/adr/0013〜0015）。選んだページをまとめて直すのは理由を返す
        r = fx.post(f"/api/decks/{body['deck_id']}/instructions",
                    {"scope": "pages", "slide_ids": [view["pages"][0]["slide_id"]], "text": "見出しを短く"})
        self.assertEqual((r.status, r.json()["error"]["code"]), (409, "imported_page"), r.body)
        # 複製しても取り込んだ画像を写す（台本の html が相対で指す）
        r = fx.post(f"/api/decks/{body['deck_id']}/duplicate", {})
        self.assertEqual(r.status, 201, r.body)
        new_id = r.json()["deck"]["id"]
        self.assertTrue((root.scripts_dir(new_id) / img[0]).is_file(), "複製に imported/ が無い")
        dup = fx.get(f"/api/decks/{new_id}").json()
        self.assertTrue(dup["pages"][1]["image_url"], "複製でも非表示のページの下見の絵を写す")
        self.assertNotEqual(dup["pages"][1]["image_url"], hidden["image_url"])


class ImportStoreTests(unittest.TestCase):
    def test_scripts_accept_import_as_writer(self) -> None:
        fx = start_server(self)
        with fx.server.store_lock:
            res = fx.store.create_import(filename="a.pptx", info={"name": "a.pptx", "bytes": 10, "sha256": "0" * 64},
                                         runtime="browser")
            s = fx.store.save_script(res["deck_id"], {"slides": [{"id": "p-001", "type": "html", "html": ""}]},
                                     written_by="import", writer_backend=None, intake_id=res["intake_id"], run_id=None)
            self.assertEqual(s["written_by"], "import")
            with self.assertRaises(sqlite3.IntegrityError):
                fx.store.conn.execute("UPDATE scripts SET written_by='other' WHERE id=?", (s["id"],))

    def test_to_script_rewrites_media_and_numbers_pages(self) -> None:
        out = importing.to_script({"title": "t", "theme": "theme.yaml", "slides": [
            {"type": "html", "html": '<img src="media/s01-a.png">'}, {"type": "html", "html": "", "id": "x"}]})
        self.assertEqual([s["id"] for s in out["slides"]], ["p-001", "p-002"])
        self.assertIn('src="imported/media/s01-a.png"', out["slides"][0]["html"])
        self.assertEqual(out["theme"], "imported/theme.yaml")
        self.assertEqual(importing.clean_filename("../../x\x00/evil.pptx"), "evil.pptx")
        # デッキの名前はファイル名から拡張子を外したもの（記号・途中の . はそのまま）
        self.assertEqual(importing.title_of("【提案】研修 v1.2.pptx"), "【提案】研修 v1.2")
        self.assertEqual(importing.title_of("a.PPTX"), "a")
        self.assertEqual(len(importing.title_of("あ" * 300 + ".pptx")), importing.MAX_TITLE_CHARS)
