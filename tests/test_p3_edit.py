"""P3（直す）: 10 ページの 1 冊のうち 7・8 枚目だけを直し、ほかのページは台本の版でも 1 文字も変わらず、
描き直されもしないこと。画面のページと描いた回の PPTX の並び・枚数が一致すること。

- 台本係は偽物（`tests/fakes/writer_common.py`。指示の語で振り分けを決める）、**エンジンは本物**、絵を描く codex は
  エンジンの偽物（呼ばれた回とページを記録する）
- ピッチ（image-slide）は描いたページを codex の呼び出しの記録から数える（描く指示に「このページは N枚目」がある）
- プレゼン（html-slide）は Node と Chrome が無ければ skip。生成の枠の絵（区画の `gen-<鍵>.png`）が替わったかで数える
"""

from __future__ import annotations

import json
import time
import unittest
from pathlib import Path

from helpers import FAKES, fake_tools, intake_payload, runnable, start_server, wait_until_idle, writer_calls

from quick_slide.scriptdoc import dumps, without_pins
from quick_slide.settings import Settings, WorkerSettings

from test_engines_e2e import html_slide_ready

CONFIG = Settings(worker=WorkerSettings(parallel=2))


def pptx_notes(path: Path) -> list[str]:
    from pptx import Presentation

    prs = Presentation(str(path))
    out = []
    for slide in prs.slides:
        out.append(slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else "")
    return out


class EditBase(unittest.TestCase):
    mode = "pitch"
    runtime = "codex"

    def setUp(self) -> None:
        varied = runnable(FAKES / "fake_codex_varied.py")   # 描くたびに違う絵（「絵だけ直した」を中身で見分ける）
        self.tools, self.log = fake_tools(self, extra={"FAKE_WRITER_PAGES": "10", "FAKE_CODEX_SLEEP": "0.05",
                                                       "IMAGE_SLIDE_CODEX": varied, "HTML_SLIDE_CODEX": varied})
        self.fx = start_server(self, tools=self.tools, settings=CONFIG)

    # --- 道具 ---
    def make_deck(self, **attrs) -> str:
        r = self.fx.post("/api/intakes", intake_payload(runtime=self.runtime, mode=self.mode, **attrs))
        self.assertEqual(r.status, 201, r.body)
        deck = r.json()["deck_id"]
        run = wait_until_idle(self.fx, deck, seconds=300)
        self.assertEqual(run["state"], "done", run["message"])
        return deck

    def view(self, deck: str) -> dict:
        return self.fx.get(f"/api/decks/{deck}").json()

    def wait(self, deck: str) -> dict:
        run = wait_until_idle(self.fx, deck, seconds=300)
        return run

    def head(self, deck: str) -> tuple[dict, bytes]:
        with self.fx.server.store_lock:
            row = self.fx.store.head_script(deck)
            path = self.fx.store.script_path(row)
        return json.loads(path.read_text(encoding="utf-8")), path.read_bytes()

    def codex_calls(self) -> set[str]:
        d = self.log / "codex-calls"
        return {p.name for p in d.glob("*.json")} if d.is_dir() else set()

    def drawn_pages(self, before: set[str]) -> list[int]:
        out = []
        for name in sorted(self.codex_calls() - before):
            rec = json.loads((self.log / "codex-calls" / name).read_text(encoding="utf-8"))
            if rec.get("cmd") == "exec" and "t1" in rec:
                out.append(rec["page"])
        return sorted(out)

    def instruct(self, deck: str, positions: list[int], text: str, **extra) -> dict:
        pages = self.view(deck)["pages"]
        ids = [pages[n - 1]["slide_id"] for n in positions]
        r = self.fx.post(f"/api/decks/{deck}/instructions", {"scope": "pages", "slide_ids": ids, "text": text, **extra})
        self.assertEqual(r.status, 200, r.body)
        return r.json()

    def assert_screen_matches_pptx(self, deck: str) -> None:
        view = self.view(deck)
        visible = [p for p in view["pages"] if not p["hidden"]]
        with self.fx.server.store_lock:
            run = self.fx.store.shown_run(deck)
        notes = pptx_notes(self.fx.root.run_dir(deck, run["id"]) / "deck.pptx")
        data, _ = self.head(deck)
        talk = {s["id"]: s["talk"][0] for s in data["slides"]}
        self.assertEqual(len(notes), len(visible), "PPTX の枚数は画面の（非表示を除く）ページ数")
        self.assertEqual([n.replace("［クリック］", "").strip() for n in notes], [talk[p["key"]] for p in visible],
                         "PPTX の並びは画面の並び")
        for p in visible:
            self.assertEqual(p["display"]["key"], "ready", p)
        self.assertFalse(view["live"]["working"])

    def duplicate_and_rebuild(self, deck: str) -> str:
        """複製して、同じ姿か（ページ・絵・書き出し）確かめ、複製を組み直す。複製の id を返す。"""
        src_view = self.view(deck)
        r = self.fx.post(f"/api/decks/{deck}/duplicate", {"title": "写し"})
        self.assertEqual(r.status, 201, r.body)
        new = r.json()["deck"]["id"]
        view = self.view(new)
        self.assertEqual([(p["key"], p["title"], p["hidden"]) for p in view["pages"]],
                         [(p["key"], p["title"], p["hidden"]) for p in src_view["pages"]], "複製はいまの姿のページを持つ")
        for p in view["pages"]:
            if not p["hidden"]:
                self.assertEqual(p["display"]["key"], "ready", p)
                self.assertEqual(len(p["versions"]), 1, "版の履歴は写さない")
        with self.fx.server.store_lock:
            a, b = self.fx.store.shown_run(deck), self.fx.store.shown_run(new)
        self.assertEqual((self.fx.root.run_dir(new, b["id"]) / "deck.pptx").read_bytes(),
                         (self.fx.root.run_dir(deck, a["id"]) / "deck.pptx").read_bytes())
        self.assert_screen_matches_pptx(new)
        calls = self.codex_calls()
        r = self.fx.post(f"/api/decks/{new}/rebuild", {})
        self.assertEqual(r.status, 200, r.body)
        run = self.wait(new)
        self.assertEqual(run["state"], "done", run["message"])
        self.assertEqual(self.drawn_pages(calls), [], "複製を組み直しても絵は描き直さない（控えから引き継ぐ）")
        self.assert_screen_matches_pptx(new)
        self.assertEqual([p["key"] for p in self.view(deck)["pages"]], [p["key"] for p in src_view["pages"]], "元は変わらない")
        return new

    def page_bytes(self, data_bytes: bytes, data: dict) -> dict[str, str]:
        """台本の版のファイルの中の、ページごとの文字（同じ書き方で書いたときの断片）。"""
        text = data_bytes.decode("utf-8")
        out = {}
        for s in data["slides"]:
            piece = json.dumps(s, ensure_ascii=False, indent=1).replace("\n", "\n  ")
            self.assertIn(piece, text)
            out[s["id"]] = piece
        return out


class PitchEditTests(EditBase):
    mode = "pitch"
    runtime = "codex"

    def test_revise_pages_7_and_8_redraws_only_them(self) -> None:
        deck = self.make_deck()
        before_data, before_bytes = self.head(deck)
        before_pieces = self.page_bytes(before_bytes, before_data)
        old = {p["key"]: p for p in self.view(deck)["pages"]}
        calls = self.codex_calls()
        r = self.instruct(deck, [7, 8], "見出しを短くして")
        self.assertIn("7・8ページ", r["message"])
        run = self.wait(deck)
        self.assertEqual(run["state"], "done", run["message"])
        self.assertEqual(self.drawn_pages(calls), [7, 8], "描き直すのは 7・8 枚目だけ（ほかはキャッシュの区画から）")
        after_data, after_bytes = self.head(deck)
        after_pieces = self.page_bytes(after_bytes, after_data)
        for key in before_pieces:
            if key in ("p-seven", "p-eight"):
                self.assertNotEqual(before_pieces[key], after_pieces[key])
            else:
                self.assertEqual(before_pieces[key], after_pieces[key], f"{key} は 1 文字も変わらない")
        self.assertEqual(after_data["slides"][6]["headline"], "立ったまま話すと短くなる（直し）")
        view = self.view(deck)
        for p in view["pages"]:
            if p["key"] in ("p-seven", "p-eight"):
                self.assertNotEqual(p["image_url"], old[p["key"]]["image_url"])
                self.assertEqual([v["version"] for v in p["versions"]], [2, 1])
                self.assertTrue(p["versions"][0]["active"])
                self.assertEqual(p["versions"][0]["revision_note"], "見出しを短くして")
            else:
                self.assertEqual(len(p["versions"]), 1, p["key"])
        self.assert_screen_matches_pptx(deck)
        # 提案（6 枚目へ）は当てずに、画面の帯と仕事の答えに出す
        job = self.fx.get(f"/api/decks/{deck}/jobs/{r['job_id']}").json()
        self.assertEqual([s["id"] for s in job["suggestions"]], ["p-rule"])
        self.assertEqual(job["routes"], {"p-seven": "script", "p-eight": "script"})
        self.assertIn("提案", view["planning"]["message"])
        self.assertEqual([p["title"] for p in job["pages"]][6], "立ったまま話すと短くなる（直し）")
        # 2 回目の回もエンジンは同じ作業フォルダに描き、回のフォルダに控える
        self.assertTrue((self.fx.root.run_dir(deck, run["id"]) / "deck.pptx").is_file())
        self.assertTrue((self.fx.root.run_dir(deck, run["id"]) / "pictures" / "p-seven.png").is_file())

    def test_duplicate_that_fails_midway_leaves_nothing(self) -> None:
        deck = self.make_deck()
        store = self.fx.store
        assets = self.fx.root.assets_dir / "decks"
        before = {p.name for p in assets.iterdir()}
        decks_before = {p.name for p in self.fx.root.decks_dir.iterdir()}
        real, calls = store.register_asset, []

        def broken(**kw):
            calls.append(1)
            if len(calls) == 3:
                raise OSError("容量不足のつもり")
            return real(**kw)

        store.register_asset = broken
        try:
            with self.fx.server.store_lock, self.assertRaises(OSError):
                store.duplicate_deck(deck)
        finally:
            store.register_asset = real
        self.assertEqual({p.name for p in assets.iterdir()}, before, "写しかけの絵が残った")
        self.assertEqual({p.name for p in self.fx.root.decks_dir.iterdir()}, decks_before, "写しかけのデッキの置き場が残った")
        with self.fx.server.store_lock:
            self.assertEqual([d for d in store.recent_decks(50) if d.get("origin_deck_id") == deck], [])

    def test_duplicate_copies_pages_and_rebuild_draws_nothing(self) -> None:
        deck = self.make_deck()
        self.instruct(deck, [3], "見出しを変えて")
        self.assertEqual(self.wait(deck)["state"], "done")
        self.duplicate_and_rebuild(deck)

    def test_second_pitch_run_with_one_changed_page_draws_only_that_page(self) -> None:
        """キャッシュの区画を回をまたいで使う（P2 は回ごとに新しい区画で、全ページを描き直していた）。"""
        deck = self.make_deck()
        calls = self.codex_calls()
        self.instruct(deck, [3], "見出しを変えて")
        run = self.wait(deck)
        self.assertEqual(run["state"], "done", run["message"])
        self.assertEqual(self.drawn_pages(calls), [3])

    def test_image_only_route_uses_engine_revise(self) -> None:
        deck = self.make_deck()
        before_data, before_bytes = self.head(deck)
        old = {p["key"]: p for p in self.view(deck)["pages"]}
        calls = self.codex_calls()
        r = self.instruct(deck, [8], "絵だけ夜の景色にして", edit_mode="minor")
        run = self.wait(deck)
        self.assertEqual(run["state"], "done", run["message"])
        self.assertEqual(self.drawn_pages(calls), [8])
        self.assertEqual(self.head(deck)[1], before_bytes, "絵だけの直しは台本を変えない")
        argv = run["record"]["argv"]
        self.assertEqual([a[2] for a in argv], ["revise"], "deck を流さず revise だけ")
        self.assertIn("--page", argv[0])
        self.assertEqual(argv[0][argv[0].index("--page") + 1], "8")
        self.assertNotIn("--remake", argv[0])
        # 微修正: 今の絵を土台に渡した（参照画像が 1 枚）
        rec = [json.loads((self.log / "codex-calls" / n).read_text(encoding="utf-8")) for n in sorted(self.codex_calls() - calls)]
        self.assertEqual([len(x["refs"]) for x in rec if x.get("cmd") == "exec" and "t1" in x], [1])
        p8 = next(p for p in self.view(deck)["pages"] if p["key"] == "p-eight")
        self.assertNotEqual(p8["image_url"], old["p-eight"]["image_url"])
        self.assertEqual([v["version"] for v in p8["versions"]], [2, 1])
        job = self.fx.get(f"/api/decks/{deck}/jobs/{r['job_id']}").json()
        self.assertEqual(job["routes"], {"p-eight": "image"})
        self.assert_screen_matches_pptx(deck)
        # 版から戻す: 絵だけが違ったので、前の絵を区画へ戻し、描かずに組み直す
        calls = self.codex_calls()
        older = p8["versions"][1]
        r = self.fx.post(f"/api/decks/{deck}/adopt", {"choices": {p8["slide_id"]: older["id"]}})
        self.assertEqual(r.status, 200, r.body)
        self.assertTrue(r.json()["started"])
        run = self.wait(deck)
        self.assertEqual(run["state"], "done", run["message"])
        self.assertEqual(self.drawn_pages(calls), [], "戻すときは描かない")
        p8 = next(p for p in self.view(deck)["pages"] if p["key"] == "p-eight")
        self.assertEqual(p8["versions"][0]["version"], 1)
        self.assertTrue(p8["versions"][0]["active"], "戻した姿が使用中")
        self.assertEqual(len(p8["versions"]), 2, "同じ姿は 1 つにまとめる")
        self.assertEqual(dumps(without_pins(self.head(deck)[0])), before_bytes, "台本は（版の印を除いて）元のまま")
        with self.fx.server.store_lock:
            rows = {r["key"]: r for r in self.fx.store.run_pages(run["id"])}
            first = {r["key"]: r for r in self.fx.store.run_pages(self.fx.store.runs_for_deck(deck)[-1]["id"])}
        self.assertEqual(rows["p-eight"]["picture_sha256"], first["p-eight"]["picture_sha256"])
        self.assert_screen_matches_pptx(deck)
        # もう使用中の版を押しても何も作らない
        r = self.fx.post(f"/api/decks/{deck}/adopt", {"choices": {p8["slide_id"]: p8["versions"][0]["id"]}})
        self.assertEqual((r.status, r.json()["started"]), (200, False))

    def test_both_route_rewrites_then_revises(self) -> None:
        deck = self.make_deck()
        calls = self.codex_calls()
        self.instruct(deck, [7], "両方: 見出しを短くして絵も明るく")
        run = self.wait(deck)
        self.assertEqual(run["state"], "done", run["message"])
        self.assertEqual([a[2] for a in run["record"]["argv"]], ["deck", "revise"])
        self.assertEqual(self.drawn_pages(calls), [7, 7], "台本の差し替えで 1 回、絵の直しで 1 回")
        self.assertEqual(self.head(deck)[0]["slides"][6]["headline"], "立ったまま話すと短くなる（直し）")
        self.assert_screen_matches_pptx(deck)

    def test_restore_script_version_draws_nothing(self) -> None:
        deck = self.make_deck()
        _, first_bytes = self.head(deck)
        self.instruct(deck, [7], "見出しを短くして")
        self.assertEqual(self.wait(deck)["state"], "done")
        p7 = next(p for p in self.view(deck)["pages"] if p["key"] == "p-seven")
        calls = self.codex_calls()
        r = self.fx.post(f"/api/decks/{deck}/adopt", {"choices": {p7["slide_id"]: p7["versions"][1]["id"]}})
        self.assertEqual(r.status, 200, r.body)
        run = self.wait(deck)
        self.assertEqual(run["state"], "done", run["message"])
        self.assertEqual(self.drawn_pages(calls), [], "前の絵を区画へ戻すので描かない")
        data, _ = self.head(deck)
        self.assertEqual(next(s for s in data["slides"] if s["id"] == "p-seven").get("_pin"), p7["versions"][1]["id"],
                         "戻したページは、その版を描いた回を指す")
        self.assertEqual(dumps(without_pins(data)), first_bytes, "台本も（版の印を除いて）1 冊目と同じ姿に戻る")
        p7 = next(p for p in self.view(deck)["pages"] if p["key"] == "p-seven")
        self.assertEqual((p7["version"], p7["title"]), (1, "立ったまま話すと短くなる"))
        self.assert_screen_matches_pptx(deck)

    def test_insert_duplicate_hide_and_reorder(self) -> None:
        deck = self.make_deck()
        pages = self.view(deck)["pages"]
        # 差し込み（台本係が 1 ページ書く）: 足したページだけ描く
        calls = self.codex_calls()
        r = self.fx.post(f"/api/decks/{deck}/pages/insert", {"after_slide_id": pages[1]["slide_id"], "text": "費用の内訳が欲しい"})
        self.assertEqual(r.status, 200, r.body)
        self.assertEqual(self.view(deck)["pending_inserts"][0]["after_slide_id"], pages[1]["slide_id"])
        self.assertEqual(self.wait(deck)["state"], "done")
        self.assertEqual(self.drawn_pages(calls), [3], "足したページ（3 枚目）だけ描く。後ろのページは区画で付け替える")
        keys = [p["key"] for p in self.view(deck)["pages"]]
        self.assertEqual(keys[:4], ["cover", "p-why", "p-added-1", "p-sec"])
        self.assertEqual(writer_calls(self.log)[-1]["instruction"], "費用の内訳が欲しい")
        self.assert_screen_matches_pptx(deck)
        # 複製: 同じ中身なので描かない。複製したページの絵は元と同じバイトだが、どのページも自分の絵の URL を保つ
        calls = self.codex_calls()
        urls = {p["key"]: p["image_url"] for p in self.view(deck)["pages"]}
        p5 = self.view(deck)["pages"][4]
        r = self.fx.post(f"/api/decks/{deck}/pages/{p5['slide_id']}/duplicate", {})
        self.assertEqual(r.status, 200, r.body)
        self.assertEqual(self.wait(deck)["state"], "done")
        self.assertEqual(self.drawn_pages(calls), [])
        keys = [p["key"] for p in self.view(deck)["pages"]]
        self.assertEqual(keys[4:6], ["p-agenda", "p-agenda-copy"])
        now = {p["key"]: p["image_url"] for p in self.view(deck)["pages"]}
        self.assertEqual({k: now[k] for k in urls}, urls, "直していないページの絵の URL は変わらない")
        self.assertNotEqual(now["p-agenda-copy"], now["p-agenda"], "複製は自分の資産")
        self.assert_screen_matches_pptx(deck)
        urls = now
        # 非表示: エンジンには渡さず、画面には残る（灰色）。描かない
        calls = self.codex_calls()
        target = self.view(deck)["pages"][1]
        r = self.fx.post(f"/api/decks/{deck}/pages/{target['slide_id']}/visibility", {"hidden": True})
        self.assertEqual(r.status, 200, r.body)
        self.assertEqual(self.wait(deck)["state"], "done")
        self.assertEqual(self.drawn_pages(calls), [])
        view = self.view(deck)
        hidden = [p for p in view["pages"] if p["hidden"]]
        self.assertEqual([p["key"] for p in hidden], ["p-why"])
        self.assertEqual({p["key"]: p["image_url"] for p in view["pages"]}, urls, "並びがずれても絵の URL は変わらない")
        self.assertTrue(hidden[0]["image_url"], "隠したページも前の絵で出す")
        self.assertEqual(view["counts"]["total"], 11)
        self.assert_screen_matches_pptx(deck)
        # 並べ替え: 描かない。答えは新しい並び
        calls = self.codex_calls()
        ids = [p["slide_id"] for p in view["pages"]]
        ids = [ids[-1]] + ids[:-1]
        r = self.fx.post(f"/api/decks/{deck}/pages/reorder", {"slide_ids": ids})
        self.assertEqual(r.status, 200, r.body)
        self.assertEqual([p["slide_id"] for p in r.json()["pages"]], ids)
        self.assertEqual(self.wait(deck)["state"], "done")
        self.assertEqual(self.drawn_pages(calls), [])
        self.assertEqual([p["slide_id"] for p in self.view(deck)["pages"]], ids)
        self.assert_screen_matches_pptx(deck)
        # 表示に戻す: 前の絵（控え）を区画に戻すので描かない
        calls = self.codex_calls()
        r = self.fx.post(f"/api/decks/{deck}/pages/{target['slide_id']}/visibility", {"hidden": False})
        self.assertEqual(self.wait(deck)["state"], "done")
        self.assertEqual(self.drawn_pages(calls), [])
        self.assert_screen_matches_pptx(deck)

    def test_two_requests_queue_and_the_second_builds_on_the_first(self) -> None:
        deck = self.make_deck()
        self.tools.extra_env["FAKE_CODEX_SLEEP"] = "1.5"
        a = self.instruct(deck, [7], "見出しを短くして")
        b = self.instruct(deck, [2], "見出しを短くして")
        self.assertTrue(b["queued"])
        self.assertIn("前の直しが終わってから", b["message"])
        deadline = time.monotonic() + 60
        seen_both_running = False
        while time.monotonic() < deadline:
            with self.fx.server.store_lock:
                ja, jb = self.fx.store.job(a["job_id"]), self.fx.store.job(b["job_id"])
            if ja["state"] == "running" and jb["state"] == "running":
                seen_both_running = True
            if jb["state"] in ("done", "failed"):
                break
            time.sleep(0.1)
        self.assertFalse(seen_both_running, "1 冊に走る仕事は 1 つだけ")
        self.assertEqual((ja["state"], jb["state"]), ("done", "done"))
        with self.fx.server.store_lock:
            ra, rb = self.fx.store.run(a["run_id"]), self.fx.store.run(b["run_id"])
            evs = self.fx.store.events(deck_id=deck, kinds=("job_started", "job_finished"))
        order = [(e["kind"], e["payload"]["job_id"]) for e in evs if e["payload"]["job_id"] in (a["job_id"], b["job_id"])]
        self.assertEqual(order, [("job_started", a["job_id"]), ("job_finished", a["job_id"]),
                                 ("job_started", b["job_id"]), ("job_finished", b["job_id"])])
        data, _ = self.head(deck)
        self.assertTrue(data["slides"][6]["headline"].endswith("（直し）"), "2 つ目は 1 つ目の直しの上に重なる")
        self.assertTrue(data["slides"][1]["headline"].endswith("（直し）"))
        self.assertEqual(rb["record"]["base"]["version"], ra["record"]["edit"]["script_version"])
        self.assert_screen_matches_pptx(deck)

    def legacy_first_run(self, deck: str) -> dict:
        """1 回目の回を P2（移行 0003 の前）の形にする: ページの中身と絵の sha256 が無い。"""
        with self.fx.server.store_lock:
            first = self.fx.store.runs_for_deck(deck)[-1]
            self.fx.store.conn.execute("UPDATE run_pages SET content_sha256=NULL,picture_sha256=NULL WHERE run_id=?",
                                       (first["id"],))
            self.fx.store.conn.execute("UPDATE script_pages SET content_sha256=NULL WHERE script_id=?",
                                       (first["script_id"],))
        return first

    def check_legacy_versions(self, deck: str) -> None:
        for p in self.view(deck)["pages"]:
            if p["key"] in ("p-why", "p-agenda"):
                self.assertEqual([v["version"] for v in p["versions"]], [2, 1], p["key"])
                self.assertEqual([v["revision_note"] for v in p["versions"]], ["見出しを短くして", None],
                                 "指示の文は、その指示が向いたページの版にだけ付く")
            else:
                self.assertEqual([(v["version"], v["active"], v["revision_note"]) for v in p["versions"]],
                                 [(1, True, None)], f"{p['key']}: 直していないページは版が 1 つ・指示も付かない")

    def test_legacy_run_without_shas_does_not_add_versions(self) -> None:
        deck = self.make_deck()
        first = self.legacy_first_run(deck)
        self.instruct(deck, [2, 4], "見出しを短くして")
        self.assertEqual(self.wait(deck)["state"], "done")
        self.check_legacy_versions(deck)
        # その回の台本の版のファイルが読めなくても、画像のバイトが同じなら同じ姿とみなす
        from quick_slide.server.page_view import LegacyContent, resolve_states

        with self.fx.server.store_lock:
            path = self.fx.store.script_path(self.fx.store.script(first["script_id"]))
        path.rename(path.with_suffix(".gone"))
        LegacyContent._memo.clear()
        self.check_legacy_versions(deck)
        rows = [{"asset_sha256": "x", "content_sha256": "c2", "picture_sha256": "p2"},
                {"asset_sha256": "x", "content_sha256": None, "picture_sha256": None},
                {"asset_sha256": "y", "content_sha256": None, "picture_sha256": None}]
        self.assertEqual(resolve_states(rows, lambda r: None), [("c2", "p2"), ("c2", "p2"), ("asset:y", "asset:y")])
        self.assertEqual(resolve_states(rows, lambda r: "c-old"), [("c2", "p2"), ("c-old", "p2"), ("c-old", "asset:y")])

    def test_a_gone_page_fails_with_a_clear_message(self) -> None:
        deck = self.make_deck()
        pages = self.view(deck)["pages"]
        with self.fx.server.store_lock:
            res = self.fx.store.create_edit(deck, {"op": "revise", "pages": ["00000000-0000-4000-8000-000000000000"],
                                                   "text": "x", "source": "screen", "positions": [99]})
        self.fx.server.wake_worker()
        run = self.wait(deck)
        self.assertEqual((run["id"], run["state"]), (res["run"]["id"], "failed"))
        self.assertIn("もうありません", run["message"])
        self.assertEqual(len(self.view(deck)["pages"]), len(pages))

    def test_writer_answer_outside_the_targets_is_sent_back(self) -> None:
        self.tools.extra_env["FAKE_WRITER_MODE"] = "bad_revision_first"
        deck = self.make_deck()
        self.instruct(deck, [7], "見出しを短くして")
        run = self.wait(deck)
        self.assertEqual(run["state"], "done", run["message"])
        calls = [c for c in writer_calls(self.log) if c["what"] == "revise"]
        self.assertEqual(len(calls), 2)
        self.assertTrue(any("直すページではない" in ln for ln in calls[1]["ng"]), calls[1]["ng"])

    def test_split_page_inserts_new_ids_in_place(self) -> None:
        deck = self.make_deck()
        calls = self.codex_calls()
        self.instruct(deck, [7], "2枚に分けて")
        self.assertEqual(self.wait(deck)["state"], "done")
        keys = [p["key"] for p in self.view(deck)["pages"]]
        self.assertEqual(keys[6:8], ["p-seven", "p-seven-b"])
        self.assertEqual(self.drawn_pages(calls), [7, 8], "分けた 2 ページだけ描く（後ろのページは付け替え）")
        self.assert_screen_matches_pptx(deck)


class PresentationEditTests(EditBase):
    mode = "presentation"
    runtime = "claude_code"

    @classmethod
    def setUpClass(cls) -> None:
        why = html_slide_ready()
        if why:
            raise unittest.SkipTest(why)

    def gen_files(self) -> dict[str, float]:
        root = Path(self.tools.extra_env["HTML_SLIDE_CACHE"])
        return {str(p): p.stat().st_ino for p in root.rglob("gen-*.png") if "history" not in p.parts}

    def test_duplicate_copies_pages_and_rebuild_draws_nothing(self) -> None:
        deck = self.make_deck()
        gens = len(self.gen_files())
        new = self.duplicate_and_rebuild(deck)
        self.assertEqual(len(self.gen_files()), gens * 2, "複製の区画に同じ絵が張られた（描いてはいない）")
        with self.fx.server.store_lock:
            ra, rb = self.fx.store.shown_run(deck), self.fx.store.shown_run(new)
        ca = json.loads((self.fx.root.run_dir(deck, ra["id"]) / "deck-info.json").read_text())["cache"]
        cb = json.loads((self.fx.root.run_dir(new, rb["id"]) / "deck-info.json").read_text())["cache"]
        self.assertNotEqual(ca, cb, "複製の回の記録は複製の区画を指す")
        a_files = {p.name: (p.stat().st_ino, p.read_bytes()) for p in Path(ca).glob("gen-*")}
        a_pics = {p.name: p.read_bytes() for p in (self.fx.root.run_dir(deck, ra["id"]) / "pictures").glob("*")}
        # 複製の 8 枚目の絵だけ直す → 複製の最初の版へ戻す。元のデッキの区画と控えは 1 バイトも変わらない
        self.instruct(new, [8], "絵だけ夜にして")
        self.assertEqual(self.wait(new)["state"], "done")
        p8 = next(p for p in self.view(new)["pages"] if p["key"] == "p-eight")
        r = self.fx.post(f"/api/decks/{new}/adopt", {"choices": {p8["slide_id"]: p8["versions"][-1]["id"]}})
        self.assertEqual(r.status, 200, r.body)
        self.assertEqual(self.wait(new)["state"], "done")
        self.assertEqual({p.name: (p.stat().st_ino, p.read_bytes()) for p in Path(ca).glob("gen-*")}, a_files,
                         "複製の版を戻しても、元のデッキの区画は変わらない")
        self.assertEqual({p.name: p.read_bytes() for p in (self.fx.root.run_dir(deck, ra["id"]) / "pictures").glob("*")},
                         a_pics, "複製の直しで、元の回の控え（絵に添えた JSON を含む）は変わらない")
        for path in (f"/api/decks/{deck}/delete", f"/api/decks/{deck}/purge"):   # 元を完全に消しても複製は残る
            r = self.fx.post(path, {})
            self.assertEqual(r.status, 200, r.body)
        for p in self.view(new)["pages"]:
            self.assertEqual(self.fx.get(p["image_url"]).status, 200)
        self.assertEqual(len(self.gen_files()), gens, "元の区画だけが消え、複製の区画は残る")
        calls = self.codex_calls()
        self.assertEqual(self.fx.post(f"/api/decks/{new}/rebuild", {}).status, 200)
        self.assertEqual(self.wait(new)["state"], "done")
        self.assertEqual(self.drawn_pages(calls), [], "元を消したあとも、複製は描き直さない")

    def test_revise_7_and_8_keeps_other_slots_and_restore(self) -> None:
        deck = self.make_deck()
        before_data, before_bytes = self.head(deck)
        before_pieces = self.page_bytes(before_bytes, before_data)
        gens = self.gen_files()
        self.assertEqual(len(gens), 4, "画像の枠は 1・5・7・8 枚目")
        calls = self.codex_calls()
        self.instruct(deck, [7, 8], "絵も変えて 見出しを短くして")
        run = self.wait(deck)
        self.assertEqual(run["state"], "done", run["message"])
        after_data, after_bytes = self.head(deck)
        after_pieces = self.page_bytes(after_bytes, after_data)
        for key in before_pieces:
            if key not in ("p-seven", "p-eight"):
                self.assertEqual(before_pieces[key], after_pieces[key], key)
        now = self.gen_files()
        kept = {k: v for k, v in gens.items() if k in now and now[k] == v}
        self.assertEqual(len(kept), 4, "前の枠の絵はそのまま（同じ inode）")
        self.assertEqual(len(now) - len(gens), 2, "描いたのは 7・8 枚目の枠だけ")
        execs = [n for n in self.codex_calls() - calls
                 if (lambda r: r.get("cmd") == "exec" and "t1" in r)(json.loads((self.log / "codex-calls" / n).read_text(encoding="utf-8")))]
        self.assertEqual(len(execs), 2)
        self.assert_screen_matches_pptx(deck)
        # 絵だけの直し（revise）→ 戻す
        p8 = next(p for p in self.view(deck)["pages"] if p["key"] == "p-eight")
        self.assertTrue(p8["can_minor"]["available"])
        self.instruct(deck, [8], "絵だけ夜にして")
        run = self.wait(deck)
        self.assertEqual(run["state"], "done", run["message"])
        self.assertEqual([a[2] for a in run["record"]["argv"]], ["revise"])
        p8 = next(p for p in self.view(deck)["pages"] if p["key"] == "p-eight")
        self.assertEqual([v["version"] for v in p8["versions"]], [3, 2, 1])
        calls = self.codex_calls()
        r = self.fx.post(f"/api/decks/{deck}/adopt", {"choices": {p8["slide_id"]: p8["versions"][1]["id"]}})
        self.assertEqual(r.status, 200, r.body)
        run = self.wait(deck)
        self.assertEqual(run["state"], "done", run["message"])
        self.assertEqual(self.codex_calls() - calls, set(), "戻すときは描かない")
        p8 = next(p for p in self.view(deck)["pages"] if p["key"] == "p-eight")
        self.assertEqual((p8["versions"][0]["version"], p8["versions"][0]["active"]), (2, True))
        self.assert_screen_matches_pptx(deck)


if __name__ == "__main__":
    unittest.main()
