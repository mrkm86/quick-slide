"""版は切り替えるだけ（docs/EDIT-PLAN.md の C・docs/adr/0001-version-pins.md）。

- 前の版を選ぶと、答えが返った時点で画面はその版の絵（描き直さない）。台本のそのページは前の版の中身と `_pin`
- 裏の組み直し（tidy）が終わっても版は増えない。書き出しの、そのページはその版のもの
  - ピッチ: その版の絵を区画に置くので、エンジンが同じ絵で PPTX まで組む
  - プレゼン: エンジンが変わっても（ここではスタイルを替えて描かせる）、ページの画像・PPTX の部品・PDF のページは、
    その版を描いた回のもの。発表はそのページだけ、その回の HTML の区切りで見せる
  - 書き出しの部品が無い回（この仕組みの前の回）の版は、画面と発表はその版の絵、書き出しだけいまの作りで描き、注を出す
- 印はページを直すと外れ、非表示・並べ替えでは残り、組み直し（rebuild）で全部外れる。複製の写しは印を持たない
"""

from __future__ import annotations

import shutil
import sys
import unittest
from pathlib import Path
from unittest import mock

from helpers import REPO

from quick_slide.engines import Engine
from quick_slide.server.vocabulary import PIN_EXPORT_REDRAW_NOTE

from test_engines_e2e import html_slide_ready
from p3_edit_helpers import EditBase

sys.path.insert(0, str(REPO / "skills" / "html-slide" / "scripts"))
import pdfpages  # noqa: E402


def slide_pictures(path: Path, no: int) -> list[bytes]:
    """PPTX の no 枚目（1 から）の画像のバイト。"""
    from pptx import Presentation

    slide = Presentation(str(path)).slides[no - 1]
    return [sh.image.blob for sh in slide.shapes if sh.shape_type == 13]


def slide_texts(path: Path, no: int) -> list[str]:
    from pptx import Presentation

    out = []
    for sh in Presentation(str(path)).slides[no - 1].shapes:
        if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
            out.append(sh.text_frame.text.strip())
    return out


class PinBase(EditBase):
    def page(self, deck: str, key: str, view: dict | None = None) -> dict:
        return next(p for p in (view or self.view(deck))["pages"] if p["key"] == key)

    def pin_of(self, deck: str, key: str) -> str | None:
        data, _ = self.head(deck)
        return next(s for s in data["slides"] if s["id"] == key).get("_pin")

    def shown(self, deck: str) -> dict:
        with self.fx.server.store_lock:
            return self.fx.store.shown_run(deck)

    def run_dir(self, deck: str, run_id: str) -> Path:
        return self.fx.root.run_dir(deck, run_id)

    def asset_sha(self, url: str) -> str:
        import hashlib

        return hashlib.sha256(self.fx.get(url).body).hexdigest()

    def adopt(self, deck: str, page: dict, version: dict) -> dict:
        r = self.fx.post(f"/api/decks/{deck}/adopt", {"choices": {page["slide_id"]: version["id"]}})
        self.assertEqual(r.status, 200, r.body)
        return r.json()


class PitchPinTests(PinBase):
    mode = "pitch"

    def test_adopt_switches_to_the_old_picture_without_a_new_version(self) -> None:
        deck = self.make_deck()
        self.instruct(deck, [7], "見出しを短くして")
        self.assertEqual(self.wait(deck)["state"], "done")
        p7 = self.page(deck, "p-seven")
        self.assertEqual([v["version"] for v in p7["versions"]], [2, 1])
        v1 = p7["versions"][1]
        calls = self.codex_calls()
        body = self.adopt(deck, p7, v1)
        # 答えが返った時点で、画面はもう v1 の絵（組み直しを待たない）
        self.assertTrue(body["applied"])
        now = self.page(deck, "p-seven", body)
        self.assertEqual((now["image_url"], now["version"], now["display"]["key"], now["pinned"]),
                         (v1["image_url"], 1, "ready", True))
        self.assertEqual(self.pin_of(deck, "p-seven"), v1["id"])
        self.assertEqual(self.wait(deck)["state"], "done")
        self.assertEqual(self.drawn_pages(calls), [], "切り替えでは描かない")
        p7 = self.page(deck, "p-seven")
        self.assertEqual([v["version"] for v in p7["versions"]], [1, 2], "組み直しの後も版は 2 つ（v1 が新しい回を含む）")
        self.assertEqual((p7["version"], p7["versions"][0]["active"], p7["image_url"]), (1, True, v1["image_url"]))
        self.assertIsNone(p7["export_note"], "ピッチは区画の絵から組むので、書き出しも v1 の絵")
        shown = self.shown(deck)
        self.assertNotEqual(shown["id"], v1["id"])
        self.assertEqual(slide_pictures(self.run_dir(deck, shown["id"]) / "deck.pptx", 7),
                         slide_pictures(self.run_dir(deck, v1["id"]) / "deck.pptx", 7), "PPTX の 7 枚目は v1 の絵")
        self.assert_screen_matches_pptx(deck)

    def save_brand(self, colors: dict, brand_id: str | None = None) -> str:
        """付属のブランドの設定から色だけ替えたブランドを作る（`brand_id` を渡せば、その新しい改訂にする）。"""
        import copy

        store = self.fx.store
        with self.fx.server.store_lock:
            base = store.brand_revision(store.preset_brand_revision_id())["config"]
            config = copy.deepcopy(base)
            config["colors"] = {**config["colors"], **colors}
            if brand_id is None:
                return store.create_brand("試験のブランド（架空）", config)["id"]
            store.add_brand_revision(brand_id, config)
            return brand_id

    def test_adopt_after_a_brand_change_keeps_the_old_picture(self) -> None:
        """ブランドの色を替えた後（組み直しは頼まない）に v1 を選んでも、7 枚目は描き直さずに v1 の絵で組む。
        描く指示（鍵）は色で変わるので、置いた絵を鍵が合わなくても使わせる（image-slide の --keep）。"""
        brand = self.save_brand({"primary": "#1D4ED8", "accent": "#B45309"})
        deck = self.make_deck(brand_id=brand)
        self.instruct(deck, [7], "見出しを短くして")
        self.assertEqual(self.wait(deck)["state"], "done")
        p7 = self.page(deck, "p-seven")
        v1 = p7["versions"][1]
        self.save_brand({"primary": "#0F766E", "accent": "#7C3AED"}, brand)   # 色を替える（組み直しは頼まない）
        calls = self.codex_calls()
        self.adopt(deck, p7, v1)
        run = self.wait(deck)
        self.assertEqual(run["state"], "done", run["message"])
        self.assertIn("--keep", run["record"]["argv"][0])
        self.assertNotIn(7, self.drawn_pages(calls), "版を切り替えたページは、色が変わっても描き直さない")
        self.assertTrue(self.drawn_pages(calls), "ほかのページは新しい色の描く指示で描き直した（色の変更が効いている）")
        p7 = self.page(deck, "p-seven")
        self.assertEqual((len(p7["versions"]), p7["version"]), (2, 1), "版は増えない")
        shown = self.shown(deck)
        self.assertEqual(slide_pictures(self.run_dir(deck, shown["id"]) / "deck.pptx", 7),
                         slide_pictures(self.run_dir(deck, v1["id"]) / "deck.pptx", 7), "PPTX の 7 枚目は v1 の絵")

    def test_the_pin_follows_the_page(self) -> None:
        deck = self.make_deck()
        self.instruct(deck, [7], "見出しを短くして")
        self.assertEqual(self.wait(deck)["state"], "done")
        p7 = self.page(deck, "p-seven")
        v1 = p7["versions"][1]
        self.adopt(deck, p7, v1)
        # 非表示・並べ替えでは残る
        other = self.page(deck, "p-why")
        self.assertEqual(self.fx.post(f"/api/decks/{deck}/pages/{other['slide_id']}/visibility", {"hidden": True}).status, 200)
        ids = [p["slide_id"] for p in self.view(deck)["pages"]]
        self.assertEqual(self.fx.post(f"/api/decks/{deck}/pages/reorder", {"slide_ids": ids[::-1]}).status, 200)
        self.assertEqual(self.wait(deck)["state"], "done")
        self.assertEqual(self.pin_of(deck, "p-seven"), v1["id"], "非表示・並べ替えでは版の印は残る")
        self.assertEqual(self.page(deck, "p-seven")["version"], 1)
        # 複製の写しは印を持たない（元は持ったまま）
        p7 = self.page(deck, "p-seven")
        self.assertEqual(self.fx.post(f"/api/decks/{deck}/pages/{p7['slide_id']}/duplicate", {}).status, 200)
        self.assertEqual(self.wait(deck)["state"], "done")
        self.assertIsNone(self.pin_of(deck, "p-seven-copy"))
        self.assertEqual(self.pin_of(deck, "p-seven"), v1["id"])
        # 直すと外れる
        pos = self.page(deck, "p-seven")["position"]
        self.instruct(deck, [pos], "見出しを短くして")
        self.assertEqual(self.wait(deck)["state"], "done")
        self.assertIsNone(self.pin_of(deck, "p-seven"), "直したページは新しい姿（前の版を指さない）")
        # 組み直しで全部外れる
        p7 = self.page(deck, "p-seven")
        self.adopt(deck, p7, next(v for v in p7["versions"] if v["version"] == 1))
        self.assertEqual(self.wait(deck)["state"], "done")
        self.assertTrue(self.pin_of(deck, "p-seven"))
        self.assertEqual(self.fx.post(f"/api/decks/{deck}/rebuild", {}).status, 200)
        self.assertEqual(self.wait(deck)["state"], "done")
        self.assertIsNone(self.pin_of(deck, "p-seven"), "組み直しは版の印を外す")
        self.assert_screen_matches_pptx(deck)


class PresentationPinTests(PinBase):
    mode = "presentation"
    runtime = "claude_code"

    @classmethod
    def setUpClass(cls) -> None:
        why = html_slide_ready()
        if why:
            raise unittest.SkipTest(why)

    @staticmethod
    def engine_changed():
        """エンジンが変わったことにする（プレゼンだけ、スタイルを visual に替えて描かせる。組み直しは頼まない）。"""
        original = Engine.style_args

        def changed(self, args):
            out = original(self, args)
            if self.mode == "presentation" and "--style" in out:
                out[out.index("--style") + 1] = "visual"
            return out

        return mock.patch.object(Engine, "style_args", changed)

    def night_version(self, deck: str) -> tuple[dict, dict]:
        """8 枚目の絵だけ直して v2 を作る。(8 枚目, v1) を返す。"""
        self.instruct(deck, [8], "絵だけ夜にして")
        self.assertEqual(self.wait(deck)["state"], "done")
        p8 = self.page(deck, "p-eight")
        self.assertEqual([v["version"] for v in p8["versions"]], [2, 1])
        return p8, p8["versions"][1]

    def test_adopt_keeps_the_old_outputs_when_the_engine_changed(self) -> None:
        deck = self.make_deck()
        p8, v1 = self.night_version(deck)
        v1_dir = self.run_dir(deck, v1["id"])
        self.assertTrue((v1_dir / "parts" / "p-eight.json").is_file(), "描いた回はページごとの書き出しの部品を残す")
        before = {p["key"]: self.asset_sha(p["image_url"]) for p in self.view(deck)["pages"]}
        with self.engine_changed():
            body = self.adopt(deck, p8, v1)
            self.assertTrue(body["applied"])
            now = self.page(deck, "p-eight", body)
            self.assertEqual((now["image_url"], now["version"], now["pinned"], now["export_note"]),
                             (v1["image_url"], 1, True, None))
            present = self.fx.get(f"/api/decks/{deck}/present").json()
            entry = next(p for p in present["pages"] if p["slide_id"] == p8["slide_id"])
            self.assertEqual(entry["html"]["url"], f"/present/{deck}/{v1['id']}/deck.html",
                             "発表はそのページだけ、v1 の回の HTML の区切りで見せる")
            self.assertEqual(self.wait(deck)["state"], "done")
        view = self.view(deck)
        after = {p["key"]: self.asset_sha(p["image_url"]) for p in view["pages"]}
        self.assertNotEqual({k for k in after if after[k] != before[k]} - {"p-eight"}, set(),
                            "ほかのページは変わったエンジン（スタイル）で描き直された")
        p8 = self.page(deck, "p-eight", view)
        self.assertEqual(after["p-eight"], self.asset_sha(v1["image_url"]), "8 枚目は v1 の絵のまま")
        self.assertEqual((len(p8["versions"]), p8["version"]), (2, 1), "版は増えない")
        shown = self.shown(deck)
        new_dir = self.run_dir(deck, shown["id"])
        self.assertEqual(shown["record"]["pinned"]["exact"], ["p-eight"])
        self.assertTrue(shown["record"]["pinned"]["assembled"])
        # 書き出し: 8 枚目は v1 の部品（PPTX）と v1 の PDF のページ。ほかのページはいまの作り
        self.assertEqual(slide_texts(new_dir / "deck.pptx", 8), slide_texts(v1_dir / "deck.pptx", 8))
        self.assertEqual(slide_pictures(new_dir / "deck.pptx", 8), slide_pictures(v1_dir / "deck.pptx", 8))
        self.assertEqual(pdfpages.page_streams(new_dir / "deck.pdf", 8), pdfpages.page_streams(v1_dir / "deck.pdf", 8),
                         "PDF の 8 ページ目は v1 の PDF のページを 1 バイトも変えずに写したもの")
        self.assertNotEqual(pdfpages.page_streams(new_dir / "deck.pdf", 2), pdfpages.page_streams(v1_dir / "deck.pdf", 2))
        self.assertEqual(pdfpages.page_count(new_dir / "deck.pdf"), 10)
        self.assertTrue((new_dir / "parts" / "p-eight.json").is_file(), "持ち込んだ部品は新しい回にも残す（次の切り替えの元）")
        present = self.fx.get(f"/api/decks/{deck}/present").json()
        self.assertTrue(present["html"])
        urls = {p["slide_id"]: p["html"]["url"] for p in present["pages"]}
        self.assertEqual(urls[p8["slide_id"]], f"/present/{deck}/{v1['id']}/deck.html")
        self.assertEqual({u for sid, u in urls.items() if sid != p8["slide_id"]}, {f"/present/{deck}/{shown['id']}/deck.html"})
        self.assertEqual(self.fx.get(urls[p8["slide_id"]]).status, 200)
        self.assert_screen_matches_pptx(deck)

    def test_a_version_without_export_parts_redraws_only_the_export(self) -> None:
        deck = self.make_deck()
        p8, v1 = self.night_version(deck)
        shutil.rmtree(self.run_dir(deck, v1["id"]) / "parts")   # この仕組みの前の回と同じ（部品が無い）
        body = self.adopt(deck, p8, v1)
        now = self.page(deck, "p-eight", body)
        self.assertEqual((now["image_url"], now["pinned"], now["export_note"]), (v1["image_url"], True, PIN_EXPORT_REDRAW_NOTE))
        self.assertEqual(self.wait(deck)["state"], "done")
        p8 = self.page(deck, "p-eight")
        self.assertEqual((len(p8["versions"]), p8["version"], p8["export_note"]), (2, 1, PIN_EXPORT_REDRAW_NOTE))
        self.assertEqual(self.asset_sha(p8["image_url"]), self.asset_sha(v1["image_url"]), "画面は v1 の絵")
        shown = self.shown(deck)
        self.assertEqual(shown["record"]["pinned"]["exact"], [])
        self.assertIsNone(shown["record"]["pinned"]["assembled"], "部品が無いので組み直さない（エンジンの書き出しのまま）")
        self.assertFalse((self.run_dir(deck, shown["id"]) / "parts" / "p-eight.json").exists(),
                         "絵と食い違う部品は残さない")
        self.assert_screen_matches_pptx(deck)


if __name__ == "__main__":
    unittest.main()
